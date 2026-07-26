import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def add_heading(doc, text, level=1):
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.name = 'Times New Roman'
        run.font.color.rgb = RGBColor(0, 51, 102)

def add_paragraph(doc, text, bold=False, italic=False, style=None):
    if style:
        p = doc.add_paragraph(style=style)
    else:
        p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    if bold:
        run.bold = True
    if italic:
        run.italic = True
    p.paragraph_format.space_after = Pt(6)
    return p

def add_code(doc, code_text):
    p = doc.add_paragraph()
    run = p.add_run(code_text)
    run.font.name = 'Courier New'
    run.font.size = Pt(10)
    # create a light grey background effect by just using font color for simplicity or a box
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(12)

doc = Document()

# Tieu de
title = doc.add_heading('BÁO CÁO NGHIÊN CỨU CHUYÊN SÂU\nTHUẬT TOÁN GRAD-CAM VÀ EXPLAINABLE AI (XAI)', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in title.runs:
    run.font.name = 'Times New Roman'

# Phan 1
add_heading(doc, '1. GIỚI THIỆU VỀ XAI VÀ SỰ CẦN THIẾT CỦA VIỆC GIẢI THÍCH MÔ HÌNH', level=1)
add_paragraph(doc, "Mạng nơ-ron tích chập (CNN) thường bị coi là các 'hộp đen' (Black-box). Khi mô hình dự đoán một lá lúa bị bệnh 'Đạo ôn' (Blast), nó chỉ đưa ra kết quả và xác suất (Confidence), nhưng không giải thích được TẠI SAO nó lại quyết định như vậy. Điều này gây ra sự thiếu tin tưởng khi áp dụng AI vào thực tế nông nghiệp.")
add_paragraph(doc, "Explainable AI (XAI) là lĩnh vực nghiên cứu giúp con người hiểu được lý do đằng sau các quyết định của AI. Trong xử lý ảnh, phương pháp phổ biến nhất là tạo ra các Bản đồ kích hoạt nhiệt (Heatmap / Saliency Map) để làm nổi bật các điểm ảnh (pixels) có ảnh hưởng lớn nhất đến kết quả dự đoán.")

# Phan 2
add_heading(doc, '2. CHI TIẾT VỀ THUẬT TOÁN GRAD-CAM', level=1)
add_heading(doc, '2.1. Grad-CAM là gì?', level=2)
add_paragraph(doc, "Grad-CAM (Gradient-weighted Class Activation Mapping) là thuật toán sử dụng đạo hàm (gradient) của một lớp đích (target class) truyền ngược (backpropagate) đến lớp chập cuối cùng (last convolutional layer) của mạng CNN để tạo ra một bản đồ định vị thô, đánh dấu các vùng quan trọng trên ảnh.")

add_heading(doc, '2.2. Cơ chế hoạt động (Giải thích dễ hiểu)', level=2)
add_paragraph(doc, "- Bước 1: Mạng CNN trích xuất đặc trưng hình ảnh qua nhiều lớp. Lớp chập cuối cùng chứa thông tin không gian (vị trí vết bệnh) và ngữ nghĩa tốt nhất.")
add_paragraph(doc, "- Bước 2: Để biết mạng đang chú ý vào đâu để đoán ra bệnh 'Đạo ôn', Grad-CAM tính đạo hàm của điểm số bệnh Đạo ôn đối với các bộ lọc (filters) ở lớp cuối cùng.")
add_paragraph(doc, "- Bước 3: Đạo hàm này thể hiện 'tầm quan trọng' của từng bộ lọc. Grad-CAM tính trung bình các đạo hàm này để ra một trọng số (weight).")
add_paragraph(doc, "- Bước 4: Nhân các trọng số này với các bản đồ đặc trưng (feature maps) tương ứng, cộng lại, và cắt bỏ phần âm (hàm ReLU) để chỉ giữ lại những chi tiết đóng góp tích cực cho kết quả bệnh Đạo ôn.")

add_heading(doc, '2.3. Công thức Toán học chi tiết', level=2)
add_paragraph(doc, "Giả sử y^c là điểm số (score) của mô hình (trước khi qua softmax) cho lớp c. A^k là bản đồ đặc trưng thứ k của lớp tích chập cuối cùng.")
add_paragraph(doc, "1. Tính trọng số tầm quan trọng (Neuron importance weights) alpha_k^c:")
add_paragraph(doc, "   α_k^c = (1 / Z) * Σ_i Σ_j ( ∂y^c / ∂A^k_ij )", bold=True)
add_paragraph(doc, "   Trong đó:")
add_paragraph(doc, "   - Z là tổng số pixel trong bản đồ đặc trưng (Z = width * height).")
add_paragraph(doc, "   - (1/Z)*ΣΣ chính là phép toán Global Average Pooling (GAP).")
add_paragraph(doc, "   - ∂y^c / ∂A^k_ij là đạo hàm (gradient) của điểm số lớp c theo từng giá trị pixel (i, j) trên bản đồ đặc trưng A^k.")

add_paragraph(doc, "2. Tính Bản đồ kích hoạt L_Grad-CAM^c:")
add_paragraph(doc, "   L_Grad-CAM^c = ReLU ( Σ_k (α_k^c * A^k) )", bold=True)
add_paragraph(doc, "   Trong đó:")
add_paragraph(doc, "   - Tổng tuyến tính của các bản đồ đặc trưng A^k nhân với trọng số α_k^c của chúng.")
add_paragraph(doc, "   - Hàm ReLU chỉ giữ lại các giá trị dương. Vì ta chỉ quan tâm đến các đặc trưng có tác động TÍCH CỰC đến lớp c. Nếu không có ReLU, bản đồ nhiệt sẽ làm nổi bật cả những đặc trưng thuộc về lớp khác (gây nhiễu).")

add_heading(doc, '2.4. Mã nguồn (Code Python / TensorFlow) của Grad-CAM', level=2)
code_gradcam = """
def make_gradcam_heatmap(img_array, model, last_conv_layer_name, pred_index=None):
    # Tạo mô hình trung gian xuất ra Feature Map và Output
    grad_model = tf.keras.models.Model(
        [model.inputs], 
        [model.get_layer(last_conv_layer_name).output, model.output]
    )

    # Theo dõi Gradient
    with tf.GradientTape() as tape:
        last_conv_layer_output, preds = grad_model(img_array)
        if pred_index is None:
            pred_index = tf.argmax(preds[0])
        class_channel = preds[:, pred_index]

    # Tính Gradient của điểm số dự đoán so với Feature Map (∂y^c / ∂A^k)
    grads = tape.gradient(class_channel, last_conv_layer_output)

    # Tính trung bình Gradient (Global Average Pooling) -> Trọng số α_k^c
    pooled_grads = tf.reduce_mean(grads, axis=(0, 1, 2))

    # Nhân Feature Map với trọng số
    last_conv_layer_output = last_conv_layer_output[0]
    heatmap = last_conv_layer_output @ pooled_grads[..., tf.newaxis]
    heatmap = tf.squeeze(heatmap)

    # Áp dụng ReLU (tf.maximum)
    heatmap = tf.maximum(heatmap, 0) / tf.math.reduce_max(heatmap)
    return heatmap.numpy()
"""
add_code(doc, code_gradcam)

doc.add_page_break()

# Phan 3
add_heading(doc, '3. SO SÁNH GRAD-CAM VỚI CÁC THUẬT TOÁN XAI KHÁC', level=1)

# CAM
add_heading(doc, '3.1. CAM (Class Activation Mapping)', level=2)
add_paragraph(doc, "- Cơ chế: Bắt buộc mô hình phải có cấu trúc Global Average Pooling (GAP) ngay trước lớp phân loại cuối cùng. Lấy trọng số của lớp Dense cuối cùng nhân với Feature Maps.")
add_paragraph(doc, "- Hạn chế: Phải thay đổi kiến trúc mô hình (xóa lớp Fully Connected, thêm GAP layer) và train lại từ đầu. Đánh đổi độ chính xác để lấy khả năng giải thích.")

# Grad-CAM
add_heading(doc, '3.2. Grad-CAM (Nghiên cứu hiện tại)', level=2)
add_paragraph(doc, "- Ưu điểm: KHÔNG cần thay đổi kiến trúc mô hình, áp dụng được ngay (Plug-and-play) cho mọi mô hình CNN có sẵn. Dùng đạo hàm thay vì trọng số lớp Dense.")
add_paragraph(doc, "- Hạn chế: Do dùng Global Average Pooling trên đạo hàm, Grad-CAM tính toán trọng số trung bình. Điều này khiến nó bị kém hiệu quả khi trên ảnh có NHIỀU vết bệnh phân tán (multiple occurrences of the same class). Nó thường chỉ nổi bật được một vết to nhất và làm nhiễu / lan mờ các vết xung quanh.")

# Grad-CAM++
add_heading(doc, '3.3. Grad-CAM++ (Giải pháp nâng cấp sắp tới)', level=2)
add_paragraph(doc, "- Cơ chế: Sử dụng đạo hàm bậc 1, bậc 2 và bậc 3 (first, second, third-order gradients) để tính trọng số thay vì chỉ tính trung bình đơn giản.")
add_paragraph(doc, "- Ưu điểm tuyệt đối: Giải quyết được nhược điểm của Grad-CAM. Nội suy cực tốt khi trong ảnh có nhiều cá thể / nhiều vết bệnh phân tán. Bản đồ nhiệt sắc nét hơn, gom gọn vào đúng đối tượng, giảm tỷ lệ loang/nhiễu (ActivArea nhỏ hơn).")

# Score-CAM
add_heading(doc, '3.4. Score-CAM', level=2)
add_paragraph(doc, "- Cơ chế: Không dùng đạo hàm (gradient-free). Chạy lan truyền xuôi (forward pass) nhiều lần trên các Feature map để lấy điểm số (Score).")
add_paragraph(doc, "- Hạn chế: Chạy rất chậm (tốn tài nguyên tính toán gấp hàng chục lần Grad-CAM) vì phải forward pass liên tục. Không phù hợp triển khai trên Web/App theo thời gian thực.")

# LIME & SHAP
add_heading(doc, '3.5. LIME và SHAP', level=2)
add_paragraph(doc, "- LIME / SHAP: Dùng kỹ thuật che giấu từng phần bức ảnh (perturbation) để xem mô hình thay đổi kết quả thế nào.")
add_paragraph(doc, "- Hạn chế: Tốc độ suy luận cực kỳ chậm. LIME mất đến vài giây hoặc vài chục giây cho 1 ảnh. Không phù hợp với bài toán cần xem Heatmap tức thì (Real-time).")


# Phan 4
add_heading(doc, '4. TẠI SAO GRAD-CAM VÀ GRAD-CAM++ LÀ LỰA CHỌN TỐT NHẤT CHO DỰ ÁN NÀY?', level=1)
add_paragraph(doc, "1. Yêu cầu thời gian thực (Real-time Web API):")
add_paragraph(doc, "   Hệ thống yêu cầu khi nông dân upload ảnh lá lúa lên web, hệ thống phải trả kết quả bệnh + ảnh nội soi (Heatmap) gần như ngay lập tức. Grad-CAM tính toán trong 1 lần Forward + Backward pass (chưa tới 0.1s). Score-CAM, LIME, hay SHAP quá chậm, làm nghẽn server.")

add_paragraph(doc, "2. Không can thiệp vào Model Accuracy:")
add_paragraph(doc, "   Bài toán bệnh lá lúa cần độ chính xác rất cao. Nếu dùng thuật toán CAM truyền thống, ta buộc phải sửa kiến trúc model và train lại, gây giảm Accuracy. Grad-CAM dùng nguyên vẹn file `best_stage2.keras` đang có độ chính xác cao nhất.")

add_paragraph(doc, "3. Đặc thù bệnh cây lúa (Lý do bắt buộc phải lên Grad-CAM++):")
add_paragraph(doc, "   Bệnh trên lá lúa (như Đạo ôn, Đốm nâu, Đốm vằn) thường không phải là một vết to ở giữa ảnh, mà là RẤT NHIỀU VẾT NHỎ nằm rải rác khắp bề mặt lá. Grad-CAM cũ gặp khó khăn với nhiều vết đốm nhỏ (hay bị loang nhiệt độ, nhiễu). Do đó, việc áp dụng phương pháp tiên tiến hơn là Grad-CAM++ (giúp khoanh vùng đa mục tiêu sắc nét) là cực kỳ phù hợp và mang tính cấp thiết về học thuật lẫn thực tiễn trong đồ án này.")

output_path = r"D:\Phát hiện bệnh cây lúa\research\Nghien_Cuu_Chuyen_Sau_GradCAM_va_XAI.docx"
doc.save(output_path)
print(f"Đã xuất file Word chuyên sâu thành công tại: {output_path}")
