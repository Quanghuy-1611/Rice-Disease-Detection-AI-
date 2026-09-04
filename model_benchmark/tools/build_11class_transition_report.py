from __future__ import annotations

import csv, json, statistics, sys, types
from collections import Counter
from pathlib import Path

import numpy as np
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

# The bundled document runtime intentionally omits matplotlib.  The imported
# report module is used only for its Word-formatting helpers, so provide a
# no-op pyplot surface before import; all figures for this report are produced
# separately with the bundled Node/sharp runtime.
_mpl=types.ModuleType('matplotlib'); _mpl.use=lambda *_a,**_k: None
_plt=types.ModuleType('matplotlib.pyplot'); _plt.rcParams={}
sys.modules.setdefault('matplotlib',_mpl); sys.modules.setdefault('matplotlib.pyplot',_plt)
import build_final_results_report_docx as base

ROOT = Path(__file__).resolve().parents[1]
FINAL = ROOT / 'runs' / 'final_test_v2'
SPLIT = ROOT.parent / 'research' / 'data' / 'splits' / 'rice_cleaned_v4_exact_cross_label_conflicts_excluded_v1.csv'
OUT = ROOT / 'docs' / 'BAO_CAO_SO_SANH_6_MO_HINH_DINH_HUONG_11_LOP_V2_CHI_TIET_THUOC_DO.docx'
ASSETS = ROOT / 'docs' / 'results_report_assets_11class'
DATA_JSON = ASSETS / 'recalculated_11class_metrics.json'

MODELS = base.MODEL_ORDER
NAMES = base.MODEL_NAMES
SHORT = base.MODEL_SHORT
COLORS = base.MODEL_COLORS
SEEDS = [42, 123, 2026]
CORE = [
    'bacterial_leaf_blight', 'blast', 'brown_planthopper', 'brown_spot',
    'dead_heart', 'golden_apple_snail', 'leaf_roller', 'normal',
    'stem_borer', 'tungro', 'weeds',
]
VI = {
    'bacterial_leaf_blight': 'Bạc lá vi khuẩn', 'blast': 'Đạo ôn',
    'brown_planthopper': 'Rầy nâu', 'brown_spot': 'Đốm nâu',
    'dead_heart': 'Dảnh héo', 'golden_apple_snail': 'Ốc bươu vàng',
    'leaf_roller': 'Sâu cuốn lá', 'normal': 'Bình thường',
    'stem_borer': 'Sâu đục thân', 'tungro': 'Tungro', 'weeds': 'Cỏ dại',
}
ABBR = ['BLB','Blast','BPH','BSpot','DeadH','GAS','LR','Normal','StemB','Tungro','Weeds']
IN_FEATURES = {'mobilenetv2':1280,'mobilenetv3large':1280,'resnet18':512,'efficientnetb0':1280,'shufflenetv2_1_0x':1024,'densenet121':1024}
PARAMS_11 = {m: base.PARAMS[m] - 8*(IN_FEATURES[m]+1) for m in MODELS}

def metric(y, p):
    per=[]
    for c in range(len(CORE)):
        tp=int(np.sum((y==c)&(p==c))); fp=int(np.sum((y!=c)&(p==c))); fn=int(np.sum((y==c)&(p!=c)))
        pr=tp/(tp+fp) if tp+fp else 0.0; rc=tp/(tp+fn) if tp+fn else 0.0
        f1=2*tp/(2*tp+fp+fn) if 2*tp+fp+fn else 0.0
        per.append({'class_name':CORE[c],'precision':pr,'recall':rc,'f1':f1,'support':int(np.sum(y==c))})
    return {'accuracy':float(np.mean(y==p)), 'macro_precision':float(np.mean([x['precision'] for x in per])),
            'macro_recall':float(np.mean([x['recall'] for x in per])), 'macro_f1':float(np.mean([x['f1'] for x in per])), 'per_class':per}

def load_predictions(model, seed):
    path=FINAL/model/f'seed_{seed}'/'test_predictions.csv'
    rows=[]
    with path.open(encoding='utf-8', newline='') as f:
        for r in csv.DictReader(f):
            if r['true_label'] in CORE: rows.append(r)
    idx={c:i for i,c in enumerate(CORE)}
    y=np.array([idx[r['true_label']] for r in rows],dtype=np.int16)
    p=np.array([idx.get(r['predicted_label'],-1) for r in rows],dtype=np.int16)
    paths=[r['relative_path'] for r in rows]
    return y,p,paths

def bootstrap_delta(y, pa, pb, seed, n=2000):
    rng=np.random.default_rng(11000+seed)
    observed=metric(y,pa)['macro_f1']-metric(y,pb)['macro_f1']
    vals=np.empty(n,dtype=float)
    for i in range(n):
        ix=rng.integers(0,len(y),len(y)); vals[i]=metric(y[ix],pa[ix])['macro_f1']-metric(y[ix],pb[ix])['macro_f1']
    lo,hi=np.percentile(vals,[2.5,97.5])
    return {'observed_delta':float(observed),'ci95_lower':float(lo),'ci95_upper':float(hi),'bootstrap_samples':n}

def build_data():
    runs={}; arrays={}; supports={s:0 for s in CORE}; split_counts={s:{'train':0,'validation':0,'test':0} for s in CORE}
    with SPLIT.open(encoding='utf-8',newline='') as f:
        for r in csv.DictReader(f):
            if r['label'] in split_counts: split_counts[r['label']][r['split']]+=1
    for model in MODELS:
        runs[model]={}; arrays[model]={}
        for seed in SEEDS:
            y,p,paths=load_predictions(model,seed); arrays[model][seed]=(y,p)
            runs[model][seed]=metric(y,p)
            for e in runs[model][seed]['per_class']: supports[e['class_name']]=e['support']
    aggregate={}
    for model in MODELS:
        aggregate[model]={}
        for k in ('accuracy','macro_precision','macro_recall','macro_f1'):
            vals=[runs[model][s][k] for s in SEEDS]
            aggregate[model][k]={'mean':statistics.mean(vals),'sample_std':statistics.stdev(vals),'values_by_seed':{str(s):runs[model][s][k] for s in SEEDS}}
    paired={}
    for other in [m for m in MODELS if m!='mobilenetv3large']:
        key=f'mobilenetv3large_minus_{other}'; paired[key]={}
        for seed in SEEDS:
            y,pa=arrays['mobilenetv3large'][seed]; y2,pb=arrays[other][seed]
            assert np.array_equal(y,y2)
            paired[key][str(seed)]=bootstrap_delta(y,pa,pb,seed)
    per_class={m:{} for m in MODELS}
    for m in MODELS:
        for c in CORE:
            vals=[next(x for x in runs[m][s]['per_class'] if x['class_name']==c)['f1'] for s in SEEDS]
            per_class[m][c]={'mean':statistics.mean(vals),'sample_std':statistics.stdev(vals),'values_by_seed':{str(s):v for s,v in zip(SEEDS,vals)}}
    data={'schema_version':'riceai-11class-transition/v1','status':'posthoc_sensitivity_from_19class_checkpoints',
          'warning':'Không phải kết quả của classifier 11 đầu ra; campaign 11 lớp chưa huấn luyện.',
          'class_selection_rule':'Giữ lớp có support test >= 100 trong split V2 đã khóa.',
          'classes':CORE,'class_names_vi':VI,'support_test':supports,'split_counts':split_counts,'test_sample_count':sum(supports.values()),
          'runs':runs,'aggregate':aggregate,'per_class':per_class,'paired_bootstrap_vs_mobilenetv3large':paired,'projected_parameters_11head':PARAMS_11}
    ASSETS.mkdir(parents=True,exist_ok=True); DATA_JSON.write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')
    return data

def save(fig,name):
    p=ASSETS/name; fig.savefig(p,dpi=360,bbox_inches='tight',pad_inches=.08,facecolor='white'); plt.close(fig); return p

def figures(d):
    agg=d['aggregate']; y=np.arange(len(MODELS))
    fig,axes=plt.subplots(2,1,figsize=(6.35,6.2),constrained_layout=True)
    for ax,k,title,lim in [(axes[0],'macro_f1','A. Macro-F1 trên 11 lớp đủ support',(90,98.5)),(axes[1],'accuracy','B. Accuracy trên 9.376 ảnh thuộc 11 lớp',(96,99.6))]:
        means=np.array([agg[m][k]['mean']*100 for m in MODELS]); sds=np.array([agg[m][k]['sample_std']*100 for m in MODELS])
        ax.barh(y,means,xerr=sds,color=[COLORS[m] for m in MODELS],capsize=3,height=.66); ax.set_yticks(y,[NAMES[m] for m in MODELS]); ax.invert_yaxis(); ax.set_xlim(*lim); ax.set_title(title,fontweight='bold'); ax.grid(axis='x',alpha=.2)
        for yi,v in enumerate(means): ax.text(lim[1]-.05,yi,f'{v:.3f}%',ha='right',va='center',fontsize=8,fontweight='bold',bbox={'facecolor':'white','alpha':.85,'edgecolor':'none'})
    fig.suptitle('Phân tích chuyển tiếp 11 lớp từ checkpoint 19 lớp',fontweight='bold'); overall=save(fig,'fig_11_overall_metrics.png')

    fig,ax=plt.subplots(figsize=(6.35,4.6),constrained_layout=True)
    for m in MODELS: ax.plot(SEEDS,[agg[m]['macro_f1']['values_by_seed'][str(s)]*100 for s in SEEDS],marker='o',label=NAMES[m],color=COLORS[m])
    ax.set_xticks(SEEDS); ax.set_ylabel('Macro-F1 (%)'); ax.set_xlabel('Training seed'); ax.grid(alpha=.22); ax.legend(ncol=2,fontsize=7.7); ax.set_title('Độ ổn định qua ba seed — phân tích 11 lớp',fontweight='bold'); stability=save(fig,'fig_11_seed_stability.png')

    fig,ax=plt.subplots(figsize=(6.35,4.25),constrained_layout=True)
    vals=[d['support_test'][c] for c in CORE]; ax.barh(np.arange(11),vals,color='#2E74B5'); ax.set_yticks(np.arange(11),[VI[c] for c in CORE]); ax.invert_yaxis(); ax.set_xlabel('Số ảnh test'); ax.set_title('Support test của 11 lớp được giữ lại',fontweight='bold'); ax.grid(axis='x',alpha=.2)
    for i,v in enumerate(vals): ax.text(v+25,i,f'{v:,}'.replace(',','.'),va='center',fontsize=8); support=save(fig,'fig_11_support.png')

    mat=np.array([[d['per_class'][m][c]['mean']*100 for c in CORE] for m in MODELS])
    fig,ax=plt.subplots(figsize=(7.1,4.5),constrained_layout=True); im=ax.imshow(mat,cmap='cividis',vmin=75,vmax=100,aspect='auto')
    ax.set_xticks(np.arange(11),[f'{a}\nn={d["support_test"][c]}' for a,c in zip(ABBR,CORE)],rotation=45,ha='right'); ax.set_yticks(np.arange(6),[NAMES[m] for m in MODELS]);
    for i in range(6):
        for j in range(11): ax.text(j,i,f'{mat[i,j]:.1f}',ha='center',va='center',fontsize=6.8,color='white' if mat[i,j]<88 else 'black')
    fig.colorbar(im,ax=ax,label='F1 (%)'); ax.set_title('Mean F1 theo lớp qua ba seed',fontweight='bold'); heat=save(fig,'fig_11_per_class_heatmap.png')

    fig,ax=plt.subplots(figsize=(6.35,4.3),constrained_layout=True)
    for m in MODELS:
        x=PARAMS_11[m]/1e6; yy=agg[m]['macro_f1']['mean']*100; ax.scatter(x,yy,s=95,color=COLORS[m]); ax.annotate(SHORT[m],(x,yy),xytext=(5,4),textcoords='offset points',fontsize=8,fontweight='bold')
    ax.set_xlabel('Tham số dự kiến khi thay classifier bằng 11 đầu ra (triệu)'); ax.set_ylabel('Macro-F1 chuyển tiếp (%)'); ax.grid(alpha=.2); ax.set_title('Chất lượng chuyển tiếp và quy mô mô hình',fontweight='bold'); pareto=save(fig,'fig_11_quality_params.png')
    return {'overall':overall,'stability':stability,'support':support,'heat':heat,'pareto':pareto}

def existing_figures():
    out={
        'overall':ASSETS/'fig_11_overall_metrics.png',
        'stability':ASSETS/'fig_11_seed_stability.png',
        'support':ASSETS/'fig_11_support.png',
        'heat':ASSETS/'fig_11_per_class_heatmap.png',
        'pareto':ASSETS/'fig_11_quality_params.png',
    }
    missing=[str(p) for p in out.values() if not p.exists()]
    if missing: raise FileNotFoundError('Thiếu assets: '+', '.join(missing))
    return out

def pct(x,n=3): return f'{x*100:.{n}f}%'.replace('.',',')
def add_cover(doc,d):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(52); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run('BÁO CÁO SƠ BỘ VÀ THIẾT KẾ NGHIÊN CỨU'); base.set_run_font(r,size=11,color=base.BLUE,bold=True)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('So sánh sáu kiến trúc CNN trong nhận diện 11 nhóm bệnh, sinh vật gây hại và trạng thái cây lúa'); base.set_run_font(r,size=23,color=base.INK,bold=True)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('Phân tích chuyển tiếp từ campaign 19 lớp • Campaign 11 lớp chính thức chưa huấn luyện'); base.set_run_font(r,size=12.5,color=base.DARK_BLUE,italic=True)
    rows=[['Hạng mục','Thông tin'],['Quy tắc chọn lớp','Giữ 11 lớp có support test ≥ 100'],['Số ảnh phân tích','9.376 ảnh test; loại 453 ảnh thuộc 8 lớp support thấp'],['Nguồn checkpoint','18 checkpoint của campaign 19 lớp (6 mô hình × 3 seed)'],['Trạng thái bằng chứng','Sơ bộ; không được gọi là mô hình 11 đầu ra']]
    base.add_table_custom(doc,rows,[2200,7160],font_size=9.1,caption_title='cover_11class',caption_description='Phạm vi và trạng thái báo cáo 11 lớp')
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(24); r=p.add_run('RiceAI Research • Phiên bản chuyển tiếp 1.0'); base.set_run_font(r,size=10,color=base.MUTED,bold=True); doc.add_page_break()

def build_doc(d,figs):
    doc=Document(); base.configure_section(doc.sections[0]); base.configure_styles(doc); base.add_running_furniture(doc.sections[0]); bullet=base.add_bullet_numbering_definition(doc)
    base.ensure_hyperlink_style(doc)
    add_cover(doc,d)
    base.add_heading(doc,'Tóm tắt điều hành',1)
    base.add_callout(doc,'Trạng thái bắt buộc: số liệu dưới đây được tái tính từ dự đoán của classifier 19 lớp trên 9.376 ảnh thuộc 11 lớp đủ support. Đây chưa phải kết quả của campaign 11 lớp được huấn luyện lại.')
    a=d['aggregate']; base.add_paragraph(doc,f"EfficientNet-B0 đứng đầu với Macro-F1 {pct(a['efficientnetb0']['macro_f1']['mean'])} ± {pct(a['efficientnetb0']['macro_f1']['sample_std'])}; MobileNetV3-Large đứng thứ hai với {pct(a['mobilenetv3large']['macro_f1']['mean'])} ± {pct(a['mobilenetv3large']['macro_f1']['sample_std'])}. Thứ hạng sáu mô hình không thay đổi so với phân tích 19 lớp.")
    for t in ['11 lớp được chọn bằng ngưỡng support test ≥ 100, không chọn theo kết quả F1.','Dự đoán rơi vào một trong tám lớp bị loại vẫn được tính là sai; không tái chuẩn hóa xác suất để làm đẹp kết quả.','Kết quả này dùng để thiết kế campaign 11 lớp; bài báo chính thức phải thay classifier thành 11 đầu ra và train lại 18 run.']: base.add_bullet(doc,t,bullet)
    base.add_figure(doc,figs['overall'],'Hình 1. Accuracy và Macro-F1 tái tính trên 11 lớp đủ support.',alt='Hai biểu đồ thanh so sánh sáu mô hình trên phân tích chuyển tiếp 11 lớp.')

    base.add_heading(doc,'1. Tên đề tài và phạm vi 11 lớp',1)
    base.add_paragraph(doc,'Tên đề xuất: “So sánh sáu kiến trúc CNN trong nhận diện 11 nhóm bệnh, sinh vật gây hại và trạng thái cây lúa”. Không dùng cụm “11 bệnh” vì danh sách gồm sâu hại, ốc, cỏ dại và lớp bình thường.')
    rows=[['#','Canonical label','Tên tiếng Việt','Train','Validation','Test','Tổng']]
    for i,c in enumerate(CORE,1):
        x=d['split_counts'][c]; rows.append([i,c,VI[c],x['train'],x['validation'],x['test'],sum(x.values())])
    base.add_table_custom(doc,rows,[400,2100,1900,1050,1100,900,950],font_size=7.8,caption_title='table_11classes',caption_description='Mười một lớp giữ lại và support theo split')
    base.add_caption(doc,'Bảng 1. Danh sách 11 lớp theo tiêu chí support test ≥ 100.')
    base.add_figure(doc,figs['support'],'Hình 2. Phân bố support test vẫn mất cân bằng trong bộ 11 lớp.',alt='Biểu đồ thanh ngang số ảnh test của từng lớp.')

    base.add_heading(doc,'2. Tám lớp tạm loại và lý do',1)
    excluded=['bacterial_leaf_streak','bacterial_panicle_blight','beetle','downy_mildew','leaf_smut','sheath_blight','southern_black_streaked_dwarf','whitebacked_planthopper']
    oldvi=base.CLASS_VI; oldsupport={}
    m=json.loads((FINAL/'efficientnetb0'/'seed_42'/'test_metrics.json').read_text(encoding='utf-8'))
    for e in m['per_class']: oldsupport[e['class_name']]=e['support']
    rows=[['Canonical label','Tên diễn giải','Test n','Lý do tạm loại']]+[[c,oldvi[c],oldsupport[c],'Support test < 100'] for c in excluded]
    base.add_table_custom(doc,rows,[2700,2300,1100,3260],font_size=8.2,caption_title='table_excluded',caption_description='Tám lớp tạm loại theo cùng một quy tắc định lượng')
    base.add_callout(doc,'Không xóa dữ liệu gốc và không sửa campaign 19 lớp. Tám lớp này được giữ trong phụ lục/pilot và có thể quay lại khi thu thập đủ dữ liệu.',warning=True)

    base.add_heading(doc,'3. Cách tái tính số liệu hiện tại',1)
    for t in ['Giữ các hàng test có nhãn thật thuộc 11 lớp được chọn.','Giữ nguyên predicted_label từ mô hình 19 đầu ra. Nếu dự đoán sang lớp bị loại, quyết định đó vẫn là sai.','Tính lại Accuracy, per-class Precision/Recall/F1 và Macro-P/R/F1 cho 11 nhãn.','Lấy mean và sample SD qua ba seed; paired bootstrap 2.000 mẫu theo từng seed cho V3-Large so với từng mô hình.']: base.add_bullet(doc,t,bullet)
    base.add_paragraph(doc,'Cách làm này đánh giá độ nhạy của kết luận xếp hạng trước việc bỏ các lớp quá ít mẫu. Nó không mô phỏng chính xác classifier 11 đầu ra và không thay thế huấn luyện lại.')

    base.add_heading(doc,'4. Khung tiêu chí dùng để so sánh mô hình',1)
    base.add_paragraph(doc,'Không xếp hạng mô hình bằng một con số duy nhất rồi bỏ qua các mặt còn lại. Nghiên cứu dùng Macro-F1 trung bình qua ba seed làm tiêu chí chất lượng chính, sau đó kiểm tra chéo bằng Accuracy, Macro-Precision, Macro-Recall, độ lệch chuẩn qua seed, F1 từng lớp, paired bootstrap và quy mô mô hình. Không tạo điểm tổng hợp có trọng số vì việc tự chọn trọng số giữa chất lượng, độ ổn định và tài nguyên sẽ làm kết luận khó kiểm chứng.')
    comparison_rows=[
        ['Khía cạnh','Thước đo','Cách tính trong nghiên cứu','Ý nghĩa khi so sánh'],
        ['Chất lượng chính','Macro-F1 mean','Tính F1 riêng cho từng lớp, lấy trung bình đều 11 lớp trong từng seed, rồi lấy mean của ba seed.','Tiêu chí xếp hạng chính vì mỗi lớp có trọng số ngang nhau dù support rất khác nhau.'],
        ['Đúng tổng thể','Accuracy mean','Tỷ lệ dự đoán đúng trên 9.376 ảnh trong từng seed, sau đó lấy mean ba seed.','Cho biết độ đúng chung nhưng có thể bị các lớp lớn như sâu đục thân và rầy nâu chi phối.'],
        ['Lỗi dương tính','Macro-Precision','Tính precision từng lớp rồi lấy trung bình đều 11 lớp.','Thấp khi mô hình thường gọi nhầm ảnh của lớp khác vào lớp đang xét.'],
        ['Khả năng phát hiện','Macro-Recall','Tính recall từng lớp rồi lấy trung bình đều 11 lớp.','Thấp khi mô hình bỏ sót nhiều ảnh thật của một hoặc nhiều lớp.'],
        ['Độ ổn định','Sample SD qua seed','Độ lệch chuẩn mẫu của metric ở seed 42, 123 và 2026, dùng mẫu số n−1.','SD nhỏ cho thấy kết quả ít dao động hơn khi huấn luyện lại; SD không phải khoảng tin cậy 95%.'],
        ['Theo từng lớp','F1 lớp và support','Báo cáo F1 trung bình ba seed kèm số ảnh test của từng lớp.','Xác định mô hình yếu ở lớp nào và tránh che khuất lớp khó bằng điểm tổng thể cao.'],
        ['Độ bất định cặp','Paired bootstrap 95% CI','Lấy mẫu lại cùng chỉ số ảnh cho hai mô hình, tính chênh lệch Macro-F1 2.000 lần theo từng seed.','Đánh giá liệu chênh lệch quan sát có ổn định trên tập ảnh; CI chứa 0 nghĩa là chưa tách biệt rõ ở phép so sánh đó.'],
        ['Khả năng triển khai','Số tham số dự kiến','Tính số tham số khi thay classifier 19 đầu ra bằng 11 đầu ra.','Chỉ là chỉ báo quy mô; chưa thay thế latency, RAM, năng lượng hoặc kích thước ứng dụng trên điện thoại.'],
    ]
    base.add_table_custom(doc,comparison_rows,[1500,1500,3100,3260],font_size=7.4,caption_title='table_comparison_framework',caption_description='Khung tiêu chí dùng để đánh giá sáu mô hình')
    base.add_paragraph(doc,'Quy tắc đọc kết quả: thứ hạng mô tả được xác định bằng Macro-F1 mean. Khi hai mô hình có điểm gần nhau, phải đọc thêm SD và paired bootstrap trước khi dùng từ “tốt hơn”. Accuracy, Macro-Precision và Macro-Recall được dùng để giải thích bản chất sai số, còn F1 từng lớp cho biết chất lượng có đồng đều hay chỉ cao ở các lớp nhiều ảnh. Quyết định triển khai mobile chỉ được đưa ra sau khi có thêm latency, RAM và năng lượng đo trên cùng thiết bị.')
    base.add_heading(doc,'4.1 Cách tính, giá trị quan sát và giới hạn của từng thước đo',2)
    base.add_paragraph(doc,'Các thước đo không có một “độ chính xác của thước đo” riêng. Với một tệp nhãn thật và nhãn dự đoán cố định, code tính chúng theo công thức xác định. Độ tin cậy của kết luận phụ thuộc vào chất lượng nhãn, độ đại diện của 9.376 ảnh test, dao động giữa các lần huấn luyện và phương pháp ước lượng bất định. Vì vậy bảng dưới đây báo cáo giá trị quan sát, đồng thời nêu phần nào đã được kiểm tra và phần nào chưa thể suy rộng.')
    metric_detail_rows=[
        ['Thước đo','Cách tính','Giá trị quan sát trong phân tích 11 lớp','Độ tin cậy và giới hạn'],
        ['Accuracy','Số ảnh dự đoán đúng chia cho N = 9.376. Tương đương tổng TP của 11 lớp chia N.','Mean sáu mô hình: 96,551% đến 99,335%. EfficientNet-B0: 99,335% ± 0,048 điểm %.','Tính chính xác từ toàn bộ tệp dự đoán. Tuy nhiên, các lớp lớn chi phối nên Accuracy cao không bảo đảm lớp hiếm được nhận diện tốt.'],
        ['Precision lớp c và Macro-Precision','Precision_c = TP_c / (TP_c + FP_c). Macro-Precision là trung bình không trọng số của 11 precision lớp.','Mean sáu mô hình: 93,074% đến 98,288%. EfficientNet-B0: 98,288% ± 0,195 điểm %.','Phản ánh mức dự đoán nhầm vào lớp. Khi mẫu số bằng 0, pipeline đặt giá trị lớp đó bằng 0. Macro average cho mỗi lớp trọng số bằng nhau.'],
        ['Recall lớp c và Macro-Recall','Recall_c = TP_c / (TP_c + FN_c). Macro-Recall là trung bình không trọng số của 11 recall lớp.','Mean sáu mô hình: 90,312% đến 97,717%. EfficientNet-B0: 97,717% ± 0,146 điểm %.','Phản ánh tỷ lệ ảnh thật được phát hiện. Recall cao có thể đi cùng precision thấp, vì vậy phải đọc hai chỉ số cùng nhau.'],
        ['F1 lớp c','F1_c = 2TP_c / (2TP_c + FP_c + FN_c), tương đương trung bình điều hòa của precision và recall.','Mean theo từng lớp và mô hình trải từ 76,711% đến 100,000%. Bảng lớp luôn ghi kèm support test.','Cân bằng precision và recall của một lớp. F1 không dùng true negative, vì vậy không mô tả toàn bộ ma trận nhầm lẫn.'],
        ['Macro-F1','Tính F1 riêng cho 11 lớp rồi lấy trung bình không trọng số. Báo cáo mean của ba seed.','Mean sáu mô hình: 91,524% đến 97,994%. EfficientNet-B0: 97,994% ± 0,159 điểm %; V3-Large: 97,585% ± 0,070 điểm %.','Tiêu chí xếp hạng chính vì dữ liệu mất cân bằng. Không được tính bằng cách lấy trung bình điều hòa của Macro-P và Macro-R.'],
        ['Sample SD qua seed','s = căn bậc hai của tổng (x_i - x̄)^2 chia (n - 1), với n = 3 seed.','SD Macro-F1 từ 0,070 đến 1,771 điểm %. Nhỏ nhất: V3-Large; lớn nhất: ShuffleNet V2.','Chỉ mô tả dao động của ba lần train. Ba seed quá ít để xem mean ± SD là khoảng tin cậy 95% hoặc khẳng định độ ổn định tổng quát.'],
        ['Paired bootstrap 95% CI','Ở mỗi lần lặp, lấy mẫu lại có hoàn lại cùng chỉ số ảnh cho hai mô hình, tính Δ Macro-F1. Dùng B = 2.000 và percentile 2,5% đến 97,5%.','15 so sánh V3-Large với năm mô hình theo ba seed: 6 CI chứa 0 và 9 CI không chứa 0.','Giữ cặp trên cùng ảnh nên phù hợp cho so sánh hai dự đoán cố định. CI chỉ phản ánh biến thiên do tập ảnh, chưa gộp biến thiên train và chưa hiệu chỉnh đa so sánh.'],
        ['Số tham số dự kiến','Lấy số tham số checkpoint 19 lớp và trừ phần classifier thừa khi đổi từ 19 xuống 11 đầu ra.','Khoảng 1,265 triệu đến 11,182 triệu tham số tùy backbone.','Đếm cấu trúc là xác định, nhưng đây là ước tính đầu 11 lớp chưa có checkpoint. Nó không dự đoán trực tiếp latency, RAM hay năng lượng.'],
    ]
    base.add_table_custom(doc,metric_detail_rows,[1450,2500,2550,2860],font_size=7.05,caption_title='table_metric_details',caption_description='Công thức, kết quả và giới hạn của các thước đo')
    base.add_heading(doc,'4.2 Nguồn phương pháp và mức độ tin cậy',2)
    source_p=doc.add_paragraph(); source_p.add_run('Accuracy: ')
    base.add_external_hyperlink(source_p,'scikit-learn accuracy_score','https://scikit-learn.org/stable/modules/generated/sklearn.metrics.accuracy_score.html',size_pt=9.0)
    source_p.add_run('. Đây là tài liệu API chính thức của thư viện dùng phổ biến để kiểm tra định nghĩa và miền giá trị của chỉ số.')
    source_p=doc.add_paragraph(); source_p.add_run('Precision, Recall và F1: ')
    base.add_external_hyperlink(source_p,'precision_score','https://scikit-learn.org/stable/modules/generated/sklearn.metrics.precision_score.html',size_pt=9.0)
    source_p.add_run(', ')
    base.add_external_hyperlink(source_p,'recall_score','https://scikit-learn.org/stable/modules/generated/sklearn.metrics.recall_score.html',size_pt=9.0)
    source_p.add_run(' và ')
    base.add_external_hyperlink(source_p,'f1_score','https://scikit-learn.org/stable/modules/generated/sklearn.metrics.f1_score.html',size_pt=9.0)
    source_p.add_run('. Các trang chính thức nêu công thức TP, FP, FN và quy tắc macro average không trọng số.')
    source_p=doc.add_paragraph(); source_p.add_run('Sample standard deviation: ')
    base.add_external_hyperlink(source_p,'NIST SEMATECH e-Handbook Measures of Scale','https://www.itl.nist.gov/div898/handbook/eda/section3/eda356.htm',size_pt=9.0)
    source_p.add_run('. NIST là cơ quan đo lường tiêu chuẩn của Hoa Kỳ; trang này nêu công thức mẫu số n - 1 và cách diễn giải độ phân tán.')
    source_p=doc.add_paragraph(); source_p.add_run('Bootstrap: ')
    base.add_external_hyperlink(source_p,'Efron 1979 Bootstrap Methods Another Look at the Jackknife','https://doi.org/10.1214/aos/1176344552',size_pt=9.0)
    source_p.add_run(' và ')
    base.add_external_hyperlink(source_p,'scikit-learn resample','https://scikit-learn.org/stable/modules/generated/sklearn.utils.resample.html',size_pt=9.0)
    source_p.add_run('. Bài của Efron là công trình gốc đã bình duyệt trên The Annals of Statistics; tài liệu scikit-learn xác nhận cơ chế lấy mẫu lại có hoàn lại. Báo cáo dùng percentile CI 95% như mô tả phương pháp, không gọi CI là xác suất mô hình đúng 95%.')

    base.add_heading(doc,'5. Kết quả tổng thể qua ba seed',1)
    rows=[['Hạng','Mô hình','Accuracy mean ± SD','Macro-P','Macro-R','Macro-F1 mean ± SD']]
    for i,m in enumerate(MODELS,1):
        rows.append([i,NAMES[m],f"{pct(a[m]['accuracy']['mean'])} ± {pct(a[m]['accuracy']['sample_std'])}",pct(a[m]['macro_precision']['mean']),pct(a[m]['macro_recall']['mean']),f"{pct(a[m]['macro_f1']['mean'])} ± {pct(a[m]['macro_f1']['sample_std'])}"])
    base.add_table_custom(doc,rows,[600,1900,1900,1450,1450,2060],font_size=7.8,caption_title='table_aggregate_11',caption_description='Metric trung bình và sample SD qua ba seed')
    base.add_paragraph(doc,'EfficientNet-B0 vẫn dẫn đầu; MobileNetV3-Large đứng thứ hai. Vì thứ hạng không đổi, kết luận so sánh backbone của pilot không phụ thuộc hoàn toàn vào tám lớp support thấp.')

    base.add_heading(doc,'6. Kết quả từng seed và độ ổn định',1)
    rows=[['Mô hình','Seed 42','Seed 123','Seed 2026','Mean ± SD']]
    for m in MODELS: rows.append([NAMES[m]]+[pct(a[m]['macro_f1']['values_by_seed'][str(s)]) for s in SEEDS]+[f"{pct(a[m]['macro_f1']['mean'])} ± {pct(a[m]['macro_f1']['sample_std'])}"])
    base.add_table_custom(doc,rows,[2100,1450,1450,1450,2910],font_size=8.1,caption_title='table_seed_11',caption_description='Macro-F1 từng seed trong phân tích chuyển tiếp 11 lớp')
    base.add_figure(doc,figs['stability'],'Hình 3. Macro-F1 của sáu mô hình qua ba training seed.',alt='Biểu đồ đường độ ổn định Macro-F1 qua seed.')

    base.add_heading(doc,'7. Phân tích theo lớp',1)
    base.add_figure(doc,figs['heat'],'Hình 4. F1 trung bình của 11 lớp qua ba seed; tên lớp và support test được trình bày riêng để dễ đọc.',width=6.75,alt='Heatmap F1 sáu mô hình trên 11 lớp với tên tiếng Việt và số ảnh test.')
    rows=[['Lớp (test n)']+[SHORT[m] for m in MODELS]]
    for c in CORE: rows.append([f"{VI[c]} ({d['support_test'][c]})"]+[f"{d['per_class'][m][c]['mean']*100:.2f}" for m in MODELS])
    base.add_table_custom(doc,rows,[2200,1190,1190,1190,1190,1190,1210],font_size=7.4,caption_title='table_perclass_11',caption_description='Mean F1 phần trăm theo lớp qua ba seed')

    base.add_heading(doc,'8. Paired bootstrap và cách diễn giải',1)
    rows=[['V3-Large trừ mô hình','Seed','Δ Macro-F1','CI percentile 95%']]
    for other in [m for m in MODELS if m!='mobilenetv3large']:
        key=f'mobilenetv3large_minus_{other}'
        for s in SEEDS:
            x=d['paired_bootstrap_vs_mobilenetv3large'][key][str(s)]
            rows.append([NAMES[other],s,f"{x['observed_delta']*100:+.3f}",f"[{x['ci95_lower']*100:+.3f}; {x['ci95_upper']*100:+.3f}]"])
    base.add_table_custom(doc,rows,[2700,900,1900,3860],font_size=7.8,caption_title='table_bootstrap_11',caption_description='Paired bootstrap 2.000 mẫu trên 9.376 ảnh theo từng seed')
    base.add_paragraph(doc,'CI chứa 0 không chứng minh hai mô hình bằng nhau; CI không chứa 0 chỉ là bằng chứng trên tập ảnh và seed tương ứng. Không hiệu chỉnh đa so sánh trong bảng này.')

    base.add_heading(doc,'9. Quy mô mô hình và triển khai',1)
    base.add_figure(doc,figs['pareto'],'Hình 5. Macro-F1 chuyển tiếp và số tham số dự kiến với classifier 11 đầu ra.',alt='Biểu đồ phân tán chất lượng và tham số dự kiến.')
    rows=[['Mô hình','Tham số checkpoint 19 lớp','Tham số dự kiến 11 đầu ra','Ghi chú']]
    for m in MODELS: rows.append([NAMES[m],f"{base.PARAMS[m]/1e6:.3f} M",f"{PARAMS_11[m]/1e6:.3f} M",'Chưa có checkpoint 11 lớp'])
    base.add_table_custom(doc,rows,[2200,2200,2200,2760],font_size=8.1,caption_title='table_params_11',caption_description='Tham số hiện có và tham số dự kiến sau khi thay classifier')
    base.add_callout(doc,'Không sử dụng thời gian train, file checkpoint hoặc latency của campaign 19 lớp như số đo chính thức cho campaign 11 lớp. Các đại lượng đó phải đo lại sau huấn luyện.',warning=True)

    base.add_heading(doc,'10. Protocol campaign 11 lớp sẽ chạy sau',1)
    steps=['Tạo manifest V3 chỉ gồm 11 nhãn và khóa SHA-256.','Đổi classifier thành 11 đầu ra; tính lại effective-number class weights trên train.','Giữ sáu backbone, ba seed 42/123/2026 và cùng recipe để bảo đảm so sánh công bằng.','Không dùng test để chọn epoch hoặc siêu tham số; chỉ mở test sau khi đủ 18 run.','Sinh lại toàn bộ bảng, hình, bootstrap, thời gian train và benchmark mobile từ campaign mới.']
    for i,t in enumerate(steps,1): base.add_bullet(doc,f'{i}. {t}',bullet)

    base.add_heading(doc,'11. Giới hạn và kết luận',1)
    base.add_paragraph(doc,'Phân tích 11 lớp hiện tại có giá trị như kiểm tra độ nhạy và hỗ trợ thiết kế nghiên cứu. Nó chưa chứng minh hiệu năng của một mô hình được huấn luyện trực tiếp cho 11 lớp. Việc chọn lớp dựa trên support sau khi campaign 19 lớp đã hoàn tất phải được công khai là quyết định thiết kế chuyển tiếp, không phải protocol định trước của pilot.')
    base.add_callout(doc,f"Kết luận sơ bộ: thứ hạng giữ nguyên trên 11 lớp đủ support; EfficientNet-B0 đạt {pct(a['efficientnetb0']['macro_f1']['mean'])} Macro-F1 và MobileNetV3-Large đạt {pct(a['mobilenetv3large']['macro_f1']['mean'])}. Kết luận cuối chỉ được cập nhật sau campaign 11 lớp chính thức.")

    base.add_heading(doc,'Phụ lục A. Kết quả 18 tổ hợp model × seed',1)
    rows=[['Mô hình','Seed','Accuracy','Macro-P','Macro-R','Macro-F1']]
    for m in MODELS:
        for s in SEEDS:
            x=d['runs'][m][str(s)]; rows.append([NAMES[m],s,pct(x['accuracy']),pct(x['macro_precision']),pct(x['macro_recall']),pct(x['macro_f1'])])
    base.add_table_custom(doc,rows,[2100,700,1600,1600,1600,1760],font_size=7.7,caption_title='appendix_runs_11',caption_description='Metric 18 tổ hợp trên phân tích chuyển tiếp 11 lớp')
    base.add_heading(doc,'Phụ lục B. Provenance',1)
    for t in [f'Dữ liệu số: {DATA_JSON}',f'Dự đoán nguồn: {FINAL}/<model>/seed_<seed>/test_predictions.csv','Quy tắc lọc: true_label thuộc 11 lớp support test ≥ 100; predicted_label ngoài 11 lớp vẫn tính sai.','Không sửa bất kỳ checkpoint, history.json, run_manifest.json hoặc file dự đoán gốc nào.']: base.add_bullet(doc,t,bullet)
    base.normalize_numbering_order(doc); OUT.parent.mkdir(parents=True,exist_ok=True); doc.save(OUT)
    print('SAVED',OUT)

if __name__=='__main__':
    import argparse
    ap=argparse.ArgumentParser(); ap.add_argument('--data-only',action='store_true'); ap.add_argument('--doc-only',action='store_true'); args=ap.parse_args()
    data=json.loads(DATA_JSON.read_text(encoding='utf-8')) if args.doc_only else build_data()
    if not args.data_only: build_doc(data,existing_figures())
    print('DATA',DATA_JSON); print('ASSETS',ASSETS)
