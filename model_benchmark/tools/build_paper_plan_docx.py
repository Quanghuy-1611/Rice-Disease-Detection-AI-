from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "KE_HOACH_BAI_BAO_SO_SANH_6_MO_HINH.md"
OUTPUT = ROOT / "docs" / "KE_HOACH_BAI_BAO_SO_SANH_6_MO_HINH.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
MUTED = "5B6573"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
GRID = "B7C5D4"
WHITE = "FFFFFF"
RED = "9B1C1C"

CONTENT_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_TOP_BOTTOM = 80
CELL_START_END = 120


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_margins(cell, top=CELL_TOP_BOTTOM, start=CELL_START_END, bottom=CELL_TOP_BOTTOM, end=CELL_START_END):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = f"w:{edge}"
        node = tc_mar.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_fill(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_border(cell, color=GRID, size="4"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def apply_table_geometry(table, widths, indent_dxa=TABLE_INDENT_DXA):
    if sum(widths) != CONTENT_DXA:
        widths[-1] += CONTENT_DXA - sum(widths)
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(widths[index] / 1440)


def choose_widths(rows):
    cols = len(rows[0])
    if cols == 2:
        return [2700, 6660]
    if cols == 3:
        return [2100, 3300, 3960]
    if cols == 4:
        return [1550, 2550, 3300, 1960]
    if cols == 5:
        return [1300, 2300, 1700, 2700, 1360]
    lengths = []
    for col in range(cols):
        longest = max(len(row[col]) for row in rows)
        lengths.append(max(8, min(longest, 45)) ** 0.5)
    total = sum(lengths)
    widths = [max(780, round(CONTENT_DXA * value / total)) for value in lengths]
    scale = CONTENT_DXA / sum(widths)
    widths = [round(width * scale) for width in widths]
    widths[-1] += CONTENT_DXA - sum(widths)
    return widths


def add_hyperlink(paragraph, text, url, color=BLUE, underline=True):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color_el = OxmlElement("w:color")
    color_el.set(qn("w:val"), color)
    r_pr.append(color_el)
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        r_pr.append(u)
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    fonts.set(qn("w:eastAsia"), "Calibri")
    r_pr.append(fonts)
    new_run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    new_run.append(text_el)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


INLINE_RE = re.compile(r"(\*\*.+?\*\*|(?<!\*)\*[^*]+?\*(?!\*)|`.+?`|\[[^\]]+\]\(https?://[^)]+\)|https?://[^\s]+)")


def add_inline(paragraph, text, size=None, color=None):
    cursor = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            set_run_font(run, size=size, color=color)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, color=color, bold=True)
        elif token.startswith("*"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=size, color=color, italic=True)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name="Consolas", size=(size or 10) - 0.5, color=DARK_BLUE)
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), LIGHT_GRAY)
            run._element.get_or_add_rPr().append(shd)
        elif token.startswith("["):
            label, url = re.match(r"\[([^\]]+)\]\((https?://[^)]+)\)", token).groups()
            add_hyperlink(paragraph, label, url)
        else:
            url = token.rstrip(".,;)")
            suffix = token[len(url):]
            add_hyperlink(paragraph, url, url)
            if suffix:
                run = paragraph.add_run(suffix)
                set_run_font(run, size=size, color=color)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run, size=size, color=color)


def next_num_id(numbering):
    ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    return max(ids, default=0) + 1


def add_numbering_definition(doc, marker, is_decimal=False):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = next_num_id(numbering)

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal" if is_decimal else "bullet")
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "%1." if is_decimal else marker)
    lvl.append(lvl_text)
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    lvl.append(suff)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "271")
    p_pr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    lvl.append(p_pr)
    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    r_pr.append(fonts)
    lvl.append(r_pr)
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_num(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num)
    p_pr.append(num_pr)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.25


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    specs = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in specs.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def configure_section(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.extend((fld_char_1, instr, fld_char_2))
    set_run_font(run, size=9, color=MUTED)


def add_running_furniture(section):
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("RiceAI  |  Benchmark 6 CNN cho nhận diện bệnh lúa")
    set_run_font(run, size=9, color=MUTED, bold=True)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(0)
    run = p.add_run("Protocol v2  |  Trang ")
    set_run_font(run, size=9, color=MUTED)
    add_page_field(p)


def add_cover(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(78)
    p.paragraph_format.space_after = Pt(18)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("KẾ HOẠCH NGHIÊN CỨU VÀ VIẾT BÀI BÁO")
    set_run_font(run, size=11, color=BLUE, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("So sánh sáu mạng nơ-ron tích chập cho nhận diện bệnh và đối tượng gây hại trên lúa")
    set_run_font(run, size=27, color=INK, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(34)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Khung thực nghiệm, thước đo, phân tích hiệu suất và cấu trúc bài báo")
    set_run_font(run, size=14, color=DARK_BLUE, italic=True)

    table = doc.add_table(rows=4, cols=2)
    rows = [
        ("Bài toán", "Phân loại ảnh RGB một nhãn, 19 lớp"),
        ("Mô hình", "MobileNetV2, MobileNetV3Large, ResNet-18, EfficientNet-B0, ShuffleNetV2 1.0×, DenseNet-121"),
        ("Protocol", "rice_disease_backbone_benchmark_v2"),
        ("Trạng thái", "Test chưa mở - chiến dịch đã khóa"),
    ]
    for r, values in enumerate(rows):
        for c, value in enumerate(values):
            cell = table.cell(r, c)
            cell.text = ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, top=110, bottom=110)
            set_cell_border(cell, color="D5DEE8", size="3")
            if c == 0:
                set_cell_fill(cell, LIGHT_BLUE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(value)
            set_run_font(run, size=9.5 if c else 10, color=INK, bold=(c == 0))
    apply_table_geometry(table, [1900, 7460])

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(42)
    p.paragraph_format.space_after = Pt(4)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("RiceAI Research")
    set_run_font(run, size=12, color=INK, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Phiên bản 1.0  •  25/08/2026")
    set_run_font(run, size=10, color=MUTED)
    doc.add_page_break()


def add_toc(doc):
    p = doc.add_paragraph("Mục lục", style="Heading 1")
    p.paragraph_format.space_before = Pt(0)
    sections = [
        "1. Định vị bài báo",
        "2. Sáu mô hình và lý do chọn",
        "3. Dữ liệu và giới hạn diễn giải",
        "4. Protocol huấn luyện cố định",
        "5. Thước đo chất lượng và cách diễn giải",
        "6. Thước đo hiệu quả",
        "7. Quy trình phân tích sau khi 18 run hoàn tất",
        "8. Bảng và hình bắt buộc",
        "9. Khung bài báo chi tiết",
        "10. Dự kiến kết quả và nguyên tắc không đoán trước kết quả",
        "11. Tiêu chí hoàn thành và lịch triển khai",
        "12. Bản đồ nguồn dẫn chứng",
        "13. Câu mẫu để viết đúng mức bằng chứng",
    ]
    for item in sections:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.12)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(item)
        set_run_font(run, size=10.5, color=DARK_BLUE)
    doc.add_page_break()


def add_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    widths = choose_widths(rows)
    for r, values in enumerate(rows):
        for c, value in enumerate(values):
            cell = table.cell(r, c)
            cell.text = ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_border(cell)
            if r == 0:
                set_cell_fill(cell, LIGHT_BLUE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            add_inline(p, value, size=8.3 if len(rows[0]) >= 4 else 9, color=INK)
            for run in p.runs:
                if r == 0:
                    run.bold = True
            if r == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_repeat_table_header(table.rows[0])
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        cant_split.set(qn("w:val"), "true")
        tr_pr.append(cant_split)
    apply_table_geometry(table, widths)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    return table


def parse_table(lines, start):
    rows = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        row = [cell.strip() for cell in lines[i].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in row):
            rows.append(row)
        i += 1
    return rows, i


def add_callout(doc, text, warning=False):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_cell_fill(cell, "FCE8E6" if warning else LIGHT_BLUE)
    set_cell_margins(cell, top=130, bottom=130, start=160, end=160)
    set_cell_border(cell, color="E4B4B0" if warning else "B7CDE2", size="4")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    add_inline(p, text, size=10.5, color=RED if warning else INK)
    apply_table_geometry(table, [CONTENT_DXA], indent_dxa=160)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def render_markdown(doc, text, bullet_id, number_id, check_id):
    lines = text.splitlines()
    i = 0
    paragraph_parts = []

    def flush_paragraph():
        nonlocal paragraph_parts
        if paragraph_parts:
            p = doc.add_paragraph()
            add_inline(p, " ".join(part.strip() for part in paragraph_parts))
            paragraph_parts = []

    while i < len(lines):
        raw = lines[i]
        line = raw.strip()
        if not line:
            flush_paragraph()
            i += 1
            continue
        if line.startswith("|") and i + 1 < len(lines) and lines[i + 1].strip().startswith("|"):
            flush_paragraph()
            rows, i = parse_table(lines, i)
            add_table(doc, rows)
            continue
        if line.startswith("### "):
            flush_paragraph()
            doc.add_paragraph(line[4:], style="Heading 3")
        elif line.startswith("## "):
            flush_paragraph()
            doc.add_paragraph(line[3:], style="Heading 2")
        elif line.startswith("# "):
            flush_paragraph()
            doc.add_paragraph(line[2:], style="Heading 1")
        elif line.startswith("> "):
            flush_paragraph()
            add_callout(doc, line[2:], warning=False)
        elif line.startswith("```powershell"):
            flush_paragraph()
            code = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            table = doc.add_table(rows=1, cols=1)
            cell = table.cell(0, 0)
            set_cell_fill(cell, "F3F5F7")
            set_cell_margins(cell, top=110, bottom=110, start=150, end=150)
            set_cell_border(cell, color="CBD5E1", size="3")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            for index, code_line in enumerate(code):
                run = p.add_run(code_line + ("\n" if index < len(code) - 1 else ""))
                set_run_font(run, name="Consolas", size=8.7, color=INK)
            apply_table_geometry(table, [CONTENT_DXA], indent_dxa=150)
        elif re.match(r"^- \[[ xX]\] ", line):
            flush_paragraph()
            text_value = re.sub(r"^- \[[ xX]\] ", "", line)
            p = doc.add_paragraph()
            apply_num(p, check_id)
            add_inline(p, text_value)
        elif line.startswith("- "):
            flush_paragraph()
            p = doc.add_paragraph()
            apply_num(p, bullet_id)
            add_inline(p, line[2:])
        elif re.match(r"^\d+\. ", line):
            flush_paragraph()
            p = doc.add_paragraph()
            apply_num(p, number_id)
            add_inline(p, re.sub(r"^\d+\. ", "", line))
        elif re.match(r"^(Phiên bản kế hoạch|Ngày lập|Trạng thái dữ liệu test|Protocol thực nghiệm chính):", line):
            pass
        else:
            paragraph_parts.append(line)
        i += 1
    flush_paragraph()


def build():
    text = SOURCE.read_text(encoding="utf-8")
    # Cover replaces the original markdown title and four metadata lines.
    body_start = text.find("## 1. Định vị bài báo")
    body = text[body_start:]

    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    configure_section(section)
    add_running_furniture(section)
    section.different_first_page_header_footer = True

    bullet_id = add_numbering_definition(doc, "•")
    number_id = add_numbering_definition(doc, "%1.", is_decimal=True)
    check_id = add_numbering_definition(doc, "☐")

    add_cover(doc)
    add_toc(doc)
    render_markdown(doc, body, bullet_id, number_id, check_id)

    props = doc.core_properties
    props.title = "Kế hoạch bài báo so sánh 6 mô hình nhận diện bệnh lúa"
    props.subject = "RiceAI backbone benchmark v2"
    props.author = "RiceAI Research"
    props.keywords = "rice disease, CNN, MobileNet, ResNet, EfficientNet, ShuffleNet, DenseNet, benchmark"
    props.comments = "Generated from the locked RiceAI benchmark paper plan."

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print("DOCX_CREATED")


if __name__ == "__main__":
    build()
