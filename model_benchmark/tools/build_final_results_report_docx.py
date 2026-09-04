from __future__ import annotations

import csv
import json
import math
import statistics
from collections import Counter, defaultdict
from pathlib import Path
from urllib.parse import urlparse

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor

from build_paper_plan_docx import (
    BLUE,
    DARK_BLUE,
    GRID,
    INK,
    LIGHT_BLUE,
    LIGHT_GRAY,
    MUTED,
    RED,
    WHITE,
    CONTENT_DXA,
    add_callout as base_add_callout,
    add_inline,
    add_numbering_definition,
    add_page_field,
    apply_num,
    apply_table_geometry,
    configure_section,
    configure_styles,
    set_cell_border,
    set_cell_fill,
    set_cell_margins,
    set_repeat_table_header,
    set_run_font,
)


ROOT = Path(__file__).resolve().parents[1]
FINAL_ROOT = ROOT / "runs" / "final_test_v2"
TRAIN_ROOT = ROOT / "runs" / "training_v2"
OUTPUT = ROOT / "docs" / "BAO_CAO_PHAN_TICH_KET_QUA_SO_SANH_6_MO_HINH.docx"
ASSET_DIR = ROOT / "docs" / "_results_report_assets"
DATASET_ROOT = ROOT.parent / "research" / "data" / "cleaned_v4"

MODEL_ORDER = [
    "efficientnetb0",
    "mobilenetv3large",
    "resnet18",
    "densenet121",
    "mobilenetv2",
    "shufflenetv2_1_0x",
]
MODEL_NAMES = {
    "mobilenetv2": "MobileNetV2",
    "mobilenetv3large": "MobileNetV3-Large",
    "resnet18": "ResNet-18",
    "efficientnetb0": "EfficientNet-B0",
    "shufflenetv2_1_0x": "ShuffleNet V2 1.0×",
    "densenet121": "DenseNet-121",
}
MODEL_SHORT = {
    "mobilenetv2": "MV2",
    "mobilenetv3large": "MV3-L",
    "resnet18": "RN18",
    "efficientnetb0": "Eff-B0",
    "shufflenetv2_1_0x": "Shuffle",
    "densenet121": "DN121",
}
MODEL_COLORS = {
    "efficientnetb0": "#2E7D32",
    "mobilenetv3large": "#2E74B5",
    "resnet18": "#6A5ACD",
    "densenet121": "#00838F",
    "mobilenetv2": "#F28E2B",
    "shufflenetv2_1_0x": "#B0413E",
}
PARAMS = {
    "mobilenetv2": 2_248_211,
    "mobilenetv3large": 4_226_371,
    "resnet18": 11_186_259,
    "efficientnetb0": 4_031_887,
    "shufflenetv2_1_0x": 1_273_079,
    "densenet121": 6_973_331,
}
WEIGHT_IDS = {
    "mobilenetv2": "MobileNet_V2_Weights.IMAGENET1K_V2",
    "mobilenetv3large": "MobileNet_V3_Large_Weights.IMAGENET1K_V2",
    "resnet18": "ResNet18_Weights.IMAGENET1K_V1",
    "efficientnetb0": "EfficientNet_B0_Weights.IMAGENET1K_V1",
    "shufflenetv2_1_0x": "ShuffleNet_V2_X1_0_Weights.IMAGENET1K_V1",
    "densenet121": "DenseNet121_Weights.IMAGENET1K_V1",
}
CLASS_NAMES = [
    "bacterial_leaf_blight",
    "bacterial_leaf_streak",
    "bacterial_panicle_blight",
    "beetle",
    "blast",
    "brown_planthopper",
    "brown_spot",
    "dead_heart",
    "downy_mildew",
    "golden_apple_snail",
    "leaf_roller",
    "leaf_smut",
    "normal",
    "sheath_blight",
    "southern_black_streaked_dwarf",
    "stem_borer",
    "tungro",
    "weeds",
    "whitebacked_planthopper",
]
CLASS_VI = {
    "bacterial_leaf_blight": "bạc lá vi khuẩn",
    "bacterial_leaf_streak": "sọc lá vi khuẩn",
    "bacterial_panicle_blight": "cháy bông vi khuẩn",
    "beetle": "bọ cánh cứng",
    "blast": "đạo ôn",
    "brown_planthopper": "rầy nâu",
    "brown_spot": "đốm nâu",
    "dead_heart": "nõn héo",
    "downy_mildew": "sương mai",
    "golden_apple_snail": "ốc bươu vàng",
    "leaf_roller": "sâu cuốn lá",
    "leaf_smut": "than lá",
    "normal": "bình thường",
    "sheath_blight": "khô vằn",
    "southern_black_streaked_dwarf": "lùn sọc đen phương Nam",
    "stem_borer": "sâu đục thân",
    "tungro": "tungro",
    "weeds": "cỏ dại",
    "whitebacked_planthopper": "rầy lưng trắng",
}
CLASS_ABBR = [
    "BLB", "BLS", "BPB", "Beetle", "Blast", "BPH", "BSpot", "DeadH", "DM",
    "GAS", "LR", "LSmut", "Normal", "SBlight", "SBSD", "StemB", "Tungro", "Weeds", "WBPH",
]
SEEDS = [42, 123, 2026]
FIGURE_WIDTH_IN = 6.35
FIGURE_DPI = 420


plt.rcParams.update({
    "font.family": "DejaVu Sans",
    "font.size": 9.2,
    "axes.titlesize": 10.6,
    "axes.labelsize": 9.2,
    "xtick.labelsize": 8.1,
    "ytick.labelsize": 8.1,
    "legend.fontsize": 8.0,
    "figure.facecolor": "white",
    "axes.facecolor": "white",
    "savefig.facecolor": "white",
})


def load_json(path: Path):
    return json.loads(path.read_text(encoding="utf-8"))


def pct(value: float, digits: int = 2) -> str:
    return f"{value * 100:.{digits}f}%"


def pp(value: float, digits: int = 3) -> str:
    return f"{value * 100:+.{digits}f} điểm %"


def fmt_int(value: int) -> str:
    return f"{value:,}".replace(",", ".")


def fmt_float_vi(value: float, digits: int = 3) -> str:
    return f"{value:.{digits}f}".replace(".", ",")


def set_row_no_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:cantSplit")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def set_repeat_table_header_safe(row):
    set_repeat_table_header(row)
    set_row_no_split(row)


def add_callout(doc, text, warning=False):
    """Create a one-cell callout and expose its first row as a header for accessibility."""
    base_add_callout(doc, text, warning=warning)
    table = doc.tables[-1]
    set_repeat_table_header_safe(table.rows[0])
    set_table_caption(
        table,
        "Cảnh báo phương pháp" if warning else "Ghi chú phương pháp",
        text[:240],
    )


def set_table_caption(table, title: str, description: str):
    tbl_pr = table._tbl.tblPr
    caption = OxmlElement("w:tblCaption")
    caption.set(qn("w:val"), title)
    tbl_pr.append(caption)
    descr = OxmlElement("w:tblDescription")
    descr.set(qn("w:val"), description)
    tbl_pr.append(descr)


def add_paragraph(doc, text: str = "", *, bold_prefix: str | None = None, italic: bool = False,
                  align=None, color=INK, size=10.5, keep_with_next=False):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.keep_with_next = keep_with_next
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_run_font(r, size=size, color=color, bold=True)
        r = p.add_run(text[len(bold_prefix):])
        set_run_font(r, size=size, color=color, italic=italic)
    else:
        add_inline(p, text, size=size, color=color)
        if italic:
            for run in p.runs:
                run.italic = True
    return p


def add_heading(doc, text: str, level: int = 1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    return p


def add_bullet(doc, text: str, num_id: int, size=10.3):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.2
    apply_num(p, num_id)
    add_inline(p, text, size=size, color=INK)
    return p


def add_number(doc, text: str, num_id: int, size=10.3, index: int = 1):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.2
    apply_num(p, num_id)
    add_inline(p, text, size=size, color=INK)
    return p


def add_bullet_numbering_definition(doc):
    """Create a Word-native bullet list with a Unicode bullet glyph."""
    num_id = add_numbering_definition(doc, "•")
    numbering = doc.part.numbering_part.element
    num = next(node for node in numbering.findall(qn("w:num")) if node.get(qn("w:numId")) == str(num_id))
    abstract_id = num.find(qn("w:abstractNumId")).get(qn("w:val"))
    abstract = next(
        node for node in numbering.findall(qn("w:abstractNum"))
        if node.get(qn("w:abstractNumId")) == abstract_id
    )
    # Keep abstractNumId and numId identical.  Desktop Word on this machine
    # otherwise resolves the freshly appended definition by the numeric id
    # rather than its abstractNumId reference and turns bullets into a
    # continuing decimal list during export.
    abstract.set(qn("w:abstractNumId"), str(num_id))
    num.find(qn("w:abstractNumId")).set(qn("w:val"), str(num_id))
    fonts = abstract.find(".//" + qn("w:rFonts"))
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Arial")
    fonts.set(qn("w:hint"), "default")
    return num_id


def normalize_numbering_order(doc):
    """Restore the OOXML schema order: abstractNum nodes must precede num nodes."""
    numbering = doc.part.numbering_part.element
    children = list(numbering)
    picture_nodes = [node for node in children if node.tag == qn("w:numPicBullet")]
    abstract_nodes = [node for node in children if node.tag == qn("w:abstractNum")]
    num_nodes = [node for node in children if node.tag == qn("w:num")]
    other_nodes = [
        node for node in children
        if node.tag not in {qn("w:numPicBullet"), qn("w:abstractNum"), qn("w:num")}
    ]
    for node in children:
        numbering.remove(node)
    for node in picture_nodes + abstract_nodes + num_nodes + other_nodes:
        numbering.append(node)


def add_caption(doc, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = False
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    set_run_font(r, size=9, color=MUTED, italic=True)
    return p


def add_table_custom(doc, rows, widths, *, font_size=8.4, header_fill=LIGHT_BLUE,
                     caption_title="", caption_description="", align_columns=None):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    for ri, values in enumerate(rows):
        row = table.rows[ri]
        set_row_no_split(row)
        for ci, value in enumerate(values):
            cell = table.cell(ri, ci)
            cell.text = ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, top=72, bottom=72, start=90, end=90)
            set_cell_border(cell, color=GRID, size="3")
            if ri == 0:
                set_cell_fill(cell, header_fill)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.04
            add_inline(p, str(value), size=font_size, color=INK)
            if ri == 0:
                for run in p.runs:
                    run.bold = True
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                # Keep a table header with at least the first data row instead
                # of leaving a lone header at the foot of a page.
                p.paragraph_format.keep_with_next = True
            elif align_columns and ci in align_columns:
                p.alignment = align_columns[ci]
    set_repeat_table_header_safe(table.rows[0])
    apply_table_geometry(table, widths, indent_dxa=90)
    if caption_title:
        set_table_caption(table, caption_title, caption_description or caption_title)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_metric_definition_table(doc):
    rows = [
        ["Thước đo", "Công thức", "Câu hỏi được trả lời"],
        ["Accuracy", "(Σᶜ₍c=1₎ TP_c) / N", "Bao nhiêu ảnh được phân loại đúng trên toàn bộ test?"],
        ["Precision lớp c", "TP_c / (TP_c + FP_c)", "Trong các ảnh dự đoán là lớp c, bao nhiêu ảnh đúng?"],
        ["Recall lớp c", "TP_c / (TP_c + FN_c)", "Trong các ảnh thật sự thuộc lớp c, mô hình tìm được bao nhiêu?"],
        ["F1 lớp c", "2 · Precision_c · Recall_c / (Precision_c + Recall_c)", "Mức cân bằng giữa precision và recall của lớp c?"],
        ["Macro-F1", "(1/C) · Σᶜ₍c=1₎ F1_c; C = 19", "Chất lượng trung bình khi mỗi lớp có trọng số bằng nhau?"],
        ["Sample SD qua seed", "√[Σⁿ₍i=1₎(x_i − x̄)² / (n − 1)]; n = 3", "Kết quả thay đổi bao nhiêu giữa ba lần huấn luyện?"],
    ]
    table = add_table_custom(
        doc,
        rows,
        [1500, 3500, 4360],
        font_size=8.15,
        caption_title="table_metric_definitions",
        caption_description="Công thức, mục đích và đơn vị diễn giải của các thước đo chính",
    )
    for row in table.rows[1:]:
        for run in row.cells[1].paragraphs[0].runs:
            run.font.name = "Cambria Math"
            run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Cambria Math")
            run.font.size = Pt(8.4)
    return table


def ensure_hyperlink_style(doc):
    try:
        style = doc.styles["Hyperlink"]
    except KeyError:
        style = doc.styles.add_style("Hyperlink", WD_STYLE_TYPE.CHARACTER)
    style.font.name = "Calibri"
    style.font.color.rgb = RGBColor(0x05, 0x63, 0xC1)
    style.font.underline = True


def add_external_hyperlink(paragraph, text: str, url: str, *, size_pt: float = 9.0):
    parsed = urlparse(url)
    if parsed.scheme != "https" or not parsed.netloc:
        raise ValueError(f"Nguồn ngoài không hợp lệ hoặc không dùng HTTPS: {url}")
    relationship_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    hyperlink.set(qn("w:history"), "1")
    hyperlink.set(qn("w:tooltip"), url)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    for tag, value in (("w:rStyle", "Hyperlink"), ("w:color", "0563C1"), ("w:u", "single")):
        element = OxmlElement(tag)
        element.set(qn("w:val"), value)
        run_properties.append(element)
    fonts = OxmlElement("w:rFonts")
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        fonts.set(qn(key), "Calibri")
    run_properties.append(fonts)
    for tag in ("w:sz", "w:szCs"):
        element = OxmlElement(tag)
        element.set(qn("w:val"), str(round(size_pt * 2)))
        run_properties.append(element)
    run.append(run_properties)
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return hyperlink


def set_image_alt(inline_shape, title: str, description: str):
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("name", title)
    doc_pr.set("title", title)
    doc_pr.set("descr", description)


def add_figure(doc, path: Path, caption: str, *, width=6.35, alt=""):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    shape = p.add_run().add_picture(str(path), width=Inches(width))
    set_image_alt(shape, caption.split(".")[0], alt or caption)
    add_caption(doc, caption)


def add_page_number_run(paragraph):
    add_page_field(paragraph)


def add_running_furniture(section):
    section.different_first_page_header_footer = True
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("RiceAI Research  |  Phân tích benchmark 6 CNN • Protocol V2")
    set_run_font(r, size=8.7, color=MUTED, bold=True)
    p_pr = p._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), BLUE)
    border.append(bottom)
    p_pr.append(border)
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("Báo cáo kết quả cuối • Trang ")
    set_run_font(r, size=8.7, color=MUTED)
    add_page_number_run(p)


def add_cover(doc, audit):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(52)
    p.paragraph_format.space_after = Pt(14)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("BÁO CÁO PHÂN TÍCH KẾT QUẢ THỰC NGHIỆM")
    set_run_font(r, size=11, color=BLUE, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("So sánh sáu kiến trúc CNN cho nhận diện 19 nhóm bệnh, dịch hại và trạng thái trên cây lúa")
    set_run_font(r, size=25, color=INK, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(26)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Kết quả khóa V2 • 18 lượt huấn luyện • 9.829 ảnh kiểm thử")
    set_run_font(r, size=13.5, color=DARK_BLUE, italic=True)

    rows = [
        ["Hạng mục", "Giá trị đã kiểm toán"],
        ["Dữ liệu", "65.516 ảnh RGB; train/validation/test = 45.858/9.829/9.829"],
        ["Thiết kế", "6 backbone × 3 seed (42, 123, 2026); cùng protocol PyTorch/torchvision"],
        ["Tiêu chí chính", "Macro-F1 trung bình ± sample SD qua 3 seed"],
        ["Trạng thái test", f"Mở theo guard một lần; audit hoàn tất {audit['completed_at_utc']}"],
    ]
    add_table_custom(doc, rows, [2150, 7210], font_size=9.2,
                     caption_title="cover_audit", caption_description="Thông tin định danh và audit của báo cáo")

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(26)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("RiceAI Research")
    set_run_font(r, size=12, color=INK, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Phiên bản 1.0 • 30/08/2026 • Asia/Ho_Chi_Minh")
    set_run_font(r, size=9.5, color=MUTED)
    doc.add_page_break()


def add_static_toc(doc):
    add_heading(doc, "Mục lục nội dung", 1)
    items = [
        "Tóm tắt điều hành",
        "1. Câu hỏi nghiên cứu và phạm vi bằng chứng",
        "2. Dữ liệu, làm sạch và chiến lược chia tập",
        "3. Sáu kiến trúc và giả thuyết kỹ thuật",
        "4. Protocol huấn luyện, chọn checkpoint và tái lập",
        "5. Quy trình mở test và thước đo đánh giá",
        "6. Kết quả tổng thể và độ ổn định qua seed",
        "7. So sánh bất định bằng paired bootstrap",
        "8. Phân tích sâu từng mô hình",
        "9. Phân tích theo lớp và mẫu lỗi",
        "10. Hiệu quả tài nguyên và hàm ý triển khai mobile",
        "11. Độ giá trị, giới hạn và nguy cơ sai lệch",
        "12. Kết luận và khuyến nghị nghiên cứu",
        "Tài liệu tham khảo",
        "Phụ lục A–D: số liệu chi tiết và provenance",
    ]
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(item)
        set_run_font(r, size=10.3, color=DARK_BLUE)
    add_callout(
        doc,
        "Cách đọc báo cáo: mọi phát biểu có số liệu được truy ngược tới final_comparison.json hoặc test_metrics.json. Những giải thích dựa trên kiến trúc được ghi rõ là giả thuyết hợp lý, không được trình bày như quan hệ nhân quả đã chứng minh.",
    )
    doc.add_page_break()


def load_campaign():
    comparison = load_json(FINAL_ROOT / "final_comparison.json")
    audit = load_json(FINAL_ROOT / "test_open_audit.json")
    metrics = defaultdict(dict)
    manifests = defaultdict(dict)
    for model in MODEL_NAMES:
        for seed in SEEDS:
            metrics[model][seed] = load_json(FINAL_ROOT / model / f"seed_{seed}" / "test_metrics.json")
            manifests[model][seed] = load_json(TRAIN_ROOT / model / f"seed_{seed}" / "run_manifest.json")
    return comparison, audit, metrics, manifests


def build_derived(comparison, metrics, manifests):
    aggregate = comparison["aggregate_across_training_seeds"]
    per_class = {}
    confusion_rate = {}
    support = {entry["class_name"]: entry["support"] for entry in metrics["efficientnetb0"][42]["per_class"]}
    for model in MODEL_NAMES:
        per_class[model] = {}
        matrices = []
        for cls in CLASS_NAMES:
            precision_values = []
            recall_values = []
            f1_values = []
            for seed in SEEDS:
                record = next(x for x in metrics[model][seed]["per_class"] if x["class_name"] == cls)
                precision_values.append(record["precision"])
                recall_values.append(record["recall"])
                f1_values.append(record["f1"])
            per_class[model][cls] = {
                "mean": statistics.mean(f1_values),
                "sd": statistics.stdev(f1_values),
                "values": f1_values,
                "precision_mean": statistics.mean(precision_values),
                "precision_sd": statistics.stdev(precision_values),
                "recall_mean": statistics.mean(recall_values),
                "recall_sd": statistics.stdev(recall_values),
            }
        for seed in SEEDS:
            cm = np.asarray(metrics[model][seed]["confusion_matrix"], dtype=float)
            row_sums = cm.sum(axis=1, keepdims=True)
            matrices.append(np.divide(cm, row_sums, out=np.zeros_like(cm), where=row_sums != 0))
        confusion_rate[model] = np.mean(np.stack(matrices), axis=0)

    resources = {}
    for model in MODEL_NAMES:
        durations = [manifests[model][s]["duration_seconds"] for s in SEEDS]
        sizes = [(TRAIN_ROOT / model / f"seed_{s}" / "best_model.pt").stat().st_size / (1024 ** 2) for s in SEEDS]
        best_val = [manifests[model][s]["best_validation_macro_f1"] for s in SEEDS]
        convergence = []
        error_counts_by_seed = {}
        for seed in SEEDS:
            history = load_json(TRAIN_ROOT / model / f"seed_{seed}" / "history.json")["epochs"]
            head_epochs = sum(1 for epoch in history if epoch["stage"] == "head")
            finetune_epochs = sum(1 for epoch in history if epoch["stage"] == "finetune")
            best_global_epoch = int(manifests[model][seed]["best_global_epoch"])
            best_epoch = next(epoch for epoch in history if int(epoch["global_epoch"]) == best_global_epoch)
            convergence.append({
                "seed": seed,
                "head_epochs": head_epochs,
                "finetune_epochs": finetune_epochs,
                "best_stage": best_epoch["stage"],
                "best_stage_epoch": int(best_epoch["stage_epoch"]),
                "best_is_last_finetune": best_epoch["stage"] == "finetune" and int(best_epoch["stage_epoch"]) == finetune_epochs,
            })
            cm = np.asarray(metrics[model][seed]["confusion_matrix"], dtype=int)
            error_counts_by_seed[seed] = int(cm.sum() - np.trace(cm))
        resources[model] = {
            "params": PARAMS[model],
            "checkpoint_mib": statistics.mean(sizes),
            "duration_hours_total": sum(durations) / 3600,
            "duration_hours_mean": statistics.mean(durations) / 3600,
            "duration_hours_sd": statistics.stdev(durations) / 3600,
            "best_val_mean": statistics.mean(best_val),
            "test_gap": aggregate[model]["macro_f1"]["mean"] - statistics.mean(best_val),
            "convergence": convergence,
            "error_counts_by_seed": error_counts_by_seed,
        }

    # Aggregate repeated decisions and images consistently misclassified by every run.
    repeated_errors = Counter()
    model_error_pairs = defaultdict(Counter)
    per_path = {}
    total_runs = len(MODEL_NAMES) * len(SEEDS)
    for model in MODEL_NAMES:
        for seed in SEEDS:
            path = FINAL_ROOT / model / f"seed_{seed}" / "test_predictions.csv"
            with path.open("r", encoding="utf-8", newline="") as handle:
                for row in csv.DictReader(handle):
                    rel = row["relative_path"]
                    truth = row["true_label"]
                    pred = row["predicted_label"]
                    record = per_path.setdefault(rel, {"truth": truth, "wrong": 0, "preds": Counter()})
                    record["preds"][pred] += 1
                    if truth != pred:
                        record["wrong"] += 1
                        repeated_errors[(truth, pred)] += 1
                        model_error_pairs[model][(truth, pred)] += 1
    unanimous_wrong = []
    for rel, record in per_path.items():
        if record["wrong"] == total_runs:
            unanimous_wrong.append((rel, record["truth"], record["preds"].most_common(3)))
    unanimous_wrong.sort()
    for model in MODEL_NAMES:
        resources[model]["error_pairs"] = model_error_pairs[model]
    return aggregate, per_class, confusion_rate, support, resources, repeated_errors, unanimous_wrong


def save_figure(fig, filename: str):
    path = ASSET_DIR / filename
    fig.savefig(
        path,
        dpi=FIGURE_DPI,
        bbox_inches="tight",
        pad_inches=0.06,
        facecolor="white",
        metadata={"Software": "RiceAI benchmark report"},
    )
    plt.close(fig)
    return path


def plot_overall(aggregate):
    models = MODEL_ORDER
    y = np.arange(len(models))
    fig, axes = plt.subplots(2, 1, figsize=(FIGURE_WIDTH_IN, 6.35), constrained_layout=True)
    for ax, metric, title, limits in [
        (axes[0], "macro_f1", "A. Macro-F1 — tiêu chí chính", (88.0, 98.6)),
        (axes[1], "accuracy", "B. Accuracy — chỉ số bổ sung", (95.0, 99.65)),
    ]:
        means = np.array([aggregate[m][metric]["mean"] * 100 for m in models])
        sds = np.array([aggregate[m][metric]["sample_std"] * 100 for m in models])
        colors = [MODEL_COLORS[m] for m in models]
        ax.barh(y, means, xerr=sds, color=colors, alpha=0.94, capsize=3.2, height=0.66,
                error_kw={"elinewidth": 1.0, "capthick": 1.0, "ecolor": "#222222"})
        ax.set_yticks(y, [MODEL_NAMES[m] for m in models])
        ax.invert_yaxis()
        ax.set_xlim(*limits)
        ax.set_title(title, fontweight="bold", pad=5)
        ax.set_xlabel("Phần trăm; thanh sai số = sample SD qua 3 seed")
        ax.grid(axis="x", alpha=0.20, linewidth=0.7)
        ax.set_axisbelow(True)
        span = limits[1] - limits[0]
        for yi, (val, sd) in enumerate(zip(means, sds)):
            # A fixed right-aligned label column keeps the numbers clear of
            # the error-bar caps, especially for models near the axis limit.
            x_text = limits[1] - span * 0.012
            ax.text(x_text, yi, f"{val:.3f}%", va="center", ha="right", fontsize=8.3,
                    fontweight="bold", color="#1F1F1F",
                    bbox={"facecolor": "white", "edgecolor": "none", "alpha": 0.88, "pad": 0.5})
    fig.suptitle("Kết quả tổng thể trên 9.829 ảnh test đã khóa", fontsize=11.4, fontweight="bold")
    return save_figure(fig, "fig_overall_metrics.png")


def plot_seed_stability(aggregate):
    fig, ax = plt.subplots(figsize=(FIGURE_WIDTH_IN, 4.75), constrained_layout=True)
    x = np.arange(len(SEEDS))
    markers = ["o", "s", "^", "D", "P", "X"]
    linestyles = ["-", "--", "-.", ":", (0, (5, 2)), (0, (3, 1, 1, 1))]
    for model, marker, linestyle in zip(MODEL_ORDER, markers, linestyles):
        vals = [aggregate[model]["macro_f1"]["values_by_seed"][str(s)] * 100 for s in SEEDS]
        ax.plot(x, vals, marker=marker, linestyle=linestyle, linewidth=1.8, markersize=5.4,
                label=MODEL_NAMES[model], color=MODEL_COLORS[model])
    ax.set_xticks(x, [str(s) for s in SEEDS])
    ax.set_xlabel("Training seed")
    ax.set_ylabel("Macro-F1 (%)")
    ax.set_ylim(88.5, 98.25)
    ax.grid(alpha=0.22, linewidth=0.7)
    ax.set_axisbelow(True)
    ax.set_title("Độ ổn định Macro-F1 qua ba lần huấn luyện", fontweight="bold", pad=7)
    ax.legend(ncol=2, loc="lower left", frameon=True, columnspacing=1.1, handlelength=2.4)
    ax.text(0.995, 0.015, "Giá trị chính xác xem Bảng 6–7", transform=ax.transAxes,
            ha="right", va="bottom", fontsize=7.5, color="#5A6573")
    return save_figure(fig, "fig_seed_stability.png")


def plot_pareto(aggregate, resources):
    fig, ax = plt.subplots(figsize=(FIGURE_WIDTH_IN, 4.55), constrained_layout=True)
    points = []
    for model in MODEL_ORDER:
        x = resources[model]["params"] / 1e6
        y = aggregate[model]["macro_f1"]["mean"] * 100
        points.append((model, x, y))
        size = 60 + resources[model]["checkpoint_mib"] * 5.2
        ax.scatter(x, y, s=size, color=MODEL_COLORS[model], alpha=0.88, edgecolor="white", linewidth=1.0, zorder=3)
    frontier = []
    for model, x, y in points:
        dominated = any((x2 <= x and y2 >= y) and (x2 < x or y2 > y) for _, x2, y2 in points)
        if not dominated:
            frontier.append((x, y))
    frontier.sort()
    ax.plot([p[0] for p in frontier], [p[1] for p in frontier], color="#5A6573", linewidth=1.2,
            linestyle="--", label="Biên Pareto (tham số, Macro-F1)", zorder=2)
    offsets = {
        "efficientnetb0": (5, 6), "mobilenetv3large": (5, -13), "resnet18": (-38, 5),
        "densenet121": (5, -11), "mobilenetv2": (5, 5), "shufflenetv2_1_0x": (5, 5),
    }
    for model, x, y in points:
        ax.annotate(MODEL_SHORT[model], (x, y), xytext=offsets[model], textcoords="offset points",
                    fontsize=8.3, fontweight="bold")
    ax.set_xlabel("Số tham số của mô hình 19 lớp (triệu)")
    ax.set_ylabel("Mean Macro-F1 trên test (%)")
    ax.set_title("Chất lượng – quy mô tham số", fontweight="bold", pad=6)
    ax.grid(alpha=0.22, linewidth=0.7)
    ax.set_axisbelow(True)
    ax.legend(loc="lower right", frameon=True)
    return save_figure(fig, "fig_quality_params.png")


def plot_class_heatmap(per_class, support):
    matrix = np.array([[per_class[m][c]["mean"] * 100 for c in CLASS_NAMES] for m in MODEL_ORDER])
    fig, axes = plt.subplots(2, 1, figsize=(FIGURE_WIDTH_IN, 6.65), constrained_layout=True)
    groups = [(0, 10), (10, 19)]
    images = []
    for panel_index, (ax, (start, stop)) in enumerate(zip(axes, groups), start=1):
        panel = matrix[:, start:stop]
        im = ax.imshow(panel, cmap="cividis", vmin=70, vmax=100, aspect="auto")
        images.append(im)
        classes = CLASS_NAMES[start:stop]
        abbreviations = CLASS_ABBR[start:stop]
        ax.set_xticks(np.arange(len(classes)), [f"{abbr}\nn={support[c]}" for abbr, c in zip(abbreviations, classes)])
        ax.set_yticks(np.arange(len(MODEL_ORDER)), [MODEL_NAMES[m] for m in MODEL_ORDER])
        ax.set_title(f"Khối {panel_index}: lớp {start + 1}–{stop}", fontweight="bold", fontsize=9.5, pad=4)
        for i in range(panel.shape[0]):
            for j in range(panel.shape[1]):
                color = "white" if panel[i, j] < 88 else "black"
                ax.text(j, i, f"{panel[i, j]:.1f}", ha="center", va="center", fontsize=7.1,
                        color=color, fontweight="bold" if panel[i, j] < 85 else "normal")
    cbar = fig.colorbar(images[0], ax=axes, pad=0.012, shrink=0.97)
    cbar.set_label("F1 (%)")
    fig.suptitle("Mean F1 theo lớp qua 3 seed; support test ghi dưới mã lớp", fontsize=10.7, fontweight="bold")
    return save_figure(fig, "fig_per_class_heatmap.png")


def plot_confusion(confusion_rate):
    selected = ["efficientnetb0", "mobilenetv3large"]
    fig, axes = plt.subplots(2, 1, figsize=(FIGURE_WIDTH_IN, 6.0), constrained_layout=True)
    for ax, model in zip(axes, selected):
        cm = confusion_rate[model].copy() * 100
        np.fill_diagonal(cm, 0)
        candidates = []
        for truth_idx, pred_idx in np.ndindex(cm.shape):
            if truth_idx != pred_idx and cm[truth_idx, pred_idx] > 0:
                candidates.append((cm[truth_idx, pred_idx], truth_idx, pred_idx))
        top = sorted(candidates, reverse=True)[:8]
        values = [value for value, _, _ in top][::-1]
        labels = [f"{CLASS_ABBR[t]} → {CLASS_ABBR[p]}" for _, t, p in top][::-1]
        y = np.arange(len(values))
        bars = ax.barh(y, values, color=MODEL_COLORS[model], alpha=0.9, height=0.64)
        ax.set_yticks(y, labels)
        ax.set_xlabel("Tỷ lệ trên lớp thật (%), mean qua 3 seed")
        ax.set_title(MODEL_NAMES[model], fontweight="bold", pad=4)
        ax.grid(axis="x", alpha=0.20, linewidth=0.7)
        ax.set_axisbelow(True)
        ax.bar_label(bars, labels=[f"{value:.2f}%" for value in values], padding=3, fontsize=7.7)
        ax.set_xlim(0, max(values) * 1.22)
    fig.suptitle("Tám hướng nhầm lẫn có tỷ lệ cao nhất (đã bỏ đường chéo đúng)", fontsize=10.6, fontweight="bold")
    return save_figure(fig, "fig_confusion_best_mobile.png")


def plot_bootstrap(comparison):
    data = comparison["paired_bootstrap_vs_mobilenetv3large"]
    keys = [
        "mobilenetv3large_minus_mobilenetv2",
        "mobilenetv3large_minus_resnet18",
        "mobilenetv3large_minus_efficientnetb0",
        "mobilenetv3large_minus_shufflenetv2_1_0x",
        "mobilenetv3large_minus_densenet121",
    ]
    labels = ["V3-Large − MobileNetV2", "V3-Large − ResNet-18", "V3-Large − EfficientNet-B0",
              "V3-Large − ShuffleNet V2", "V3-Large − DenseNet-121"]
    fig, ax = plt.subplots(figsize=(FIGURE_WIDTH_IN, 4.9), constrained_layout=True)
    positions = []
    pos_labels = []
    y = 0
    offsets = {42: -0.22, 123: 0.0, 2026: 0.22}
    seed_colors = {42: "#2E74B5", 123: "#F28E2B", 2026: "#2E7D32"}
    seed_markers = {42: "o", 123: "s", 2026: "^"}
    bounds = []
    for key, label in zip(keys, labels):
        positions.append(y)
        pos_labels.append(label)
        for seed in SEEDS:
            rec = data[key][str(seed)]
            mean = rec["observed_delta"] * 100
            lo = rec["ci95_lower"] * 100
            hi = rec["ci95_upper"] * 100
            bounds.extend([lo, hi])
            ax.errorbar(mean, y + offsets[seed], xerr=[[mean - lo], [hi - mean]], fmt=seed_markers[seed],
                        markersize=5.0, capsize=3, linewidth=1.2, color=seed_colors[seed],
                        label=f"seed {seed}" if y == 0 else None)
        y += 1
    x_min = math.floor((min(bounds) - 0.4) * 2) / 2
    x_max = math.ceil((max(bounds) + 0.4) * 2) / 2
    ax.axvspan(x_min, 0, color="#FCE8E6", alpha=0.45, zorder=0)
    ax.axvspan(0, x_max, color="#E6F4EA", alpha=0.45, zorder=0)
    ax.axvline(0, color="#333333", linewidth=1.2)
    ax.set_yticks(positions, pos_labels)
    ax.invert_yaxis()
    ax.set_xlim(x_min, x_max)
    ax.set_xlabel("Δ Macro-F1 (điểm %); dương ủng hộ V3-Large")
    ax.set_title("Paired bootstrap theo từng seed — 2.000 mẫu, CI percentile 95%", fontweight="bold", pad=6)
    ax.grid(axis="x", alpha=0.20, linewidth=0.7)
    ax.set_axisbelow(True)
    ax.legend(ncol=3, loc="upper center", bbox_to_anchor=(0.5, -0.16), frameon=True)
    return save_figure(fig, "fig_paired_bootstrap.png")


def plot_error_atlas(unanimous_wrong):
    items = unanimous_wrong[:12]
    class_abbreviations = dict(zip(CLASS_NAMES, CLASS_ABBR))
    columns = 3
    rows = math.ceil(len(items) / columns)
    fig, axes = plt.subplots(rows, columns, figsize=(FIGURE_WIDTH_IN, rows * 1.72), constrained_layout=True)
    axes = np.asarray(axes).reshape(-1)
    for ax, (rel, truth, preds) in zip(axes, items):
        image_path = DATASET_ROOT / Path(rel)
        ax.imshow(plt.imread(image_path))
        pred_text = "; ".join(f"{class_abbreviations.get(pred, pred)}:{count}" for pred, count in preds[:2])
        ax.set_title(f"{image_path.name}\nThật: {truth}\nDự đoán: {pred_text}", fontsize=6.8, pad=3)
        ax.axis("off")
    for ax in axes[len(items):]:
        ax.axis("off")
    fig.suptitle("11 ảnh bị cả 18 tổ hợp model × seed dự đoán sai", fontsize=10.5, fontweight="bold")
    return save_figure(fig, "fig_unanimous_error_atlas.png")


def generate_figures(comparison, aggregate, per_class, confusion_rate, support, resources, unanimous_wrong):
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    return {
        "overall": plot_overall(aggregate),
        "stability": plot_seed_stability(aggregate),
        "pareto": plot_pareto(aggregate, resources),
        "classes": plot_class_heatmap(per_class, support),
        "confusion": plot_confusion(confusion_rate),
        "bootstrap": plot_bootstrap(comparison),
        "atlas": plot_error_atlas(unanimous_wrong),
    }


def add_exec_summary(doc, aggregate, resources, support, bullet_id):
    add_heading(doc, "Tóm tắt điều hành", 1)
    add_callout(
        doc,
        "Kết luận chính: trong protocol đã khóa này, EfficientNet-B0 đạt mean Macro-F1 cao nhất (97,764% ± 0,064%) và đứng đầu cả ba seed. MobileNetV3-Large đứng thứ hai (97,323% ± 0,244%) và vẫn là ứng viên mobile hợp lý, nhưng quyết định triển khai cuối cùng cần benchmark độ trễ, RAM và năng lượng của cả sáu model trên cùng thiết bị/runtime.",
    )
    add_paragraph(
        doc,
        "Báo cáo phân tích 18 checkpoint (6 mô hình × 3 seed) trên cùng tập test 9.829 ảnh đã khóa. Tiêu chí chính được chốt trước là Macro-F1 vì phân bố lớp rất lệch: stem_borer, brown_planthopper và leaf_roller chiếm 76,539% test. Accuracy vẫn được báo cáo nhưng không được dùng một mình để kết luận chất lượng trên lớp hiếm.",
    )
    rows = [["Hạng", "Mô hình", "Accuracy", "Macro-F1", "SD F1", "Tham số", "Checkpoint"]]
    for rank, model in enumerate(MODEL_ORDER, 1):
        rows.append([
            str(rank),
            MODEL_NAMES[model],
            pct(aggregate[model]["accuracy"]["mean"], 3),
            pct(aggregate[model]["macro_f1"]["mean"], 3),
            f"{aggregate[model]['macro_f1']['sample_std']*100:.3f} đ.%",
            f"{PARAMS[model]/1e6:.3f} M",
            f"{resources[model]['checkpoint_mib']:.2f} MiB",
        ])
    add_table_custom(doc, rows, [520, 1900, 1230, 1230, 1050, 1200, 2230], font_size=7.6,
                     caption_title="table_executive_ranking", caption_description="Xếp hạng kết quả trung bình qua ba seed và quy mô checkpoint")
    add_caption(doc, "Bảng 1. Xếp hạng mô tả theo mean Macro-F1. SD là sample SD qua 3 training seed, không phải khoảng tin cậy 95%.")
    add_bullet(doc, "EfficientNet-B0 có hiệu năng trung bình cao nhất và độ biến thiên seed thấp nhất; hơn MobileNetV3-Large **0,441 điểm % Macro-F1** theo mean.", bullet_id)
    add_bullet(doc, "MobileNetV3-Large đứng #2 ở seed 123 và 2026, #3 ở seed 42, sát ResNet-18; lợi thế mobile của kiến trúc là cơ sở chọn ứng viên, chưa phải bằng chứng runtime trên OPPO A54.", bullet_id)
    add_bullet(doc, "ShuffleNet V2 1.0× nhỏ nhất (1,273 triệu tham số) nhưng Macro-F1 thấp nhất và SD lớn nhất; đây là trade-off chất lượng–độ gọn, không tự động đồng nghĩa chậm/nhanh trên thiết bị.", bullet_id)
    add_bullet(doc, "Các lớp khó nhất xuyên mô hình là downy_mildew, tungro, blast, bacterial_leaf_streak và weeds. Hai lớp đạt 100% ở mọi run có support rất nhỏ (n=12 và n=29), vì vậy không nên quảng bá như kết quả chắc chắn.", bullet_id)
    add_bullet(doc, "Kết quả chỉ có giá trị nội bộ trên split ảnh phân tầng. Split không nhóm theo cây, ruộng hoặc phiên chụp; chưa có external/geographic/temporal test.", bullet_id)


def add_scope_and_data(doc, support, bullet_id):
    add_heading(doc, "1. Câu hỏi nghiên cứu và phạm vi bằng chứng", 1)
    add_paragraph(doc, "Nghiên cứu trả lời câu hỏi: khi giữ cố định dữ liệu, cách tiền xử lý, augmentation, loss, lịch huấn luyện, tiêu chí chọn checkpoint và framework, sáu backbone CNN khác nhau đem lại trade-off nào giữa chất lượng phân loại, độ ổn định và quy mô mô hình cho hệ RiceAI?")
    add_heading(doc, "1.1 Đối tượng so sánh", 2)
    add_bullet(doc, "MobileNetV2 và MobileNetV3-Large: họ mạng nhẹ định hướng thiết bị biên/mobile [1], [2].", bullet_id)
    add_bullet(doc, "ResNet-18: baseline residual cổ điển, có identity shortcut [3].", bullet_id)
    add_bullet(doc, "EfficientNet-B0: baseline được xây dựng quanh compound scaling depth–width–resolution [4].", bullet_id)
    add_bullet(doc, "ShuffleNet V2 1.0×: kiến trúc nhẹ dùng channel split/shuffle và nhấn mạnh chi phí truy cập bộ nhớ [5].", bullet_id)
    add_bullet(doc, "DenseNet-121: kết nối dày đặc để tái sử dụng đặc trưng [6].", bullet_id)
    add_heading(doc, "1.2 Phạm vi kết luận", 2)
    add_callout(doc, "Phát biểu đúng: “EfficientNet-B0 đạt mean Macro-F1 cao nhất trong campaign RiceAI V2 trên split test cố định này.” Phát biểu không được phép: “EfficientNet-B0 tốt nhất tuyệt đối cho mọi ruộng lúa, điện thoại và điều kiện chụp.”", warning=True)

    add_heading(doc, "2. Dữ liệu, làm sạch và chiến lược chia tập", 1)
    add_paragraph(doc, "Tập dữ liệu sau làm sạch gồm 65.516 ảnh RGB thuộc 19 nhãn canonical. Đây không phải 19 bệnh thuần túy: ngoài bệnh còn có sâu hại/đối tượng (beetle, planthoppers, snail, leaf roller, stem borer), cỏ dại và trạng thái normal. Vì vậy báo cáo dùng cụm “19 nhóm bệnh, dịch hại/đối tượng và trạng thái lúa”.")
    rows = [
        ["Bước", "Số lượng", "Diễn giải"],
        ["File quét ban đầu", "67.342", "Nguồn trước curation"],
        ["Exact duplicate bị loại", "102", "Loại bản sao chính xác"],
        ["Duplicate chéo nhãn bị loại", "1.724 file / 862 nhóm", "Loại toàn bộ nhóm xung đột nhãn trước split"],
        ["Ảnh duy nhất dùng", "65.516", "Đã full SHA verification trước khóa V2"],
        ["Train / validation / test", "45.858 / 9.829 / 9.829", "70% / 15% / 15%, stratified image-level, seed 2026"],
    ]
    add_table_custom(doc, rows, [2500, 2100, 4760], font_size=8.7,
                     caption_title="table_dataset_curation", caption_description="Luồng làm sạch và chia tập dữ liệu")
    add_caption(doc, "Bảng 2. Provenance dữ liệu và số lượng dùng trong campaign.")
    largest = sorted(support.items(), key=lambda x: x[1], reverse=True)[:3]
    rare = sorted(support.items(), key=lambda x: x[1])[:5]
    add_paragraph(doc, f"Ba lớp lớn nhất là {largest[0][0]} (n={largest[0][1]}), {largest[1][0]} (n={largest[1][1]}) và {largest[2][0]} (n={largest[2][1]}), tổng cộng 7.523/9.829 ảnh test (76,539%). Năm lớp nhỏ nhất là " + ", ".join(f"{name} (n={n})" for name, n in rare) + ".")
    add_callout(doc, "Giới hạn trọng yếu: split được phân tầng ở cấp ảnh, không group theo cây, ruộng hay phiên chụp vì metadata không chứa các group ID đó. Near-duplicate hoặc ảnh tương quan cùng bối cảnh có thể nằm ở nhiều split, làm kết quả nội bộ lạc quan. Roberts et al. khuyến nghị blocking khi dữ liệu có cấu trúc phụ thuộc [15]; Kapoor và Narayanan cũng chỉ ra leakage có thể thổi phồng kết quả khoa học ML [16].", warning=True)


def add_architectures(doc, aggregate, per_class, resources):
    add_heading(doc, "3. Sáu kiến trúc và giả thuyết kỹ thuật", 1)
    add_paragraph(doc, "Tất cả backbone dùng official ImageNet-1K weights từ torchvision 0.17.1, classifier được thay bằng 19 đầu ra và dropout 0,4. Vì phiên bản weight enum không đồng nhất (V1/V2), so sánh thực tế là “kiến trúc + official pretrained checkpoint + common recipe”, không phải phép cô lập topology thuần túy.")
    rows = [["Mô hình", "Thành phần kiến trúc nổi bật", "Official weight", "Tham số 19 lớp", "Giả thuyết trước kết quả"]]
    architecture_rows = {
        "mobilenetv2": ("Depthwise separable; inverted residual; linear bottleneck", "Gọn, baseline mobile; có thể thiếu capacity ở lớp hình thái gần"),
        "mobilenetv3large": ("Hardware-aware NAS/NetAdapt; SE; hard-swish", "Trade-off chất lượng–latency tốt trên mobile, cần đo runtime thật"),
        "resnet18": ("Basic residual blocks và identity shortcut", "Tối ưu ổn định, baseline mạnh nhưng nhiều tham số"),
        "efficientnetb0": ("MBConv + compound scaling cân bằng", "Độ chính xác/hiệu quả cao nhờ phân bổ capacity cân bằng"),
        "shufflenetv2_1_0x": ("Channel split, channel shuffle, pointwise groups", "Rất gọn nhưng có nguy cơ under-capacity cho 19 lớp"),
        "densenet121": ("Dense connectivity, concatenate feature maps", "Tái sử dụng đặc trưng tốt; có thể tốn memory traffic"),
    }
    for model in ["mobilenetv2", "mobilenetv3large", "resnet18", "efficientnetb0", "shufflenetv2_1_0x", "densenet121"]:
        feature, hypothesis = architecture_rows[model]
        rows.append([MODEL_NAMES[model], feature, WEIGHT_IDS[model].split(".")[-1], fmt_int(PARAMS[model]), hypothesis])
    add_table_custom(doc, rows, [1400, 2300, 1150, 1250, 3260], font_size=7.3,
                     caption_title="table_architecture", caption_description="Đặc trưng kiến trúc, pretrained weight và giả thuyết kỹ thuật")
    add_caption(doc, "Bảng 3. Lý do chọn sáu backbone. Giả thuyết kiến trúc không thay thế số liệu thực nghiệm.")
    add_heading(doc, "3.1 Nguyên tắc diễn giải nguyên nhân", 2)
    add_paragraph(doc, "Khi EfficientNet-B0 hoặc MobileNetV3-Large đạt kết quả cao, compound scaling, MBConv, squeeze-and-excitation hay hardware-aware search là các cơ chế kiến trúc hợp lý để thảo luận [2], [4]. Tuy nhiên campaign không thực hiện ablation từng thành phần; do đó chỉ được viết “có thể liên quan/phù hợp với”, không được viết “do thành phần X gây ra”.")
    add_paragraph(doc, "Tương tự, ShuffleNet V2 nhấn mạnh rằng FLOPs không đủ dự báo tốc độ vì memory access cost và đặc tính nền tảng cũng quan trọng [5]. DenseNet có tái sử dụng feature [6], nhưng concatenate nhiều feature map có thể làm traffic bộ nhớ đáng kể; điều này cần đo profiler thay vì suy từ số tham số.")


def add_protocol(doc, audit, manifests, resources, bullet_id, number_id):
    add_heading(doc, "4. Protocol huấn luyện, chọn checkpoint và tái lập", 1)
    add_heading(doc, "4.1 Tiền xử lý và augmentation", 2)
    add_bullet(doc, "Resize trực tiếp 224×224 bằng bilinear + antialias; chuyển RGB; ToTensor về [0,1].", bullet_id)
    add_bullet(doc, "Chuẩn hóa ImageNet mean [0,485; 0,456; 0,406] và std [0,229; 0,224; 0,225].", bullet_id)
    add_bullet(doc, "Train augmentation: horizontal flip p=0,5; RandomAffine rotation ±20,16°, translation 0,1, scale 0,85–1,15; ColorJitter brightness/contrast 0,15.", bullet_id)
    add_bullet(doc, "Validation/test không augmentation; test shuffle=false và giữ thứ tự manifest.", bullet_id)
    add_heading(doc, "4.2 Hai giai đoạn transfer learning", 2)
    steps = [
        "Stage head: khóa backbone, huấn luyện classifier bằng Adam, LR=1e-3, tối đa 15 epoch.",
        "Stage fine-tune: mở toàn mạng trừ BatchNorm, khởi tạo lại optimizer/scheduler, Adam LR=1e-5, tối đa 30 epoch.",
        "Weighted cross-entropy theo effective number of samples, β=0,999; chuẩn hóa mean rồi clip [0,25; 5], dựa trên Cui et al. [7].",
        "ReduceLROnPlateau theo validation Macro-F1: factor=0,2, patience=2, min LR=1e-7.",
        "Early stopping patience=5; bộ đếm reset khi đổi stage nhưng best score/checkpoint là global xuyên hai stage.",
        "Checkpoint được chọn duy nhất theo validation Macro-F1; test không tham gia huấn luyện hoặc chọn model.",
    ]
    for index, step in enumerate(steps, start=1):
        add_number(doc, step, number_id, index=index)
    rows = [
        ["Thành phần", "Cấu hình khóa", "Ghi chú kiểm toán"],
        ["Input / batch / loader", "RGB 224×224; batch=32; workers=4", "workers lấy từ command trong 18 run_manifest"],
        ["Classifier", "19 đầu ra; dropout=0,4", "Tham số báo cáo đã gồm classifier mới"],
        ["Optimizer", "Adam; head LR=1e−3; fine-tune LR=1e−5", "Không truyền weight_decay trong mã; mặc định Adam = 0"],
        ["BatchNorm", "Đóng băng running statistics và tham số ở cả hai stage", "Giữ cùng chính sách cho sáu backbone"],
        ["Loss", "Weighted cross-entropy; effective number β=0,999", "Weight chuẩn hóa mean, clip [0,25; 5]"],
        ["Scheduler", "ReduceLROnPlateau; factor=0,2; patience=2; min LR=1e−7", "Theo validation Macro-F1"],
        ["Dừng / chọn", "Early stop patience=5; checkpoint global tốt nhất", "Test không tham gia chọn checkpoint"],
        ["Lặp", "Seed 42, 123, 2026", "Ba lần huấn luyện độc lập/mô hình"],
    ]
    add_table_custom(doc, rows, [1900, 3510, 3950], font_size=7.8,
                     caption_title="table_locked_training_protocol", caption_description="Cấu hình huấn luyện khóa dùng chung cho sáu kiến trúc")
    add_caption(doc, "Bảng 4a. Cấu hình thực thi được đối chiếu giữa contract, source code và run manifest.")
    add_paragraph(doc, "Revision V2 tăng ngân sách fine-tune từ 15 lên 30 epoch sau pilot chỉ dựa trên validation của MobileNetV2; pilot V1 bị loại khỏi kết quả chính và thay đổi được khóa trước khi mở test. Đây là protocol revision informed by validation, không phải test-driven tuning.")
    add_heading(doc, "4.3 Kiểm soát tái lập", 2)
    add_paragraph(doc, "Campaign đặt seed Python/NumPy/PyTorch/CUDA, bật deterministic algorithms, tắt cuDNN benchmark, bật cuDNN deterministic, đặt CUBLAS_WORKSPACE_CONFIG=:4096:8 và tắt TF32. PyTorch vẫn lưu ý không thể bảo đảm tái lập tuyệt đối qua release, platform hoặc CPU/GPU khác [14].")
    rows = [
        ["Thành phần khóa", "Định danh / SHA-256"],
        ["Benchmark contract", "e54ea9e3adb1fbb6a82ea29781dc6e8b066390f279aebb175778086f6b7e4514"],
        ["Campaign lock V2", "63f9c337986b3d4cfaf20c1745df85e2cfd36a95e7557e894decc781e87646b0"],
        ["Split manifest", "bc2c7c8c72a14968c402cbb06fc0eebdde57a8b00a5ac5c88962716f06024160"],
        ["Runtime", "Python 3.11.9; torch 2.2.1+cu121; torchvision 0.17.1+cu121; CUDA"],
        ["Training seeds", "42, 123, 2026"],
    ]
    add_table_custom(doc, rows, [2600, 6760], font_size=8.0,
                     caption_title="table_reproducibility", caption_description="Khóa provenance và môi trường campaign")
    add_caption(doc, "Bảng 4b. Các khóa cần công bố để tái lập campaign.")
    add_heading(doc, "4.4 Diễn biến hội tụ của 18 run", 2)
    rows = [["Mô hình", "Seed", "Epoch head / fine-tune", "Checkpoint tốt nhất", "Cảnh báo budget"]]
    for model in ["mobilenetv2", "mobilenetv3large", "resnet18", "efficientnetb0", "shufflenetv2_1_0x", "densenet121"]:
        for record in resources[model]["convergence"]:
            warning = "Best ở epoch fine-tune cuối" if record["best_is_last_finetune"] else "Best trước epoch cuối"
            rows.append([
                MODEL_NAMES[model],
                str(record["seed"]),
                f"{record['head_epochs']} / {record['finetune_epochs']}",
                f"{record['best_stage']} {record['best_stage_epoch']}",
                warning,
            ])
    add_table_custom(doc, rows, [1750, 620, 1800, 1900, 3290], font_size=7.3,
                     caption_title="table_convergence_runs", caption_description="Số epoch thực chạy và vị trí checkpoint validation tốt nhất ở đủ 18 run")
    add_caption(doc, "Bảng 4c. Mọi checkpoint tốt nhất đều nằm ở fine-tune; ba run đạt best đúng epoch fine-tune cuối nên chưa đủ bằng chứng rằng mọi backbone đã hội tụ hoàn toàn.")
    add_paragraph(doc, "Ba trường hợp MobileNetV2–2026, ShuffleNet V2–42 và DenseNet-121–2026 đạt best validation Macro-F1 đúng epoch fine-tune cuối. Vì vậy common budget vẫn là một giới hạn công bằng theo ngân sách, không phải bằng chứng rằng mỗi kiến trúc đã đạt tối ưu riêng; mô hình hội tụ chậm có thể bị bất lợi.")
    add_callout(doc, "Audit nuance: full hash verification 65.516/65.516 ảnh được chạy trước V2 trên cùng manifest; V2 lock khóa SHA của report đó. Report preflight mang provenance contract V1, nên cách viết chính xác là “report kiểm tra cùng manifest được V2 lock tham chiếu”, không phải “report được sinh trực tiếp từ V2 contract”.")


def add_test_metrics(doc, audit, bullet_id):
    add_heading(doc, "5. Quy trình mở test và thước đo đánh giá", 1)
    add_heading(doc, "5.1 Quy trình test khóa", 2)
    add_paragraph(doc, "Script huấn luyện chỉ nạp train/validation. Script final_evaluate_v2.py tách riêng và trước khi mở test đã xác minh đủ 18 manifest complete, test_evaluated=false, provenance model/seed/contract/split/lock khớp và SHA-256 của checkpoint khớp.")
    rows = [
        ["Mốc audit", "UTC", "Ý nghĩa"],
        ["Mở test", audit["opened_at_utc"], "Guard chấp nhận cụm unlock I_UNDERSTAND_TEST_IS_FINAL"],
        ["Hoàn tất", audit["completed_at_utc"], "18 metrics JSON + 18 predictions CSV + aggregate/paired bootstrap"],
        ["Trạng thái", audit["status"], "Mỗi CSV đúng 9.829 hàng; mọi confusion matrix tổng 9.829"],
    ]
    add_table_custom(doc, rows, [1700, 3100, 4560], font_size=8.4,
                     caption_title="table_test_audit", caption_description="Mốc mở test và kiểm toán đầu ra")
    add_caption(doc, "Bảng 5. Audit mở test. Cơ chế là procedural guard; không phải khóa mật mã không thể chạy lại.")

    add_heading(doc, "5.2 Định nghĩa metric", 2)
    add_metric_definition_table(doc)
    add_caption(doc, "Bảng 5a. Công thức dùng đúng theo mã đánh giá; nếu mẫu số bằng 0, implementation đặt precision/recall/F1 tương ứng bằng 0.")
    add_paragraph(doc, "Macro averaging cho mỗi lớp trọng số bằng nhau, vì vậy phù hợp làm endpoint chính khi class support mất cân bằng [8], [9]. Accuracy vẫn hữu ích để mô tả tổng số ảnh đúng, nhưng có thể bị ba lớp lớn chi phối.")
    add_heading(doc, "5.3 Hai nguồn bất định không được trộn", 2)
    add_bullet(doc, "Mean ± sample SD qua 3 training seed mô tả biến thiên do huấn luyện. Với n=3, đây là thống kê mô tả, không phải 95% CI. Việc báo nhiều lần chạy tránh phụ thuộc một seed may mắn [11], [12].", bullet_id)
    add_bullet(doc, "Bootstrap 2.000 mẫu trên ảnh test ước lượng sampling uncertainty có điều kiện trên checkpoint cố định; dùng percentile 2,5%–97,5% [10].", bullet_id)
    add_bullet(doc, "Paired bootstrap giữ cặp dự đoán trên cùng ảnh nhưng chỉ thực hiện MobileNetV3-Large so với từng model, riêng theo seed. probability_delta_gt_zero là tỷ lệ bootstrap Δ>0, không phải p-value hay posterior probability.", bullet_id)
    add_callout(doc, "Không được báo các chỉ số chưa tính: Weighted-F1, balanced accuracy, ROC-AUC, specificity, calibration/ECE, FLOPs, latency, năng lượng hoặc peak RAM. Chúng chỉ có thể xuất hiện trong mục “cần đo tiếp”.", warning=True)
    add_heading(doc, "5.4 Cách đọc và giới hạn của từng bằng chứng", 2)
    rows = [
        ["Bằng chứng", "Dùng để trả lời", "Không được suy ra"],
        ["Mean Macro-F1", "Chất lượng cân bằng qua 19 lớp", "Ý nghĩa thống kê hoặc tổng quát ngoài split"],
        ["Accuracy", "Tỷ lệ đúng trên 9.829 ảnh", "Chất lượng tốt ở lớp hiếm"],
        ["Precision / Recall theo lớp", "Dương tính giả và bỏ sót của từng nhãn", "Nguyên nhân lỗi nếu chưa audit ảnh/nhãn"],
        ["Sample SD qua 3 seed", "Biến thiên do ba lần huấn luyện", "CI95; phân phối đầy đủ của seed"],
        ["Bootstrap ảnh test", "Bất định lấy mẫu có điều kiện trên checkpoint", "Bất định huấn luyện; p-value; posterior probability"],
        ["Số tham số / checkpoint", "Quy mô mạng và file PyTorch", "FLOPs, latency, RAM, pin hoặc kích thước APK"],
    ]
    add_table_custom(doc, rows, [1900, 3440, 4020], font_size=7.8,
                     caption_title="table_evidence_interpretation", caption_description="Câu hỏi được phép trả lời và giới hạn suy luận của mỗi thước đo")
    add_caption(doc, "Bảng 5b. Khung diễn giải bắt buộc để tránh trộn các nguồn bất định hoặc suy diễn quá dữ liệu.")
    rows = [
        ["Nội dung dùng làm thước đo", "Nguồn phương pháp", "Vai trò trong báo cáo"],
        ["Precision, recall, F1, macro averaging", "Sokolova & Lapalme [8]; scikit-learn [9]", "Định nghĩa metric và cách đọc theo lớp"],
        ["Bootstrap percentile", "Efron & Tibshirani [10]", "CI lấy mẫu ảnh test và paired delta"],
        ["Báo phân bố qua nhiều lần chạy", "Reimers & Gurevych [11]; Bouthillier et al. [12]", "Báo từng seed, mean và sample SD"],
        ["Tách chọn model khỏi đánh giá cuối", "Cawley & Talbot [13]", "Không dùng final test để tuning"],
        ["Giới hạn tái lập theo nền tảng", "PyTorch official docs [14]", "Không hứa bitwise reproducibility qua máy/release"],
    ]
    add_table_custom(doc, rows, [2700, 3020, 3640], font_size=7.7,
                     caption_title="table_metric_source_map", caption_description="Ánh xạ trực tiếp giữa thước đo, nguồn chính thức và cách sử dụng")
    add_caption(doc, "Bảng 5c. Các link URL/DOI bấm được được in đầy đủ ở mục Tài liệu tham khảo.")


def add_overall_results(doc, aggregate, figures, resources):
    add_heading(doc, "6. Kết quả tổng thể và độ ổn định qua seed", 1)
    add_figure(doc, figures["overall"], "Hình 1. Mean Macro-F1 và Accuracy trên hai panel xếp dọc; thanh sai số là sample SD qua ba training seed.", alt="Hai biểu đồ thanh ngang cỡ chữ lớn xếp hạng sáu mô hình theo Macro-F1 và Accuracy. EfficientNet-B0 đứng đầu, ShuffleNet V2 thấp nhất.")
    rows = [["Mô hình", "Accuracy", "Macro-P", "Macro-R", "Macro-F1", "Khoảng seed F1"]]
    for model in MODEL_ORDER:
        agg = aggregate[model]
        vals = list(agg["macro_f1"]["values_by_seed"].values())
        rows.append([
            MODEL_NAMES[model],
            f"{agg['accuracy']['mean']*100:.3f} ± {agg['accuracy']['sample_std']*100:.3f}",
            f"{agg['macro_precision']['mean']*100:.3f} ± {agg['macro_precision']['sample_std']*100:.3f}",
            f"{agg['macro_recall']['mean']*100:.3f} ± {agg['macro_recall']['sample_std']*100:.3f}",
            f"{agg['macro_f1']['mean']*100:.3f} ± {agg['macro_f1']['sample_std']*100:.3f}",
            f"{(max(vals)-min(vals))*100:.3f} đ.%",
        ])
    add_table_custom(doc, rows, [1750, 1520, 1520, 1520, 1520, 1530], font_size=7.7,
                     caption_title="table_aggregate_results", caption_description="Kết quả mean và sample SD qua ba training seed")
    add_caption(doc, "Bảng 6. Kết quả tổng hợp (%). Mỗi ô metric trình bày mean ± sample SD qua 3 seed.")
    add_paragraph(doc, "EfficientNet-B0 đứng đầu đồng thời ở Accuracy, Macro-Precision, Macro-Recall và Macro-F1. Mean Macro-F1 97,764% cao hơn MobileNetV3-Large 0,441 điểm %, ResNet-18 0,887 điểm %, DenseNet-121 1,107 điểm %, MobileNetV2 1,580 điểm % và ShuffleNet V2 6,793 điểm %.")
    add_paragraph(doc, "Khoảng seed của EfficientNet-B0 chỉ 0,124 điểm %, nhỏ nhất trong sáu model. MobileNetV3-Large cũng ổn định (0,483 điểm %). Ngược lại ShuffleNet V2 có khoảng 3,379 điểm % và sample SD 1,691 điểm %, cho thấy kết quả nhạy seed nhất trong protocol này.")
    add_figure(doc, figures["stability"], "Hình 2. Macro-F1 của từng seed. EfficientNet-B0 đứng đầu cả ba; ShuffleNet V2 đứng cuối cả ba.", alt="Biểu đồ đường cho Macro-F1 theo seed 42, 123 và 2026 của sáu mô hình.")
    add_heading(doc, "6.1 Kết quả từng seed", 2)
    rows = [["Seed", "#1", "#2", "#3", "#4", "#5", "#6"]]
    for seed in SEEDS:
        ordered = sorted(MODEL_NAMES, key=lambda m: aggregate[m]["macro_f1"]["values_by_seed"][str(seed)], reverse=True)
        rows.append([str(seed)] + [f"{MODEL_SHORT[m]} {aggregate[m]['macro_f1']['values_by_seed'][str(seed)]*100:.3f}%" for m in ordered])
    add_table_custom(doc, rows, [700, 1440, 1440, 1440, 1440, 1440, 1460], font_size=7.5,
                     caption_title="table_seed_ranking", caption_description="Xếp hạng Macro-F1 trong từng training seed")
    add_caption(doc, "Bảng 7. Xếp hạng theo từng seed; không chọn “best seed” làm kết quả chính.")
    add_heading(doc, "6.2 Mức chênh thực tế, thứ hạng và cân bằng precision–recall", 2)
    rank_by_seed = {}
    for seed in SEEDS:
        ordered = sorted(MODEL_NAMES, key=lambda m: aggregate[m]["macro_f1"]["values_by_seed"][str(seed)], reverse=True)
        rank_by_seed[seed] = {model: rank + 1 for rank, model in enumerate(ordered)}
    rows = [["Mô hình", "Đúng / sai: seed 42", "Đúng / sai: seed 123", "Đúng / sai: seed 2026", "Hạng F1", "Macro-R − Macro-P"]]
    for model in MODEL_ORDER:
        error_counts = resources[model]["error_counts_by_seed"]
        rows.append([
            MODEL_NAMES[model],
            f"{9829-error_counts[42]} / {error_counts[42]}",
            f"{9829-error_counts[123]} / {error_counts[123]}",
            f"{9829-error_counts[2026]} / {error_counts[2026]}",
            "/".join(str(rank_by_seed[seed][model]) for seed in SEEDS),
            f"{(aggregate[model]['macro_recall']['mean']-aggregate[model]['macro_precision']['mean'])*100:+.3f} đ.%",
        ])
    add_table_custom(doc, rows, [1580, 1560, 1560, 1560, 1100, 2000], font_size=7.25,
                     caption_title="table_correct_error_rank_pr", caption_description="Số ảnh đúng sai theo checkpoint, thứ hạng Macro-F1 và chênh lệch macro recall precision")
    add_caption(doc, "Bảng 7a. Mỗi seed dự đoán cùng 9.829 ảnh. “Hạng F1” lần lượt ứng với seed 42/123/2026.")
    eff_errors = statistics.mean(resources["efficientnetb0"]["error_counts_by_seed"].values())
    v3_errors = statistics.mean(resources["mobilenetv3large"]["error_counts_by_seed"].values())
    add_paragraph(doc, f"EfficientNet-B0 tạo trung bình {eff_errors:.1f} quyết định sai/run, còn MobileNetV3-Large {v3_errors:.1f}; chênh khoảng {v3_errors-eff_errors:.1f} ảnh/run. Tuy nhiên chênh Macro-F1 giữa hai model là 0,441 điểm %, lớn hơn chênh Accuracy vì Macro-F1 cho mỗi lớp trọng số bằng nhau và phản ánh rõ hơn chất lượng ở lớp nhỏ.")
    gaps = {
        model: (aggregate[model]["macro_recall"]["mean"] - aggregate[model]["macro_precision"]["mean"]) * 100
        for model in MODEL_ORDER
    }
    widest_model = max(gaps, key=gaps.get)
    add_paragraph(doc, f"Macro-Recall cao hơn Macro-Precision ở cả sáu mô hình. Chênh lớn nhất thuộc {MODEL_NAMES[widest_model]} ({gaps[widest_model]:.3f} điểm %), gợi ý dương tính giả đáng chú ý ở một số lớp. Đây là mô tả error profile; không được quy nguyên nhân cho weighted loss nếu chưa có ablation loss.")
    add_heading(doc, "6.3 Validation–test gap nội bộ", 2)
    add_paragraph(doc, "Cả 18 checkpoint có test Macro-F1 thấp hơn best validation Macro-F1. Mean gap theo mô hình lần lượt là " + "; ".join(
        f"{MODEL_SHORT[model]} {resources[model]['test_gap']*100:+.3f} điểm %" for model in MODEL_ORDER
    ) + ". Đây là chênh lệch nội bộ giữa checkpoint đã chọn trên validation và test cố định; có thể chứa selection optimism hoặc khác biệt lấy mẫu, nhưng không tự chứng minh overfitting nhân quả.")


def add_bootstrap_results(doc, comparison, figures):
    add_heading(doc, "7. So sánh bất định bằng paired bootstrap", 1)
    add_figure(doc, figures["bootstrap"], "Hình 3. Chênh lệch Macro-F1 MobileNetV3-Large trừ mô hình đối chiếu; CI95 theo từng seed.", alt="Forest plot của 15 chênh lệch paired bootstrap. Điểm dương ủng hộ MobileNetV3-Large, điểm âm ủng hộ mô hình đối chiếu.")
    add_heading(doc, "7.1 Cách đọc forest plot", 2)
    rows = [
        ["Dấu / vị trí CI", "Diễn giải được phép", "Không được viết"],
        ["Δ > 0", "Checkpoint V3-Large có Macro-F1 quan sát cao hơn đối chiếu", "V3 tốt hơn tuyệt đối"],
        ["CI hoàn toàn dương", "Bất định lấy mẫu ảnh test trong run ủng hộ V3", "Ý nghĩa thống kê tổng thể qua seed"],
        ["CI chứa 0", "Dữ liệu test/run chưa tách biệt rõ hai checkpoint", "Hai model tương đương"],
        ["CI hoàn toàn âm", "Bất định lấy mẫu ảnh test trong run ủng hộ đối chiếu", "Đối chiếu thắng trên mọi domain"],
    ]
    add_table_custom(doc, rows, [1700, 3900, 3760], font_size=7.8,
                     caption_title="table_bootstrap_reading", caption_description="Quy tắc diễn giải dấu chênh lệch và khoảng tin cậy paired bootstrap")
    add_caption(doc, "Bảng 8a. Pairing nằm ở dự đoán của hai checkpoint trên cùng ảnh test; không phải pairing của quá trình huấn luyện.")
    add_paragraph(doc, "Cùng số seed giữa hai kiến trúc không biến chúng thành hai training replicate được ghép cặp. Việc ghép cặp chỉ xuất hiện khi resample cùng 9.829 ảnh và lấy chênh lệch metric của hai vector dự đoán. Có 15 contrast (5 đối chiếu × 3 seed), chưa hiệu chỉnh multiplicity và chưa có mô hình phân cấp gộp seed.")
    data = comparison["paired_bootstrap_vs_mobilenetv3large"]
    mapping = [
        ("mobilenetv3large_minus_mobilenetv2", "V3 − MobileNetV2"),
        ("mobilenetv3large_minus_resnet18", "V3 − ResNet-18"),
        ("mobilenetv3large_minus_efficientnetb0", "V3 − EfficientNet-B0"),
        ("mobilenetv3large_minus_shufflenetv2_1_0x", "V3 − ShuffleNet V2"),
        ("mobilenetv3large_minus_densenet121", "V3 − DenseNet-121"),
    ]
    rows = [["Đối chiếu", "Seed", "Δ F1 (đ.%)", "CI95 (đ.%)", "Tỷ lệ Δ>0", "Đọc đúng"]]
    for key, label in mapping:
        for seed in SEEDS:
            rec = data[key][str(seed)]
            lo = rec["ci95_lower"] * 100
            hi = rec["ci95_upper"] * 100
            reading = "CI loại 0" if lo > 0 or hi < 0 else "CI chứa 0"
            rows.append([
                label,
                str(seed),
                f"{rec['observed_delta']*100:+.3f}",
                f"[{lo:+.3f}; {hi:+.3f}]",
                f"{rec['probability_delta_gt_zero']:.4f}",
                reading,
            ])
    add_table_custom(doc, rows, [1900, 620, 1150, 1700, 1250, 2740], font_size=7.1,
                     caption_title="table_paired_bootstrap", caption_description="Paired bootstrap của MobileNetV3-Large so với từng mô hình theo cùng seed")
    add_caption(doc, "Bảng 8b. Paired bootstrap 2.000 mẫu/đối chiếu/seed. Không có hiệu chỉnh multiple comparisons và không có phép gộp ba seed.")
    add_paragraph(doc, "MobileNetV3-Large hơn MobileNetV2 ở cả ba giá trị quan sát; CI loại 0 ở seed 42 và 123 nhưng chứa 0 ở seed 2026. So với ShuffleNet V2, V3 hơn rõ ở cả ba seed với CI hoàn toàn dương. So với ResNet-18, dấu và độ chắc chắn không đồng nhất: seed 42 hơi nghiêng về ResNet, seed 123 gần nhau, seed 2026 ủng hộ V3.")
    add_paragraph(doc, "EfficientNet-B0 có observed Macro-F1 cao hơn V3 ở cả ba seed. Tuy nhiên CI của V3−EffB0 chỉ hoàn toàn âm ở seed 2026; hai seed còn lại chứa 0. Do không có mô hình phân cấp hoặc kiểm định gộp qua seed, báo cáo chỉ kết luận “EfficientNet-B0 có mean cao hơn và thắng quan sát ở cả ba seed”, không tuyên bố significance tổng thể.")


def add_model_analysis(doc, aggregate, per_class, resources):
    add_heading(doc, "8. Phân tích sâu từng mô hình", 1)
    analyses = {
        "efficientnetb0": [
            "Kết quả: #1 theo mean Macro-F1 (97,764% ± 0,064%), Accuracy 99,298% ± 0,027%; đứng đầu cả ba seed.",
            "Ổn định: range seed chỉ 0,124 điểm %, nhỏ nhất campaign. Đây là bằng chứng mô tả mạnh nhất về độ ổn định huấn luyện trong sáu model, dù n=3 còn nhỏ.",
            "Lớp khó: downy_mildew 92,939%, tungro 94,142%, blast 94,946%, weeds 95,149%. Đây là kết quả cân bằng tốt nhất hoặc gần tốt nhất ở nhóm lớp khó.",
            "Kiến trúc: compound scaling/MBConv là cơ chế hợp lý để giải thích hiệu quả phân bổ capacity [4], nhưng campaign không có ablation nên không chứng minh quan hệ nhân quả.",
            "Triển khai: 4,032 triệu tham số và checkpoint 15,67 MiB; thậm chí ít tham số hơn V3-Large nhưng không được suy ra latency thấp hơn nếu chưa benchmark runtime chung.",
        ],
        "mobilenetv3large": [
            "Kết quả: #2, Macro-F1 97,323% ± 0,244%, Accuracy 99,132% ± 0,058%; kém EfficientNet-B0 0,441 điểm % theo mean.",
            "Lớp khó: weeds 95,563% là cao nhất trong sáu model theo mean; downy_mildew 91,909%, tungro 93,071%, blast 94,054%.",
            "Kiến trúc: hardware-aware NAS, NetAdapt, SE và hard-swish được thiết kế cho accuracy–latency trade-off trên mobile [2]. Đây là cơ sở chọn ứng viên sản phẩm, không thay phép đo trên OPPO A54.",
            "Bất định: paired bootstrap không cho thấy V3 chắc chắn hơn EfficientNet ở cấp model tổng thể; V3 hơn Shuffle nhất quán và hơn MobileNetV2 rõ ở 2/3 seed.",
            "Triển khai: 4,226 triệu tham số, checkpoint 16,32 MiB; cần TFLite/ONNX conversion đồng nhất trước benchmark.",
        ],
        "resnet18": [
            "Kết quả: #3, Macro-F1 96,878% ± 0,923%; seed 42 đạt 97,592% nhưng seed 2026 giảm còn 95,835%.",
            "Ổn định: SD cao thứ hai sau ShuffleNet; vì vậy một run ResNet đẹp có thể gây kết luận quá lạc quan nếu không chạy nhiều seed.",
            "Kiến trúc: residual shortcut hỗ trợ tối ưu mạng sâu [3], nhưng ResNet-18 không được thiết kế đặc biệt cho mobile.",
            "Tài nguyên: 11,186 triệu tham số và checkpoint 42,75 MiB, lớn nhất campaign. Dù chất lượng khá cao, tỷ lệ chất lượng–quy mô kém hơn EfficientNet/V3.",
            "Lưu ý: không có paired bootstrap ResNet với EfficientNet hoặc DenseNet; không được nói #3 có ý nghĩa thống kê so với #4.",
        ],
        "densenet121": [
            "Kết quả: #4, Macro-F1 96,658% ± 0,320%, chỉ thấp hơn ResNet-18 0,220 điểm % theo mean.",
            "Lớp: leaf_smut 94,500% là cao nhất theo mean trong sáu model; tungro 93,284% và bacterial_leaf_streak 92,841%.",
            "Ổn định: tốt hơn ResNet-18/MobileNetV2/ShuffleNet theo SD; xu hướng tăng đều từ seed 42 đến 2026 chỉ là mô tả, không phải quan hệ với giá trị seed.",
            "Kiến trúc: dense connectivity giúp feature propagation/reuse [6]; có thể phù hợp tổn thương đa mức, nhưng concatenate feature map có chi phí bộ nhớ chưa đo.",
            "Tài nguyên: 6,973 triệu tham số, checkpoint 27,18 MiB và thời gian train ghi nhận dài nhất; training time không phải inference latency.",
        ],
        "mobilenetv2": [
            "Kết quả: #5, Macro-F1 96,184% ± 0,922%; seed 123 thấp hơn rõ so với hai seed còn lại.",
            "So với V3: V3 tăng mean Macro-F1 1,139 điểm % và observed delta dương ở cả ba seed; đây là bằng chứng hợp lý để nâng backbone mobile từ V2 lên V3.",
            "Kiến trúc: inverted residual và linear bottleneck tạo baseline mobile gọn [1], nhưng trong bài toán 19 lớp này capacity/attention có thể hạn chế phân biệt tổn thương gần nhau.",
            "Lớp khó: downy_mildew 86,516%, tungro 89,403%, weeds 90,811%; là các điểm cần xem khi dùng V2 làm bản Basic.",
            "Tài nguyên: 2,248 triệu tham số, checkpoint 8,81 MiB. Nhỏ hơn V3/EfficientNet nhưng chưa đủ để kết luận trải nghiệm mobile tốt hơn.",
        ],
        "shufflenetv2_1_0x": [
            "Kết quả: #6, Macro-F1 90,971% ± 1,691%; thấp hơn model #5 MobileNetV2 5,213 điểm % và biến thiên lớn nhất.",
            "Lớp khó: bacterial_leaf_streak 71,217%, downy_mildew 74,558%, tungro 75,876%, blast 79,901%, weeds 81,939%.",
            "Diễn giải: kết quả cho thấy cấu hình 1.0× dưới common recipe có thể thiếu capacity/độ ổn định cho lớp hình thái gần; chưa chứng minh mọi biến thể ShuffleNet đều yếu.",
            "Tài nguyên: nhỏ nhất, 1,273 triệu tham số và checkpoint 5,02 MiB; đây là đầu Pareto về quy mô nhưng phải đánh đổi chất lượng đáng kể.",
            "Tốc độ: paper ShuffleNet V2 nhấn mạnh actual speed phụ thuộc memory access/platform [5]; không được gọi đây là model nhanh nhất nếu chưa đo.",
        ],
    }
    order = ["efficientnetb0", "mobilenetv3large", "resnet18", "densenet121", "mobilenetv2", "shufflenetv2_1_0x"]
    for idx, model in enumerate(order, 1):
        heading = add_heading(doc, f"8.{idx} {MODEL_NAMES[model]}", 2)
        if idx == 4:
            heading.paragraph_format.page_break_before = True
        for item in analyses[model]:
            add_paragraph(doc, item, bold_prefix=item.split(":", 1)[0] + ":" if ":" in item else None)
    diagnostic_heading = add_heading(doc, "8.7 Bảng chẩn đoán chéo sáu mô hình", 2)
    diagnostic_heading.paragraph_format.page_break_before = True
    rows = [["Mô hình", "Macro-P / Macro-R", "Lớp F1 thấp nhất", "Lớp bất ổn nhất", "Nhầm lẫn nhiều nhất / 3 seed"]]
    for model in MODEL_ORDER:
        hardest_class = min(CLASS_NAMES, key=lambda cls: per_class[model][cls]["mean"])
        unstable_class = max(CLASS_NAMES, key=lambda cls: per_class[model][cls]["sd"])
        (truth, pred), count = resources[model]["error_pairs"].most_common(1)[0]
        rows.append([
            MODEL_NAMES[model],
            f"{aggregate[model]['macro_precision']['mean']*100:.2f} / {aggregate[model]['macro_recall']['mean']*100:.2f}",
            f"{hardest_class}\n{per_class[model][hardest_class]['mean']*100:.2f}%",
            f"{unstable_class}\nSD {per_class[model][unstable_class]['sd']*100:.2f} đ.%",
            f"{truth} → {pred}: {count}",
        ])
    add_table_custom(doc, rows, [1550, 1700, 2050, 2050, 2010], font_size=7.15,
                     caption_title="table_model_diagnostics", caption_description="Macro precision recall, lớp yếu, bất ổn theo seed và mẫu nhầm riêng của từng mô hình")
    add_caption(doc, "Bảng 8c. Số nhầm là tổng quyết định của ba seed trên cùng test, không phải số ảnh độc lập.")
    sh_bls = per_class["shufflenetv2_1_0x"]["bacterial_leaf_streak"]
    sh_dm = per_class["shufflenetv2_1_0x"]["downy_mildew"]
    v2_dm = per_class["mobilenetv2"]["downy_mildew"]
    add_paragraph(doc, f"Error profile theo lớp cho thấy recall có thể cao nhưng precision thấp: ShuffleNet ở bacterial_leaf_streak đạt precision {sh_bls['precision_mean']*100:.2f}%, recall {sh_bls['recall_mean']*100:.2f}% và F1 {sh_bls['mean']*100:.2f}%; ở downy_mildew là {sh_dm['precision_mean']*100:.2f}% / {sh_dm['recall_mean']*100:.2f}% / {sh_dm['mean']*100:.2f}%. MobileNetV2 ở downy_mildew cũng có precision {v2_dm['precision_mean']*100:.2f}% nhưng recall {v2_dm['recall_mean']*100:.2f}%. Điều này định vị dương tính giả, nhưng chưa xác định nguyên nhân nếu chưa audit ảnh và ablation loss.")


def add_class_errors(doc, aggregate, per_class, support, resources, repeated_errors, unanimous_wrong, figures, bullet_id):
    add_heading(doc, "9. Phân tích theo lớp và mẫu lỗi", 1)
    add_figure(doc, figures["classes"], "Hình 4. Mean F1 của 19 lớp, tách thành hai khối để giữ cỡ chữ; mọi ô đều ghi giá trị và support.", alt="Hai heatmap sáu mô hình theo 19 lớp với giá trị F1 trong mọi ô. Các ô thấp tập trung ở downy mildew, tungro, blast, weeds và bacterial leaf streak, đặc biệt với ShuffleNet V2.")
    hardest = ["downy_mildew", "tungro", "blast", "weeds", "bacterial_leaf_streak"]
    rows = [["Lớp (support)"] + [MODEL_SHORT[m] for m in MODEL_ORDER]]
    for cls in hardest:
        rows.append([f"{cls} (n={support[cls]})"] + [f"{per_class[m][cls]['mean']*100:.2f} ± {per_class[m][cls]['sd']*100:.2f}" for m in MODEL_ORDER])
    add_table_custom(doc, rows, [2050, 1215, 1215, 1215, 1215, 1215, 1235], font_size=7.2,
                     caption_title="table_hard_classes", caption_description="Mean và sample SD F1 của năm lớp khó nhất")
    add_caption(doc, "Bảng 9. Năm lớp khó nhất; mỗi ô là mean ± sample SD F1 (%) qua 3 seed.")
    add_paragraph(doc, "downy_mildew là lớp khó nhất toàn campaign. EfficientNet-B0 đạt mean 92,939%, V3-Large 91,909%, ResNet-18 90,156%, DenseNet-121 89,183%, MobileNetV2 86,516% và ShuffleNet 74,558%. Các khoảng cách này cho thấy năng lực phân biệt tổn thương hình thái gần là nơi kiến trúc tách nhau rõ hơn so với lớp lớn/dễ.")
    add_paragraph(doc, "southern_black_streaked_dwarf (n=12) và whitebacked_planthopper (n=29) đạt 100% ở cả 18 run. Nhưng chỉ cần một ảnh sai sẽ làm recall thay đổi 8,33 hoặc 3,45 điểm %, nên 100% ở support thấp không đồng nghĩa bài toán đã giải quyết.")

    add_heading(doc, "9.1 Mô hình dẫn đầu theo từng lớp", 2)
    rows = [["Lớp", "Support", "Mô hình có mean F1 cao nhất", "Mean F1", "Biên với mức kế tiếp"]]
    unique_wins = Counter()
    for cls in CLASS_NAMES:
        values = {model: per_class[model][cls]["mean"] for model in MODEL_ORDER}
        best_value = max(values.values())
        winners = [model for model, value in values.items() if math.isclose(value, best_value, rel_tol=0, abs_tol=1e-12)]
        if len(winners) == 1:
            unique_wins[winners[0]] += 1
        distinct = sorted({round(value, 14) for value in values.values()}, reverse=True)
        margin = best_value - distinct[1] if len(distinct) > 1 else 0.0
        rows.append([
            cls,
            str(support[cls]),
            ", ".join(MODEL_SHORT[model] for model in winners),
            f"{best_value*100:.3f}%",
            f"{margin*100:.3f} đ.%" if len(winners) == 1 else "Đồng hạng",
        ])
    add_table_custom(doc, rows, [2250, 800, 2550, 1500, 2260], font_size=7.1,
                     caption_title="table_class_leaders", caption_description="Mô hình dẫn mean F1 theo từng lớp, có xử lý đồng hạng và support")
    add_caption(doc, "Bảng 9a. Mean metric được tính qua ba seed; biên không phải CI hay kiểm định theo lớp.")
    add_paragraph(doc, "Số lớp có người thắng duy nhất: " + "; ".join(
        f"{MODEL_NAMES[model]} {unique_wins[model]}" for model in MODEL_ORDER if unique_wins[model]
    ) + ". Bốn lớp còn lại có đồng hạng tuyệt đối, phần lớn do support thấp hoặc tất cả mô hình đều dự đoán đúng toàn bộ; vì vậy không gán điểm thắng tùy tiện cho model xuất hiện trước.")

    add_heading(doc, "9.2 Bất ổn theo lớp qua seed", 2)
    rows = [["Mô hình", "Lớp có SD F1 lớn nhất", "Support", "Mean F1", "Sample SD F1"]]
    for model in MODEL_ORDER:
        cls = max(CLASS_NAMES, key=lambda class_name: per_class[model][class_name]["sd"])
        rows.append([
            MODEL_NAMES[model], cls, str(support[cls]),
            f"{per_class[model][cls]['mean']*100:.3f}%",
            f"{per_class[model][cls]['sd']*100:.3f} đ.%",
        ])
    add_table_custom(doc, rows, [1700, 2920, 850, 1600, 2290], font_size=7.6,
                     caption_title="table_class_instability", caption_description="Lớp có biến thiên F1 lớn nhất trong mỗi mô hình qua ba training seed")
    add_caption(doc, "Bảng 9b. Aggregate ổn định có thể che bất ổn ở một lớp nhỏ; n=3 nên SD chỉ mang tính mô tả.")

    add_heading(doc, "9.3 Mẫu nhầm lẫn", 2)
    add_figure(doc, figures["confusion"], "Hình 5. Tám hướng nhầm có tỷ lệ theo lớp thật cao nhất của EfficientNet-B0 và MobileNetV3-Large; đường chéo đúng đã được loại.", alt="Hai biểu đồ thanh chỉ hiển thị các nhầm lẫn ngoài đường chéo lớn nhất của EfficientNet-B0 và MobileNetV3-Large, trung bình qua ba seed.")
    top_pairs = repeated_errors.most_common(12)
    rows = [["Nhãn thật", "Dự đoán", "Số quyết định sai / 3 seed × 6 model", "Diễn giải"]]
    for (truth, pred), count in top_pairs:
        note = "Ưu tiên audit" if truth in {"blast", "tungro", "weeds", "downy_mildew"} else "Theo dõi"
        rows.append([truth, pred, str(count), note])
    add_table_custom(doc, rows, [2350, 2350, 2150, 2510], font_size=7.8,
                     caption_title="table_error_pairs", caption_description="Các cặp nhầm lẫn có tổng số repeated decisions cao nhất")
    add_caption(doc, "Bảng 10. Tổng lỗi qua 18 run. Đây là số quyết định lặp lại, không phải số ảnh độc lập.")
    add_bullet(doc, "Mẫu phổ biến: blast→tungro; weeds↔normal; tungro↔downy_mildew/blast; blast↔brown_spot/downy_mildew.", bullet_id)
    add_bullet(doc, "Tổng repeated errors trên 3×9.829 quyết định/model: EfficientNet-B0 207; V3-Large 256; ResNet-18 325; DenseNet-121 328; MobileNetV2 407; ShuffleNet V2 1.022.", bullet_id)
    rows = [["Mô hình", "Ba hướng nhầm có tổng quyết định lớn nhất qua 3 seed"]]
    for model in MODEL_ORDER:
        summary = "; ".join(
            f"{truth}→{pred}: {count}"
            for (truth, pred), count in resources[model]["error_pairs"].most_common(3)
        )
        rows.append([MODEL_NAMES[model], summary])
    add_table_custom(doc, rows, [1900, 7460], font_size=7.6,
                     caption_title="table_model_error_signatures", caption_description="Ba hướng nhầm có tổng số quyết định lớn nhất trong từng mô hình qua ba seed")
    add_caption(doc, "Bảng 10a. Error signature riêng từng mô hình; đơn vị là quyết định lặp lại trên cùng tập ảnh.")
    add_heading(doc, "9.4 Ảnh sai nhất quán", 2)
    add_paragraph(doc, f"Có {len(unanimous_wrong)} ảnh bị cả 18 tổ hợp model×seed dự đoán sai. Đây là ứng viên audit chất lượng ảnh/nhãn hoặc trường hợp hình thái đặc biệt; không phải bằng chứng tự động rằng nhãn sai.")
    add_figure(doc, figures["atlas"], "Hình 6. Error atlas của 11 ảnh sai ở toàn bộ 18 run; nhãn thật và phân bố dự đoán được ghi dưới mỗi ảnh.", alt="Lưới 11 ảnh kiểm thử bị cả sáu mô hình ở ba seed dự đoán sai. Phần lớn là ảnh blast bị nhầm sang tungro; hình chỉ dùng cho audit hậu nghiệm, không kết luận nhãn sai.")
    rows = [["Ảnh", "Nhãn thật", "Các dự đoán phổ biến (số run)"]]
    for rel, truth, preds in unanimous_wrong[:12]:
        rows.append([Path(rel).name, truth, "; ".join(f"{p}:{n}" for p, n in preds)])
    add_table_custom(doc, rows, [3100, 2000, 4260], font_size=7.7,
                     caption_title="table_unanimous_errors", caption_description="Các ảnh bị dự đoán sai ở toàn bộ 18 run")
    add_caption(doc, "Bảng 11. Danh sách audit hậu nghiệm; không dùng để sửa test hoặc huấn luyện lại.")


def add_efficiency(doc, aggregate, resources, figures):
    add_heading(doc, "10. Hiệu quả tài nguyên và hàm ý triển khai mobile", 1)
    add_figure(doc, figures["pareto"], "Hình 7. Quan hệ mean Macro-F1 với số tham số; đường đứt thể hiện biên Pareto trên đúng hai trục này.", alt="Biểu đồ phân tán chất lượng theo số tham số. ShuffleNet V2, MobileNetV2 và EfficientNet-B0 nằm trên biên Pareto tham số–Macro-F1; chưa phản ánh latency hay RAM.")
    rows = [["Mô hình", "Tham số", "Checkpoint", "Train tổng 3 seed", "Mean/run", "Test−best val F1"]]
    for model in MODEL_ORDER:
        r = resources[model]
        rows.append([
            MODEL_NAMES[model],
            fmt_int(r["params"]),
            f"{r['checkpoint_mib']:.3f} MiB",
            f"{r['duration_hours_total']:.3f} h",
            f"{r['duration_hours_mean']:.3f} h",
            f"{r['test_gap']*100:+.3f} đ.%",
        ])
    add_table_custom(doc, rows, [1800, 1500, 1450, 1600, 1350, 1660], font_size=7.8,
                     caption_title="table_resources", caption_description="Quy mô checkpoint và thời gian huấn luyện ghi nhận")
    add_caption(doc, "Bảng 12. Thời gian train là số ghi trong manifest, bị ảnh hưởng bởi early stopping và tải máy; không phải inference latency.")
    add_heading(doc, "10.1 Biên Pareto chỉ theo tham số và Macro-F1", 2)
    pareto_models = {"shufflenetv2_1_0x", "mobilenetv2", "efficientnetb0"}
    rows = [["Mô hình", "Tham số", "Mean Macro-F1", "Trạng thái trong mặt phẳng 2D"]]
    for model in sorted(MODEL_NAMES, key=lambda item: resources[item]["params"]):
        rows.append([
            MODEL_NAMES[model], fmt_int(resources[model]["params"]),
            f"{aggregate[model]['macro_f1']['mean']*100:.3f}%",
            "Trên biên Pareto" if model in pareto_models else "Bị mô hình khác chi phối",
        ])
    add_table_custom(doc, rows, [1950, 1550, 1850, 4010], font_size=7.8,
                     caption_title="table_parameter_f1_pareto", caption_description="Trạng thái Pareto chỉ xét số tham số và mean Macro-F1")
    add_caption(doc, "Bảng 12a. Biên Pareto thay đổi nếu thêm latency, RAM, năng lượng hoặc precision; các đại lượng đó chưa được đo.")
    add_paragraph(doc, "Trong không gian hai mục tiêu “ít tham số hơn” và “Macro-F1 cao hơn”, EfficientNet-B0 chiếm ưu thế so với MobileNetV3-Large vì vừa ít tham số hơn vừa có mean F1 cao hơn. Tuy vậy V3 được thiết kế hardware-aware và có thể chạy nhanh hơn trên một số runtime; vì thế không được gọi EfficientNet là lựa chọn mobile tốt hơn trước khi đo latency/RAM/năng lượng.")
    add_paragraph(doc, "EfficientNet-B0 có ít hơn MobileNetV3-Large 194.484 tham số (4,602%) và cao hơn 0,441 điểm % Macro-F1 trong campaign. Đây là dominance theo hai trục đã đo, không phải dominance triển khai. Checkpoint .pt chứa tensor/trạng thái PyTorch và cũng không bằng kích thước TFLite/ONNX/package cuối.")
    add_heading(doc, "10.2 Quyết định Basic và Plus", 2)
    rows = [
        ["Phiên bản", "Backbone đề xuất", "Giải thích", "Điều kiện trước phát hành"],
        ["Basic", "MobileNetV3-Large + Grad-CAM chuẩn", "V3 có chất lượng #2 và định hướng mobile; hiện tốt hơn V2 trong campaign", "Benchmark cùng runtime; UX explainability tối giản; không tuyên bố calibrated confidence"],
        ["Plus", "MobileNetV3-Large + Hybrid CAM", "Giữ cùng classifier để cô lập giá trị của phương pháp giải thích", "Đánh giá localization/user study/runtime riêng; không trộn kết quả CAM với benchmark backbone"],
        ["Ứng viên chất lượng", "EfficientNet-B0", "Mean Macro-F1 #1, ổn định nhất, quy mô gần V3", "Chuyển đổi cùng precision và benchmark A/B trên OPPO A54"],
    ]
    add_table_custom(doc, rows, [1250, 2000, 2850, 3260], font_size=7.7,
                     caption_title="table_product_recommendation", caption_description="Khuyến nghị tách benchmark backbone khỏi đánh giá explainability và mobile")
    add_caption(doc, "Bảng 13. Khuyến nghị sản phẩm hiện tại; quyết định cuối phụ thuộc benchmark mobile chuẩn hóa.")
    add_callout(doc, "Benchmark mobile bắt buộc tiếp theo: cùng OPPO A54, cùng ảnh/input 224×224, batch=1, cùng runtime/delegate, cùng precision/quantization, cùng số thread, cùng warm-up; báo p50/p95/trimmed mean, peak PSS/RSS, năng lượng/nhiệt và kích thước package. Không so model bằng training time hoặc FLOPs đơn lẻ [5], [18].")


def add_validity(doc, bullet_id):
    add_heading(doc, "11. Độ giá trị, giới hạn và nguy cơ sai lệch", 1)
    rows = [
        ["Nguy cơ", "Tác động", "Cách viết / khắc phục"],
        ["Split image-level", "Ảnh tương quan cùng cây/ruộng/phiên có thể ở nhiều split", "Chỉ kết luận nội bộ; xây group-aware split khi có metadata"],
        ["Provenance thu thập/nhãn còn thiếu", "Chưa xác minh đầy đủ địa điểm, thiết bị, thời gian và chuyên gia gán nhãn", "Không gọi kết quả là xác nhận nông học/lâm sàng; bổ sung datasheet và audit chuyên gia"],
        ["Không external test", "Chưa biết domain shift theo vùng, mùa, thiết bị, ánh sáng", "Thu thập field holdout độc lập theo không gian/thời gian"],
        ["n=3 seed", "SD ước lượng thô; không suy luận distribution mạnh", "Báo từng seed + mean±sample SD; tăng seed nếu tài nguyên cho phép"],
        ["Common recipe", "Không phải cấu hình tối ưu riêng từng model", "Kết luận “under common protocol”; có thể làm tuning budget công bằng ở nghiên cứu sau"],
        ["Weight enum V1/V2", "So sánh gồm cả pretrained checkpoint", "Ghi rõ weight ID; không gọi là topology-only ablation"],
        ["Class support cực thấp", "F1/100% dễ biến động bởi 1 ảnh", "Luôn báo support và CI; tăng dữ liệu lớp hiếm"],
        ["Bootstrap iid ảnh", "Không phản ánh clustering theo ruộng/cây", "Dùng cluster/bootstrap hoặc grouped test khi có group ID"],
        ["Không calibration", "Softmax confidence không phải xác suất tin cậy", "Đo ECE/Brier, reliability diagram và temperature scaling trên validation"],
        ["Không mobile benchmark 6 model", "Chưa kết luận latency/RAM/năng lượng", "Benchmark cùng thiết bị/runtime sau conversion"],
    ]
    add_table_custom(doc, rows, [1650, 3080, 4630], font_size=7.5,
                     caption_title="table_validity", caption_description="Các mối đe dọa độ giá trị và phương án khắc phục")
    add_caption(doc, "Bảng 14. Threats to validity bắt buộc đưa vào bài báo.")
    add_heading(doc, "11.1 Những gì campaign chứng minh", 2)
    add_bullet(doc, "Dưới một common recipe, cùng split nội bộ và cùng budget tối đa, EfficientNet-B0 có mean Macro-F1 cao nhất và biến thiên seed thấp nhất.", bullet_id)
    add_bullet(doc, "MobileNetV3-Large tốt hơn MobileNetV2 về mean F1 và có lợi thế quan sát nhất quán ở ba seed; paired CI loại 0 ở hai seed.", bullet_id)
    add_bullet(doc, "ShuffleNet V2 1.0× đánh đổi chất lượng lớn để đạt quy mô nhỏ trong cấu hình đã thử.", bullet_id)
    add_heading(doc, "11.2 Những gì campaign chưa chứng minh", 2)
    add_bullet(doc, "Khả năng tổng quát ngoài thực địa, sang tỉnh/mùa/camera khác.", bullet_id)
    add_bullet(doc, "Model nào nhanh/ít RAM/ít pin nhất trên Android.", bullet_id)
    add_bullet(doc, "Hybrid CAM tốt hơn Grad-CAM về localization hoặc lòng tin người dùng.", bullet_id)
    add_bullet(doc, "Confidence là xác suất đã hiệu chỉnh hoặc phù hợp quyết định nông nghiệp có rủi ro.", bullet_id)
    add_bullet(doc, "EfficientNet-B0 tốt nhất tuyệt đối; ResNet chắc chắn hơn DenseNet; mọi biến thể ShuffleNet đều yếu.", bullet_id)


def add_conclusion(doc, aggregate, bullet_id, number_id):
    add_heading(doc, "12. Kết luận và khuyến nghị nghiên cứu", 1)
    add_paragraph(doc, "Campaign RiceAI V2 đã hoàn tất 18/18 run và mở test theo guard sau khi khóa contract/split/checkpoint. EfficientNet-B0 đạt kết quả tốt nhất trong phạm vi thí nghiệm: Accuracy 99,298% ± 0,027% và Macro-F1 97,764% ± 0,064%. MobileNetV3-Large đứng thứ hai với 97,323% ± 0,244%, đồng thời là ứng viên mobile hợp lý nhờ mục tiêu thiết kế hardware-aware. ResNet-18 và DenseNet-121 tạo nhóm giữa; MobileNetV2 là baseline mobile nhỏ nhưng kém V3; ShuffleNet V2 1.0× nhỏ nhất nhưng chịu giảm chất lượng và độ ổn định rõ.")
    add_paragraph(doc, "Kết luận trung thực nhất không phải “model nào thắng tuyệt đối”, mà là: EfficientNet-B0 tối ưu chất lượng nội bộ; MobileNetV3-Large có hồ sơ cân bằng tốt cho mobile nhưng cần xác nhận runtime; nhóm lớp downy_mildew/tungro/blast/weeds là nút thắt chất lượng; và giá trị ngoài thực địa vẫn phải được kiểm tra bằng grouped/external holdout.")
    add_heading(doc, "12.1 Thứ tự công việc đề xuất", 2)
    steps = [
        "Đóng băng final_test_v2; không dùng test để điều chỉnh model/hyperparameter [13].",
        "Tạo benchmark Android chung cho cả sáu model; nếu conversion khác nhau, ghi rõ runtime/delegate/precision.",
        "Thu thập metadata cây–ruộng–phiên–thiết bị và tạo grouped external split.",
        "Audit 11 ảnh sai nhất quán mà không thay đổi test hiện tại; kết quả audit chỉ dùng thiết kế dataset phiên bản mới.",
        "Đánh giá calibration, abstention/uncertainty và threshold theo chi phí sai lầm nông nghiệp.",
        "Thực hiện thí nghiệm Hybrid CAM vs Grad-CAM độc lập: localization, pointing game/IoU nếu có mask, thời gian sinh heatmap và user study.",
    ]
    for index, step in enumerate(steps, start=1):
        add_number(doc, step, number_id, index=index)


def add_references(doc):
    add_heading(doc, "Tài liệu tham khảo", 1)
    ensure_hyperlink_style(doc)
    add_callout(
        doc,
        "Cách mở nguồn: trong Word desktop dùng Ctrl + nhấp vào URL màu xanh; trong Word mobile/Reading Mode chỉ cần chạm. URL/DOI thật được hiển thị để vẫn có thể sao chép nếu trình xem không hỗ trợ hyperlink DOCX.",
    )
    refs = [
        ("[1]", "M. Sandler et al., “MobileNetV2: Inverted Residuals and Linear Bottlenecks,” CVPR, 2018. doi: 10.1109/CVPR.2018.00474.", "https://openaccess.thecvf.com/content_cvpr_2018/html/Sandler_MobileNetV2_Inverted_Residuals_CVPR_2018_paper.html"),
        ("[2]", "A. Howard et al., “Searching for MobileNetV3,” ICCV, 2019, pp. 1314–1324. doi: 10.1109/ICCV.2019.00140.", "https://openaccess.thecvf.com/content_ICCV_2019/html/Howard_Searching_for_MobileNetV3_ICCV_2019_paper.html"),
        ("[3]", "K. He et al., “Deep Residual Learning for Image Recognition,” CVPR, 2016. doi: 10.1109/CVPR.2016.90.", "https://openaccess.thecvf.com/content_cvpr_2016/html/He_Deep_Residual_Learning_CVPR_2016_paper.html"),
        ("[4]", "M. Tan and Q. V. Le, “EfficientNet: Rethinking Model Scaling for Convolutional Neural Networks,” ICML/PMLR 97, 2019, pp. 6105–6114.", "https://proceedings.mlr.press/v97/tan19a.html"),
        ("[5]", "N. Ma et al., “ShuffleNet V2: Practical Guidelines for Efficient CNN Architecture Design,” ECCV, 2018. doi: 10.1007/978-3-030-01264-9_8.", "https://openaccess.thecvf.com/content_ECCV_2018/html/Ningning_Light-weight_CNN_Architecture_ECCV_2018_paper.html"),
        ("[6]", "G. Huang et al., “Densely Connected Convolutional Networks,” CVPR, 2017. doi: 10.1109/CVPR.2017.243.", "https://openaccess.thecvf.com/content_cvpr_2017/html/Huang_Densely_Connected_Convolutional_CVPR_2017_paper.html"),
        ("[7]", "Y. Cui et al., “Class-Balanced Loss Based on Effective Number of Samples,” CVPR, 2019, pp. 9268–9277. doi: 10.1109/CVPR.2019.00949.", "https://openaccess.thecvf.com/content_CVPR_2019/html/Cui_Class-Balanced_Loss_Based_on_Effective_Number_of_Samples_CVPR_2019_paper.html"),
        ("[8]", "M. Sokolova and G. Lapalme, “A systematic analysis of performance measures for classification tasks,” Information Processing & Management, 45(4), 2009. doi: 10.1016/j.ipm.2009.03.002.", "https://doi.org/10.1016/j.ipm.2009.03.002"),
        ("[9]", "Scikit-learn, “Precision, recall, F-measure and support” and “Confusion matrix,” official documentation.", "https://scikit-learn.org/stable/modules/generated/sklearn.metrics.precision_recall_fscore_support.html"),
        ("[10]", "B. Efron and R. Tibshirani, “Bootstrap Methods for Standard Errors, Confidence Intervals, and Other Measures of Statistical Accuracy,” Statistical Science, 1(1), 1986. doi: 10.1214/ss/1177013815.", "https://doi.org/10.1214/ss/1177013815"),
        ("[11]", "N. Reimers and I. Gurevych, “Reporting Score Distributions Makes a Difference,” EMNLP, 2017. doi: 10.18653/v1/D17-1035.", "https://aclanthology.org/D17-1035/"),
        ("[12]", "X. Bouthillier et al., “Accounting for Variance in Machine Learning Benchmarks,” MLSys, 2021.", "https://proceedings.mlsys.org/paper_files/paper/2021/hash/0184b0cd3cfb185989f858a1d9f5c1eb-Abstract.html"),
        ("[13]", "G. C. Cawley and N. L. C. Talbot, “On Over-fitting in Model Selection and Subsequent Selection Bias in Performance Evaluation,” JMLR, 11, 2010, pp. 2079–2107.", "https://www.jmlr.org/papers/v11/cawley10a.html"),
        ("[14]", "PyTorch 2.2, “Reproducibility,” official documentation.", "https://docs.pytorch.org/docs/2.2/notes/randomness.html"),
        ("[15]", "D. R. Roberts et al., “Cross-validation strategies for data with temporal, spatial, hierarchical, or phylogenetic structure,” Ecography, 40, 2017, pp. 913–929. doi: 10.1111/ecog.02881.", "https://doi.org/10.1111/ecog.02881"),
        ("[16]", "S. Kapoor and A. Narayanan, “Leakage and the reproducibility crisis in machine-learning-based science,” Patterns, 4(9), 2023, 100804. doi: 10.1016/j.patter.2023.100804.", "https://doi.org/10.1016/j.patter.2023.100804"),
        ("[17]", "TorchVision 0.17, “Models and pre-trained weights,” official documentation.", "https://docs.pytorch.org/vision/0.17/models.html"),
        ("[18]", "P. Chao et al., “HarDNet: A Low Memory Traffic Network,” ICCV, 2019.", "https://openaccess.thecvf.com/content_ICCV_2019/html/Chao_HarDNet_A_Low_Memory_Traffic_Network_ICCV_2019_paper.html"),
    ]
    for number, text, url in refs:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.first_line_indent = Inches(-0.28)
        p.paragraph_format.space_after = Pt(5)
        r = p.add_run(number + " ")
        set_run_font(r, size=9, color=DARK_BLUE, bold=True)
        r = p.add_run(text + " ")
        set_run_font(r, size=9, color=INK)
        r.add_break()
        label = "URL chính thức: "
        r = p.add_run(label)
        set_run_font(r, size=8.8, color=MUTED, bold=True)
        add_external_hyperlink(p, url, url, size_pt=8.8)
        if number == "[9]":
            confusion_url = "https://scikit-learn.org/stable/modules/generated/sklearn.metrics.confusion_matrix.html"
            r = p.add_run()
            r.add_break()
            r = p.add_run("URL confusion matrix: ")
            set_run_font(r, size=8.8, color=MUTED, bold=True)
            add_external_hyperlink(p, confusion_url, confusion_url, size_pt=8.8)
    add_heading(doc, "Nguồn nội bộ có thể tái lập", 2)
    internal = [
        "model_benchmark/runs/final_test_v2/final_comparison.json — aggregate qua seed và paired bootstrap.",
        "model_benchmark/runs/final_test_v2/test_open_audit.json — audit mở test.",
        "model_benchmark/runs/final_test_v2/<model>/seed_<seed>/test_metrics.json — metric, per-class, confusion, CI per run.",
        "model_benchmark/runs/final_test_v2/<model>/seed_<seed>/test_predictions.csv — 9.829 dự đoán/run.",
        "model_benchmark/runs/training_v2/<model>/seed_<seed>/run_manifest.json — provenance, validation best, duration, checkpoint SHA.",
        "model_benchmark/config/benchmark_contract_v2.json và campaign_lock_v2.json — protocol/lock.",
        "research/data/splits/rice_cleaned_v4_exact_cross_label_conflicts_excluded_v1.metadata.json — curation/split provenance.",
    ]
    for item in internal:
        add_paragraph(doc, item, size=9.2)


def add_appendices(doc, comparison, aggregate, per_class, support, resources):
    add_heading(doc, "Phụ lục A. Kết quả từng seed", 1)
    rows = [["Mô hình", "Seed", "Accuracy", "Macro-P", "Macro-R", "Macro-F1", "CI95 bootstrap F1"]]
    per_run = {(r["model_id"], r["seed"]): r for r in comparison["per_run"]}
    for model in ["mobilenetv2", "mobilenetv3large", "resnet18", "efficientnetb0", "shufflenetv2_1_0x", "densenet121"]:
        for seed in SEEDS:
            r = per_run[(model, seed)]
            ci = load_json(FINAL_ROOT / model / f"seed_{seed}" / "test_metrics.json")["bootstrap_ci95_macro_f1"]
            rows.append([
                MODEL_NAMES[model], str(seed), pct(r["accuracy"], 3), pct(r["macro_precision"], 3),
                pct(r["macro_recall"], 3), pct(r["macro_f1"], 3), f"[{pct(ci['lower'],3)}; {pct(ci['upper'],3)}]",
            ])
    add_table_custom(doc, rows, [1600, 620, 1180, 1180, 1180, 1180, 2420], font_size=7.1,
                     caption_title="appendix_per_seed", caption_description="Đầy đủ kết quả test và CI bootstrap của 18 run")
    add_caption(doc, "Phụ lục Bảng A1. CI bootstrap là bất định ảnh test trong từng checkpoint, không phải CI của mean ba seed.")

    add_heading(doc, "Phụ lục B. Mean F1 theo 19 lớp", 1)
    rows = [["Canonical class", "n"] + [MODEL_SHORT[m] for m in MODEL_ORDER]]
    for cls in CLASS_NAMES:
        rows.append([cls, str(support[cls])] + [f"{per_class[m][cls]['mean']*100:.2f}" for m in MODEL_ORDER])
    add_table_custom(doc, rows, [2450, 520, 1060, 1060, 1060, 1060, 1060, 1090], font_size=6.8,
                     caption_title="appendix_per_class", caption_description="Mean F1 theo lớp qua ba training seed")
    add_caption(doc, "Phụ lục Bảng B1. Mean F1 (%) qua ba seed; không có paired CI theo lớp.")

    add_heading(doc, "Phụ lục C. Class support và trọng số loss", 1)
    weights = [0.376080, 1.153245, 1.281286, 0.769298, 0.390291, 0.269471, 0.342932, 0.425185, 0.767873, 0.623243, 0.269542, 2.942473, 0.380669, 0.738151, 4.945561, 0.269471, 0.506482, 0.401034, 2.147714]
    rows = [["#", "Canonical class", "Tên diễn giải tiếng Việt", "Test n", "Effective-number weight"]]
    for idx, (cls, weight) in enumerate(zip(CLASS_NAMES, weights)):
        rows.append([str(idx), cls, CLASS_VI[cls], str(support[cls]), f"{weight:.6f}"])
    add_table_custom(doc, rows, [520, 2900, 2800, 900, 2240], font_size=7.5,
                     caption_title="appendix_class_weights", caption_description="Class index, support test và trọng số weighted cross-entropy")
    add_caption(doc, "Phụ lục Bảng C1. Trọng số tính từ train split bằng effective number, β=0,999, chuẩn hóa mean và clip [0,25; 5].")

    add_heading(doc, "Phụ lục D. Checklist tái lập và viết bài", 1)
    checklist = [
        "Giữ nguyên split/contract/campaign lock SHA trong bản thảo và code release.",
        "Báo đầy đủ 18 run; không chỉ chọn seed tốt nhất.",
        "Ghi mean ± sample SD qua seed; không gọi SD là CI95.",
        "Ghi test-image bootstrap CI riêng; không trộn với seed variability.",
        "Ghi rõ paired bootstrap chỉ V3-Large vs năm model, theo từng seed, chưa multiplicity correction.",
        "Luôn báo class support cạnh per-class F1.",
        "Ghi split image-level và chưa có external/grouped validation.",
        "Không đưa Weighted-F1/AUC/ECE/FLOPs/latency nếu chưa chạy phân tích bổ sung.",
        "Tách nghiên cứu backbone khỏi nghiên cứu Grad-CAM/Hybrid CAM.",
        "Không dùng final test để điều chỉnh campaign V2 sau khi đã mở.",
    ]
    for item in checklist:
        p = doc.add_paragraph()
        r = p.add_run("☐ ")
        set_run_font(r, size=10.2, color=BLUE, bold=True)
        r = p.add_run(item)
        set_run_font(r, size=10.2, color=INK)


def build():
    comparison, audit, metrics, manifests = load_campaign()
    aggregate, per_class, confusion_rate, support, resources, repeated_errors, unanimous_wrong = build_derived(comparison, metrics, manifests)
    figures = generate_figures(comparison, aggregate, per_class, confusion_rate, support, resources, unanimous_wrong)

    doc = Document()
    configure_styles(doc)
    configure_section(doc.sections[0])
    add_running_furniture(doc.sections[0])
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.styles["Normal"].paragraph_format.space_after = Pt(5.5)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.22
    bullet_id = add_bullet_numbering_definition(doc)
    protocol_number_id = add_numbering_definition(doc, "%1.", is_decimal=True)
    conclusion_number_id = add_numbering_definition(doc, "%1.", is_decimal=True)
    normalize_numbering_order(doc)

    add_cover(doc, audit)
    add_static_toc(doc)
    add_exec_summary(doc, aggregate, resources, support, bullet_id)
    add_scope_and_data(doc, support, bullet_id)
    add_architectures(doc, aggregate, per_class, resources)
    add_protocol(doc, audit, manifests, resources, bullet_id, protocol_number_id)
    add_test_metrics(doc, audit, bullet_id)
    add_overall_results(doc, aggregate, figures, resources)
    add_bootstrap_results(doc, comparison, figures)
    add_model_analysis(doc, aggregate, per_class, resources)
    add_class_errors(doc, aggregate, per_class, support, resources, repeated_errors, unanimous_wrong, figures, bullet_id)
    add_efficiency(doc, aggregate, resources, figures)
    add_validity(doc, bullet_id)
    add_conclusion(doc, aggregate, bullet_id, conclusion_number_id)
    add_references(doc)
    add_appendices(doc, comparison, aggregate, per_class, support, resources)

    props = doc.core_properties
    props.title = "Báo cáo phân tích kết quả so sánh 6 mô hình CNN cho RiceAI"
    props.subject = "Kết quả final test khóa V2: 6 backbone × 3 seed"
    props.author = "RiceAI Research"
    props.keywords = "rice disease, CNN, MobileNet, EfficientNet, ResNet, ShuffleNet, DenseNet, benchmark, Macro-F1"
    props.comments = "Generated only from locked campaign V2 artifacts; test results are not used for tuning."

    settings = doc.settings._element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)
    if settings.find(qn("w:doNotCompressPictures")) is None:
        settings.append(OxmlElement("w:doNotCompressPictures"))
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print("DOCX_CREATED")


if __name__ == "__main__":
    build()
