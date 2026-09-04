"""Create a concise supervisor review from frozen V2 results; no training/evaluation."""
from __future__ import annotations

import ast
import hashlib
import json
import statistics as stats
import zipfile
from pathlib import Path

import numpy as np
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / 'docs'
SOURCE = DOCS / 'BAO_CAO_PHAN_TICH_KET_QUA_SO_SANH_6_MO_HINH.docx'
OUTPUT = DOCS / 'BAO_CAO_SO_BO_SO_SANH_6_MO_HINH.docx'
QA = DOCS / '_qa_preliminary_report'
ASSETS = DOCS / '_results_report_assets'
FINAL = ROOT / 'runs' / 'final_test_v2'
TRAIN = ROOT / 'runs' / 'training_v2'
SEEDS = [42, 123, 2026]
BLUE, INK, MUTED, GRID = '2E74B5', '17324D', '596675', 'BCC8D5'

# Resolve stable labels from the existing report without importing its plotting dependencies.
tree = ast.parse((ROOT / 'tools/build_final_results_report_docx.py').read_text(encoding='utf-8'))
LABELS = {}
for node in tree.body:
    if isinstance(node, ast.Assign):
        for target in node.targets:
            if isinstance(target, ast.Name) and target.id in {
                'MODEL_ORDER', 'MODEL_NAMES', 'MODEL_SHORT', 'PARAMS',
                'CLASS_NAMES', 'CLASS_VI', 'CLASS_ABBR',
            }:
                LABELS[target.id] = ast.literal_eval(node.value)
globals().update(LABELS)


def load(path):
    return json.loads(path.read_text(encoding='utf-8'))


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def pc(x, n=3):
    return f'{100*x:.{n}f}'


def signed(x):
    return f'{100*x:+.3f}'


def elm(tag, **attrs):
    e = OxmlElement('w:' + tag)
    for k, v in attrs.items():
        e.set(qn('w:' + k), str(v))
    return e


def font(run, size=None, bold=None, color=None):
    run.font.name = 'Calibri'
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    rf = run._element.get_or_add_rPr().get_or_add_rFonts()
    for k in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        rf.set(qn('w:' + k), 'Calibri')


def set_style(doc, name, size, before, after, line=1.1, color=INK, bold=False):
    s = doc.styles[name] if name in doc.styles else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    s.font.name = 'Calibri'
    s.font.size = Pt(size)
    s.font.color.rgb = RGBColor.from_string(color)
    s.font.bold = bold
    p = s.paragraph_format
    p.space_before, p.space_after = Pt(before), Pt(after)
    p.line_spacing = line
    p.widow_control = True
    return s


def setup():
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Inches(8.5), Inches(11)
    s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Inches(1)
    s.header_distance = s.footer_distance = Inches(.492)
    set_style(doc, 'Normal', 11, 0, 6)
    set_style(doc, 'Title', 23, 0, 4, color=INK, bold=True)
    set_style(doc, 'Subtitle', 13, 0, 8, color=MUTED)
    for i, size, before, after in [(1,16,16,8),(2,13,12,6),(3,12,8,4)]:
        st = set_style(doc, f'Heading {i}', size, before, after, color=BLUE if i<3 else '1F4D78', bold=True)
        st.paragraph_format.keep_with_next = True
    set_style(doc, 'BriefTable', 8.5, 0, 0, 1.05)
    set_style(doc, 'Caption', 9, 4, 4, 1.05, MUTED)
    set_style(doc, 'BriefNote', 9.5, 4, 6, 1.1, MUTED)
    set_style(doc, 'BriefSource', 8, 0, 3, 1.0)
    set_style(doc, 'BriefFlow', 10, 3, 3, 1.05, BLUE, True)
    for nm in ('Header','Footer'):
        set_style(doc, nm, 8, 0, 0, 1.0, MUTED)
    s.header.paragraphs[0].text = 'RiceAI | Báo cáo sơ bộ • So sánh 6 mô hình • Protocol V2'
    p = s.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run('Bản gửi giảng viên xem trước | Trang ')
    for tag, attrs in [('fldChar',{'fldCharType':'begin'}),('instrText',{}),('fldChar',{'fldCharType':'end'})]:
        e = elm(tag, **attrs)
        if tag == 'instrText':
            e.text = ' PAGE '
        p.add_run()._r.append(e)
    settings = doc.settings.element
    settings.append(elm('doNotCompressPictures'))
    settings.append(elm('updateFields', val='true'))
    defaults = doc.styles.element.find(qn('w:docDefaults'))
    if defaults is not None:
        rpr = defaults.find('.//' + qn('w:rPr'))
        if rpr is not None:
            rpr.append(elm('lang', val='vi-VN'))
    doc.core_properties.title = 'Báo cáo sơ bộ so sánh 6 mô hình nhận diện đối tượng trên lúa'
    doc.core_properties.subject = 'Tóm lược kết quả cuối campaign V2 để giảng viên xem trước'
    doc.core_properties.author = 'RiceAI Research'
    return doc


def para(doc, text, style=None):
    return doc.add_paragraph(text, style)


def heading(doc, text, new=False):
    p = doc.add_paragraph(text, 'Heading 1')
    if new:
        p.paragraph_format.page_break_before = True
        p.paragraph_format.space_before = Pt(0)  # Named override: page-leading heading.
    return p


def table(doc, rows, widths, title, size=8.5):
    assert sum(widths) == 9360
    t = doc.add_table(rows=0, cols=len(widths))
    t.autofit = False
    pr = t._tbl.tblPr
    pr.find(qn('w:tblW')).set(qn('w:w'),'9360')
    pr.find(qn('w:tblW')).set(qn('w:type'),'dxa')
    pr.append(elm('tblInd', w=120, type='dxa'))
    pr.append(elm('tblCaption', val=title))
    pr.append(elm('tblDescription', val=title + '; số liệu từ kết quả V2 đã khóa'))
    margins = elm('tblCellMar')
    for side, val in [('top',80),('bottom',80),('start',120),('end',120)]:
        margins.append(elm(side,w=val,type='dxa'))
    pr.append(margins)
    borders = elm('tblBorders')
    for side in ('top','left','bottom','right','insideH','insideV'):
        borders.append(elm(side,val='single',sz=4,color=GRID))
    pr.append(borders)
    for child in list(t._tbl.tblGrid):
        t._tbl.tblGrid.remove(child)
    for w in widths:
        t._tbl.tblGrid.append(elm('gridCol',w=w))
    for idx, values in enumerate(rows):
        row = t.add_row()
        row._tr.get_or_add_trPr().append(elm('cantSplit'))
        if idx == 0:
            row._tr.get_or_add_trPr().append(elm('tblHeader'))
        for j,(c,v,w) in enumerate(zip(row.cells,values,widths)):
            c.width = Inches(w/1440)
            tcw = c._tc.get_or_add_tcPr().find(qn('w:tcW'))
            tcw.set(qn('w:w'),str(w)); tcw.set(qn('w:type'),'dxa')
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if idx == 0:
                c._tc.get_or_add_tcPr().append(elm('shd',fill='F2F4F7',val='clear'))
            p = c.paragraphs[0]; p.style = doc.styles['BriefTable']
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
            font(p.add_run(str(v)),size=size,bold=idx==0)
    return t


def figure(doc, name, caption, width=6.35):
    path = QA / name if name == 'brief_seed_stability.png' else ASSETS / name
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sh = p.add_run().add_picture(str(path),width=Inches(width))
    sh._inline.docPr.set('descr',caption)
    sh._inline.docPr.set('title',caption.split('.')[0])
    para(doc,caption,'Caption')


def link(p, label, url):
    node = elm('hyperlink')
    node.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    node.set(qn('w:tooltip'),url)
    r = elm('r'); rp = elm('rPr')
    rp.append(elm('color',val='0563C1')); rp.append(elm('u',val='single'))
    rp.append(elm('sz',val=16)); r.append(rp)
    tx=elm('t'); tx.text=label; r.append(tx); node.append(r); p._p.append(node)


def draw_seed_chart(metrics):
    """Redraw the scientific plot from numeric results, with no stale table cross-reference."""
    width, height = 2667, 1995
    image = Image.new('RGB', (width, height), 'white')
    d = ImageDraw.Draw(image)
    regular = 'C:/Windows/Fonts/arial.ttf'
    bold = 'C:/Windows/Fonts/arialbd.ttf'
    f = ImageFont.truetype(regular, 49)
    title = ImageFont.truetype(bold, 61)
    legend = ImageFont.truetype(regular, 44)
    left, right, top, bottom = 235, 2560, 155, 1485
    ymin, ymax = 88, 98.5
    xs = [left+100, (left+right)//2, right-100]
    y = lambda value: bottom - (value-ymin)/(ymax-ymin)*(bottom-top)
    d.text((width/2, 54), 'Độ ổn định Macro-F1 qua ba seed',font=title,fill='#202020',anchor='mm')
    for tick in range(88,99,2):
        yy=y(tick)
        d.line((left,yy,right,yy),fill='#E4E8EC',width=3)
        d.text((left-35,yy),str(tick),font=f,fill='#202020',anchor='rm')
    for xx,seed in zip(xs,SEEDS):
        d.line((xx,top,xx,bottom),fill='#EEF0F2',width=2)
        d.text((xx,bottom+67),str(seed),font=f,fill='#202020',anchor='mm')
    d.rectangle((left,top,right,bottom),outline='#454545',width=3)
    d.text(((left+right)/2,bottom+147),'Training seed',font=f,fill='#202020',anchor='mm')
    label=Image.new('RGBA',(520,90),(255,255,255,0))
    ImageDraw.Draw(label).text((260,45),'Macro-F1 (%)',font=f,fill='#202020',anchor='mm')
    label=label.rotate(90,expand=True)
    image.paste(label,(22,int((top+bottom-label.height)/2)),label)
    colors=['#2E7D32','#2E74B5','#6A5ACD','#00838F','#F28E2B','#B0413E']
    def marker(xv,yv,i,color):
        r=18
        if i==0:
            d.ellipse((xv-r,yv-r,xv+r,yv+r),fill=color)
        elif i==1:
            d.rectangle((xv-r,yv-r,xv+r,yv+r),fill=color)
        elif i==2:
            d.polygon([(xv,yv-r-4),(xv-r-4,yv+r),(xv+r+4,yv+r)],fill=color)
        elif i==3:
            d.polygon([(xv,yv-r-4),(xv-r-4,yv),(xv,yv+r+4),(xv+r+4,yv)],fill=color)
        elif i==4:
            d.line((xv-r,yv,xv+r,yv),fill=color,width=12)
            d.line((xv,yv-r,xv,yv+r),fill=color,width=12)
        else:
            d.line((xv-r,yv-r,xv+r,yv+r),fill=color,width=12)
            d.line((xv-r,yv+r,xv+r,yv-r),fill=color,width=12)
    for i,(model,color) in enumerate(zip(MODEL_ORDER,colors)):
        points=[(xv,y(100*metrics[model][seed]['macro_f1'])) for xv,seed in zip(xs,SEEDS)]
        d.line(points,fill=color,width=8)
        for xv,yv in points:
            marker(xv,yv,i,color)
        col=i%2; row=i//2
        lx=320+col*1240; ly=1760+row*84
        d.line((lx,ly,lx+115,ly),fill=color,width=8)
        marker(lx+57,ly,i,color)
        d.text((lx+145,ly),MODEL_NAMES[model],font=legend,fill='#202020',anchor='lm')
    image.save(QA/'brief_seed_stability.png',dpi=(420,420))


def verify_data(comparison, audit, metrics):
    assert audit['status']=='complete'
    assert comparison['test_sample_count']==9829 and len(comparison['per_run'])==18
    for model in MODEL_ORDER:
        for seed in SEEDS:
            rec=metrics[model][seed]
            cm=np.asarray(rec['confusion_matrix'],dtype=float)
            assert cm.shape==(19,19) and cm.sum()==9829
            tp=cm.diagonal(); pred=cm.sum(axis=0); actual=cm.sum(axis=1)
            precision=np.divide(tp,pred,out=np.zeros_like(tp),where=pred!=0)
            recall=tp/actual
            f1=np.divide(2*precision*recall,precision+recall,out=np.zeros_like(tp),where=(precision+recall)!=0)
            vals={'accuracy':tp.sum()/cm.sum(),'macro_precision':precision.mean(),'macro_recall':recall.mean(),'macro_f1':f1.mean()}
            for k,v in vals.items():
                assert abs(v-rec[k])<1e-12,(model,seed,k)
            for i,r in enumerate(rec['per_class']):
                assert r['class_name']==CLASS_NAMES[i] and r['support']==int(actual[i])
                assert abs(r['f1']-f1[i])<1e-12
        for metric in ('accuracy','macro_precision','macro_recall','macro_f1'):
            vals=[metrics[model][s][metric] for s in SEEDS]
            a=comparison['aggregate_across_training_seeds'][model][metric]
            assert abs(a['mean']-stats.mean(vals))<1e-12
            assert abs(a['sample_std']-stats.stdev(vals))<1e-12


def build():
    QA.mkdir(parents=True,exist_ok=True)
    source_sha=sha(SOURCE)
    comparison=load(FINAL/'final_comparison.json'); audit=load(FINAL/'test_open_audit.json')
    aggregate=comparison['aggregate_across_training_seeds']
    metrics={m:{s:load(FINAL/m/f'seed_{s}'/'test_metrics.json') for s in SEEDS} for m in MODEL_ORDER}
    manifests={m:{s:load(TRAIN/m/f'seed_{s}'/'run_manifest.json') for s in SEEDS} for m in MODEL_ORDER}
    verify_data(comparison,audit,metrics)
    draw_seed_chart(metrics)
    support={r['class_name']:r['support'] for r in metrics['efficientnetb0'][42]['per_class']}
    perclass={m:{c:stats.mean(next(r['f1'] for r in metrics[m][s]['per_class'] if r['class_name']==c) for s in SEEDS) for c in CLASS_NAMES} for m in MODEL_ORDER}
    doc=setup()

    # 1. Lead summary: compact, no ornamental cover page.
    para(doc,'BÁO CÁO SƠ BỘ','Title')
    para(doc,'So sánh 6 mô hình CNN trong nhận diện bệnh và đối tượng trên lúa','Subtitle')
    para(doc,'Gửi giảng viên xem trước | Tổng hợp ngày 03/09/2026','BriefNote')
    para(doc,'Bản rút gọn từ báo cáo phân tích chi tiết; sử dụng kết quả test cuối đã hoàn tất, không phải kết quả tạm trong khi huấn luyện. Đã hoàn thành 6 mô hình × 3 seed = 18 lần huấn luyện và đánh giá trên cùng 9.829 ảnh test.')
    heading(doc,'1. Kết quả chính')
    rows=[['Mô hình','Accuracy\n(%)','Macro-P\n(%)','Macro-R\n(%)','Macro-F1\n(%)']]
    for m in MODEL_ORDER:
        rows.append([MODEL_NAMES[m]]+[f"{pc(aggregate[m][k]['mean'])}\n± {pc(aggregate[m][k]['sample_std'])}" for k in ('accuracy','macro_precision','macro_recall','macro_f1')])
    table(doc,rows,[2240,1780,1780,1780,1780],'Bảng 1. Kết quả tổng thể qua ba seed',8.7)
    para(doc,'Bảng 1. Mean ± sample SD qua seed 42, 123, 2026; xếp theo mean Macro-F1 giảm dần. SD không phải khoảng tin cậy 95%.','Caption')
    para(doc,'EfficientNet-B0 đứng đầu về mean Macro-F1 (97.764%) và ổn định nhất qua ba seed. MobileNetV3-Large đứng thứ hai (97.323%), cao hơn MobileNetV2 1.139 điểm phần trăm. ShuffleNet V2 1.0× có ít tham số nhất nhưng chất lượng thấp nhất trong cấu hình đã thử.')
    heading(doc,'Phạm vi kết luận')
    para(doc,'Đây là so sánh phân loại ảnh đơn nhãn 19 lớp, bao gồm bệnh, sinh vật gây hại, cỏ dại và trạng thái bình thường; không phải định vị vùng bệnh. Kết quả chỉ áp dụng cho tập dữ liệu nội bộ và quy trình chung V2. Chưa có kiểm thử độc lập ngoài thực địa hoặc đo tốc độ, RAM, năng lượng trên điện thoại.')
    para(doc,'Tiêu chí chính: Macro-F1 — tính F1 riêng từng lớp rồi lấy trung bình đều. Accuracy phản ánh tỷ lệ ảnh dự đoán đúng; Macro-P/Macro-R bổ sung góc nhìn precision/recall cân bằng giữa các lớp [7].','BriefNote')

    # 2. Methods and architecture comparison.
    heading(doc,'2. Thiết kế so sánh và quy trình kiểm thử',True)
    para(doc,'65.516 ảnh → Train 45.858 | Validation 9.829 | Test 9.829. Chia phân tầng theo nhãn ở cấp ảnh (70/15/15); manifest được khóa và dùng chung cho tất cả mô hình.')
    para(doc,'Khóa dữ liệu/quy chuẩn → Huấn luyện 6 mô hình × 3 seed','BriefFlow')
    para(doc,'→ Chọn checkpoint bằng validation Macro-F1 → Mở test cuối','BriefFlow')
    para(doc,'→ Tổng hợp mean ± SD, F1 theo lớp, confusion và bootstrap CI','BriefFlow')
    para(doc,'Sơ đồ 1. Test không tham gia chọn checkpoint hoặc tinh chỉnh siêu tham số.','Caption')
    rows=[['Mô hình','Đặc trưng kiến trúc','Trọng số ImageNet']]
    features={
        'mobilenetv2':'Inverted residual; linear bottleneck [1]',
        'mobilenetv3large':'Thiết kế hướng phần cứng; SE; h-swish [2]',
        'resnet18':'Kết nối cộng residual; 18 lớp [3]',
        'efficientnetb0':'MBConv, SE; mô hình cơ sở họ EfficientNet [4]',
        'shufflenetv2_1_0x':'Channel split/shuffle; biến thể 1.0× [5]',
        'densenet121':'Nối đặc trưng giữa các lớp; 121 lớp [6]',
    }
    for m in MODEL_ORDER:
        wid=manifests[m][42]['provenance']['weights_id'].split('.')[-1]
        rows.append([MODEL_NAMES[m],features[m],wid])
    table(doc,rows,[2260,4800,2300],'Bảng 2. Kiến trúc và trọng số tiền huấn luyện',9)
    para(doc,'Bảng 2. Cùng TorchVision nhưng V1/V2 là các checkpoint tiền huấn luyện khác nhau; đây không phải thí nghiệm chỉ thay topology.','Caption')
    heading(doc,'Cấu hình dùng chung')
    para(doc,'Ảnh RGB 224×224, chuẩn hóa ImageNet; batch 32; Adam; dropout 0,4; head tối đa 15 epoch (LR 10⁻³), fine-tune tối đa 30 epoch (LR 10⁻⁵); early stopping patience 5. Dùng trọng số lớp effective number (β=0,999), cùng augmentation và đóng băng BatchNorm ở cả hai giai đoạn [9, 11].')
    para(doc,'Chỉ validation quyết định checkpoint tốt nhất. Các run có thể dừng ở số epoch khác nhau theo cùng quy tắc early stopping. Quy trình V2 và test đã khóa không bị thay đổi khi lập bản tóm tắt này.','BriefNote')

    # 3. Overall chart, kept at the original readable final size.
    heading(doc,'3. Biểu đồ kết quả tổng thể',True)
    figure(doc,'fig_overall_metrics.png','Hình 1. Mean Macro-F1 và Accuracy; thanh sai số là sample SD qua ba seed.')
    para(doc,'Macro-F1 được ưu tiên vì số ảnh giữa 19 lớp rất chênh lệch. Accuracy cao không thay thế được đánh giá các lớp ít ảnh. Biểu đồ dùng trục rút gọn để thể hiện khác biệt, không dùng độ dài thanh để suy ra tỷ lệ hơn/kém.','BriefNote')

    # 4. All 18 seed results in a six-row matrix.
    heading(doc,'4. Độ ổn định qua ba seed',True)
    figure(doc,'brief_seed_stability.png','Hình 2. Macro-F1 từng seed; nối điểm chỉ để dễ đối chiếu, seed không phải thời gian. Giá trị chính xác ở Bảng 3.')
    rows=[['Mô hình','Seed 42','Seed 123','Seed 2026','SD\n(điểm %)']]
    for m in MODEL_ORDER:
        rows.append([MODEL_NAMES[m]]+[pc(metrics[m][s]['macro_f1']) for s in SEEDS]+[pc(aggregate[m]['macro_f1']['sample_std'])])
    table(doc,rows,[2640,1680,1680,1680,1680],'Bảng 3. Macro-F1 của 18 run',8.8)
    para(doc,'Bảng 3. Giá trị Macro-F1 tính theo %. SD được tính với mẫu số n−1, n=3.','Caption')
    para(doc,'EfficientNet-B0 đứng đầu ở cả ba seed; ShuffleNet V2 đứng cuối cả ba. Ba seed chỉ phản ánh một phần biến thiên do huấn luyện, không đại diện mọi biến thiên dữ liệu hoặc triển khai [10].','BriefNote')

    # 5. Per-class visualization.
    heading(doc,'5. Chất lượng theo 19 lớp',True)
    figure(doc,'fig_per_class_heatmap.png','Hình 3. Mean F1 (%) qua ba seed; n là số ảnh test của lớp. Bảng 4 giải nghĩa mã lớp.')
    para(doc,'Các lớp khó nổi bật gồm sương mai (DM), tungro, đạo ôn (Blast), cỏ dại (Weeds) và sọc lá vi khuẩn (BLS). F1 100% ở lớp rất ít ảnh không có nghĩa mô hình nhận diện hoàn hảo ngoài thực địa.','BriefNote')

    # 6. Complete per-class comparison, concise and interpretable.
    heading(doc,'6. Bảng đối chiếu F1 theo lớp',True)
    rows=[['Mã / lớp (diễn giải)','n','Eff-B0','MV3-L','RN18','DN121','MV2','Shuffle']]
    for c,abbr in zip(CLASS_NAMES,CLASS_ABBR):
        rows.append([f'{abbr} — {CLASS_VI[c]}',str(support[c])]+[pc(perclass[m][c],2) for m in MODEL_ORDER])
    table(doc,rows,[2500,740,1020,1020,1020,1020,1020,1020],'Bảng 4. Mean F1 toàn bộ 19 lớp',8)
    para(doc,'Bảng 4. Mean F1 (%), không phải mean precision/recall. n cộng lại bằng 9.829; kết quả cùng lớp được lấy trung bình qua 3 seed.','Caption')
    para(doc,'Eff-B0: EfficientNet-B0; MV3-L: MobileNetV3-Large; RN18: ResNet-18; DN121: DenseNet-121; MV2: MobileNetV2; Shuffle: ShuffleNet V2 1.0×.','BriefNote')
    para(doc,'Hai lớp SBSD (n=12) và WBPH (n=29) đạt F1 100% ở cả sáu mô hình, nhưng support thấp làm bằng chứng còn hạn chế. So sánh theo lớp cần đọc cùng n và bảng tổng thể, không chỉ chọn các lớp có kết quả đẹp.','BriefNote')

    # 7. Confusion chart and short model-level table.
    heading(doc,'7. Nhầm lẫn cần chú ý',True)
    figure(doc,'fig_confusion_best_mobile.png','Hình 4. Tám hướng nhầm có tỷ lệ lớn nhất của EfficientNet-B0 và MobileNetV3-Large; đã bỏ dự đoán đúng.',width=5.8)
    rows=[['Mô hình','Lớp có mean F1 thấp nhất','F1 (%)']]
    for m in MODEL_ORDER:
        c=min(CLASS_NAMES,key=lambda c:perclass[m][c])
        rows.append([MODEL_NAMES[m],CLASS_VI[c],pc(perclass[m][c],2)])
    table(doc,rows,[3000,4500,1860],'Bảng 5. Lớp yếu nhất của từng mô hình',8.3)
    para(doc,'Bảng 5. Lớp có mean F1 thấp nhất trong 19 lớp của từng mô hình.','Caption')
    para(doc,'Tỷ lệ nhầm trong hình lấy trên số ảnh thuộc lớp thật, rồi trung bình qua seed; không phải tỷ lệ lỗi trên toàn bộ test. Chỉ diễn giải lỗi quan sát, chưa kết luận nguyên nhân hoặc nhãn dữ liệu sai.','BriefNote')

    # 8. Paired bootstrap, condensed but preserve all comparisons and seeds.
    heading(doc,'8. So sánh chênh lệch và khoảng tin cậy',True)
    figure(doc,'fig_paired_bootstrap.png','Hình 5. Δ = Macro-F1(V3-Large) − Macro-F1(đối chiếu), đơn vị điểm %. CI95 theo từng seed.')
    rows=[['Đối chiếu với V3-Large','Seed 42','Seed 123','Seed 2026']]
    for m in ['mobilenetv2','resnet18','efficientnetb0','shufflenetv2_1_0x','densenet121']:
        recs=comparison['paired_bootstrap_vs_mobilenetv3large']['mobilenetv3large_minus_'+m]
        row=[MODEL_NAMES[m]]
        for s in SEEDS:
            r=recs[str(s)]
            row.append(f"{signed(r['observed_delta'])}\n[{signed(r['ci95_lower'])}; {signed(r['ci95_upper'])}]")
        rows.append(row)
    table(doc,rows,[2400,2320,2320,2320],'Bảng 6. Chênh lệch Macro-F1 và CI95 của 15 đối chiếu',8)
    para(doc,'Bảng 6. Dòng đầu: Δ quan sát; ngoặc vuông: CI95. Bootstrap ghép cặp trên cùng ảnh test, 2.000 lần lấy mẫu/đối chiếu/seed [8].','Caption')
    para(doc,'CI hoàn toàn dương ủng hộ V3 trong checkpoint đó; hoàn toàn âm ủng hộ đối chiếu; chứa 0 chưa tách biệt rõ. CI này không phải CI của mean ba seed. Chưa hiệu chỉnh 15 so sánh đồng thời nên không tuyên bố ưu thế thống kê tổng thể.','BriefNote')

    # 9. Resource footprint is not runtime performance.
    heading(doc,'9. Quy mô mô hình và chi phí huấn luyện',True)
    figure(doc,'fig_quality_params.png','Hình 6. Chất lượng theo số tham số; đường đứt là biên Pareto chỉ trên hai trục này.')
    rows=[['Mô hình','Tham số\n(triệu)','Checkpoint\n(MiB)','Train tổng\n3 seed (giờ)','Mean/run\n(giờ)']]
    for m in MODEL_ORDER:
        total=sum(manifests[m][s]['duration_seconds'] for s in SEEDS)/3600
        size=stats.mean((TRAIN/m/f'seed_{s}'/'best_model.pt').stat().st_size/1024**2 for s in SEEDS)
        rows.append([MODEL_NAMES[m],f'{PARAMS[m]/1e6:.3f}',f'{size:.3f}',f'{total:.3f}',f'{total/3:.3f}'])
    table(doc,rows,[2640,1440,1800,1800,1680],'Bảng 7. Quy mô và thời gian huấn luyện ghi trong manifest',8.5)
    para(doc,'Bảng 7. Checkpoint .pt trung bình qua 3 seed; thời gian từ run_manifest, chịu ảnh hưởng tải máy và early stopping.','Caption')
    para(doc,'Biên Pareto tham số–Macro-F1 gồm ShuffleNet V2, MobileNetV2 và EfficientNet-B0. EfficientNet-B0 ít hơn V3-Large 194.484 tham số và cao hơn 0.441 điểm % mean F1. Tuy nhiên chưa thể suy ra model nào nhanh hơn hoặc tiết kiệm RAM/pin hơn trên Android.','BriefNote')

    # 10. Short handoff conclusion and a compact real hyperlink bibliography.
    heading(doc,'10. Kết luận sơ bộ và bước tiếp theo',True)
    para(doc,'Trong cùng quy trình V2, EfficientNet-B0 dẫn đầu chất lượng và độ ổn định; MobileNetV3-Large đứng thứ hai và là ứng viên triển khai cần đo runtime thực tế. ResNet-18/DenseNet-121 ở nhóm giữa; ShuffleNet V2 nhỏ nhất nhưng giảm chất lượng rõ trong cấu hình đã thử.')
    para(doc,'Đề xuất thảo luận với giảng viên: giữ nguyên kết quả test đã khóa; bổ sung kiểm thử ngoài thực địa theo cây/ruộng/phiên chụp; đo latency, RAM và năng lượng trên cùng thiết bị/runtime. Nghiên cứu Hybrid CAM so với Grad-CAM cần thí nghiệm riêng, không suy ra từ benchmark classifier này.','BriefNote')
    para(doc,'Giới hạn bắt buộc: split cấp ảnh, chưa có external test; chỉ 3 seed; cùng recipe không đồng nghĩa tối ưu riêng từng kiến trúc; có khác biệt trọng số ImageNet V1/V2. Báo cáo sơ bộ không khẳng định mô hình tốt nhất tuyệt đối.','BriefNote')
    heading(doc,'Nguồn và tài liệu chính')
    refs=[
        ('[1] Sandler et al. (2018). MobileNetV2. CVPR.','https://openaccess.thecvf.com/content_cvpr_2018/html/Sandler_MobileNetV2_Inverted_Residuals_CVPR_2018_paper.html'),
        ('[2] Howard et al. (2019). Searching for MobileNetV3. ICCV.','https://openaccess.thecvf.com/content_ICCV_2019/html/Howard_Searching_for_MobileNetV3_ICCV_2019_paper.html'),
        ('[3] He et al. (2016). Deep Residual Learning for Image Recognition. CVPR.','https://openaccess.thecvf.com/content_cvpr_2016/html/He_Deep_Residual_Learning_CVPR_2016_paper.html'),
        ('[4] Tan & Le (2019). EfficientNet. ICML.','https://proceedings.mlr.press/v97/tan19a.html'),
        ('[5] Ma et al. (2018). ShuffleNet V2. ECCV.','https://openaccess.thecvf.com/content_ECCV_2018/html/Ningning_Light-weight_CNN_Architecture_ECCV_2018_paper.html'),
        ('[6] Huang et al. (2017). Densely Connected Convolutional Networks. CVPR.','https://openaccess.thecvf.com/content_cvpr_2017/html/Huang_Densely_Connected_Convolutional_CVPR_2017_paper.html'),
        ('[7] Scikit-learn. Precision, recall, F-measure and support. Tài liệu chính thức.','https://scikit-learn.org/stable/modules/generated/sklearn.metrics.precision_recall_fscore_support.html'),
        ('[8] Efron & Tibshirani (1986). Bootstrap Methods for Standard Errors, Confidence Intervals, and Other Measures of Statistical Accuracy.','https://doi.org/10.1214/ss/1177013815'),
        ('[9] Cui et al. (2019). Class-Balanced Loss Based on Effective Number of Samples. CVPR.','https://openaccess.thecvf.com/content_CVPR_2019/html/Cui_Class-Balanced_Loss_Based_on_Effective_Number_of_Samples_CVPR_2019_paper.html'),
        ('[10] Reimers & Gurevych (2017). Reporting Score Distributions Makes a Difference. EMNLP.','https://aclanthology.org/D17-1035/'),
        ('[11] TorchVision 0.17. Models and pre-trained weights. Tài liệu chính thức.','https://docs.pytorch.org/vision/0.17/models.html'),
        ('[12] Cawley & Talbot (2010). On Over-fitting in Model Selection and Subsequent Selection Bias in Performance Evaluation. JMLR.','https://www.jmlr.org/papers/v11/cawley10a.html'),
    ]
    for text,url in refs:
        p=para(doc,text+' ','BriefSource')
        link(p,'Mở nguồn chính thức',url)
    para(doc,'Trong Word desktop: Ctrl + nhấp liên kết. Hướng dẫn giữ test độc lập xem [12]. Nguồn mở rộng và phương pháp chi tiết nằm trong báo cáo phân tích đầy đủ.','BriefSource')
    para(doc,'Nguồn số liệu: final_test_v2/final_comparison.json, test_metrics.json của 18 run; training_v2/run_manifest.json; benchmark_contract_v2.json. Biểu đồ sử dụng cùng dữ liệu với bản chi tiết đã kiểm tra.','BriefSource')

    # Ensure plot bytes match the reviewed detailed report, not merely the filenames.
    with zipfile.ZipFile(SOURCE) as archive:
        source_media={hashlib.sha256(archive.read(n)).hexdigest() for n in archive.namelist() if n.startswith('word/media/')}
    figures=['fig_overall_metrics.png','fig_seed_stability.png','fig_per_class_heatmap.png','fig_confusion_best_mobile.png','fig_paired_bootstrap.png','fig_quality_params.png']
    for name in figures:
        assert sha(ASSETS/name) in source_media, 'Chart differs from reviewed source: '+name
    doc.save(OUTPUT)
    assert sha(SOURCE)==source_sha,'Detailed source was changed'
    audit_out={'source_sha256':source_sha,'output_sha256':sha(OUTPUT),'validated_runs':18,'test_images_per_run':9829,'classes':19,'aggregate_recomputed_from_confusion':True,'source_figures_verified':figures,'redrawn_figure':'seed stability redrawn from validated metrics to remove stale source table numbers','tables':len(doc.tables),'hyperlinks':12,'style_preset':'standard_business_brief','header_pattern':'memo_masthead','named_overrides':{'page_leading_heading_before_pt':0,'BriefTable':'8–9pt, line1.05, after0','BriefSource':'8pt, line1, after3','BriefFlow':'10pt blue, before3 after3','figures':'420dpi,6.35in;confusion5.8in'},'source_retained_unchanged':True}
    (QA/'content_audit.json').write_text(json.dumps(audit_out,ensure_ascii=False,indent=2),encoding='utf-8')
    print('PRELIMINARY REPORT CREATED')
    print('18 runs checked against confusion matrices; aggregate mean/sample SD checked.')
    print('Source figures hash-verified; seed plot redrawn from validated metrics; source DOCX unchanged.')
    print(OUTPUT)


if __name__=='__main__':
    build()
