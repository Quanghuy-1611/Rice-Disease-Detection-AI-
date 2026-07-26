import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def add_heading(doc, text, level=1):
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.name = 'Times New Roman'
        # Word defaults to some blue for headings, let's keep it default or force black
        if level > 0:
            run.font.color.rgb = docx.shared.RGBColor(0, 0, 0)
    return heading

def add_paragraph(doc, text, bold=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(13)
    run.bold = bold
    return p

def create_thesis_word():
    doc = docx.Document()

    # Thêm tiêu đề trang bìa
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("\n\nBÁO CÁO ĐỒ ÁN TỐT NGHIỆP\n")
    title_run.font.name = 'Times New Roman'
    title_run.font.size = Pt(20)
    title_run.bold = True

    project_name = doc.add_paragraph()
    project_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_run = project_name.add_run("ỨNG DỤNG TRÍ TUỆ NHÂN TẠO TRONG NHẬN DIỆN\nVÀ CHẨN ĐOÁN BỆNH TRÊN LÁ LÚA\n")
    name_run.font.name = 'Times New Roman'
    name_run.font.size = Pt(24)
    name_run.bold = True
    
    doc.add_page_break()

    # CHƯƠNG 1
    add_heading(doc, "CHƯƠNG 1: TỔNG QUAN ĐỀ TÀI", level=1)
    
    add_heading(doc, "1.1. Đặt vấn đề và lý do chọn đề tài", level=2)
    add_paragraph(doc, "Việt Nam là một trong những quốc gia xuất khẩu gạo hàng đầu thế giới, với diện tích đất nông nghiệp trồng lúa nước chiếm tỷ trọng vô cùng lớn. Tuy nhiên, hàng năm ngành nông nghiệp luôn phải đối mặt với những thiệt hại kinh tế nặng nề do dịch bệnh trên lúa gây ra. Điển hình là các bệnh nguy hiểm như Đạo ôn (Blast), Bạc lá vi khuẩn (Bacterial Leaf Blight), hay bệnh Than đen (Leaf Smut). Nếu không được phát hiện và xử lý kịp thời, dịch bệnh có thể lây lan nhanh chóng, làm giảm năng suất và chất lượng hạt gạo.")
    add_paragraph(doc, "Hiện nay, việc chẩn đoán bệnh thường dựa vào kinh nghiệm mắt thường của người nông dân hoặc phải chờ đợi sự tư vấn từ các chuyên gia nông nghiệp, kỹ sư bảo vệ thực vật. Quy trình này bộc lộ nhiều hạn chế: tốc độ chậm, khả năng chẩn đoán sai sót cao và thiếu tính phổ quát, đặc biệt ở các vùng sâu vùng xa.")
    add_paragraph(doc, "Trước sự bùng nổ của Công nghệ thông tin và Trí tuệ nhân tạo (AI), việc ứng dụng công nghệ Deep Learning vào nông nghiệp thông minh đang trở thành xu hướng tất yếu. Một hệ thống chẩn đoán tự động thông qua hình ảnh chụp từ điện thoại di động sẽ giúp bà con nông dân phát hiện bệnh ngay lập tức, từ đó đưa ra phác đồ điều trị chính xác, giảm thiểu thiệt hại và bảo vệ mùa màng. Đó chính là lý do và động lực để nhóm thực hiện đề tài này.")

    add_heading(doc, "1.2. Mục tiêu đề tài", level=2)
    add_paragraph(doc, "- Xây dựng và huấn luyện thành công mô hình Trí tuệ nhân tạo (CNN) có khả năng phân loại chính xác 11 trường hợp trên lá lúa (10 loại bệnh phổ biến và 1 trường hợp lá khỏe mạnh).")
    add_paragraph(doc, "- Xây dựng thuật toán xác định và khoanh vùng tự động (Bounding Box) vị trí vết bệnh trên lá lúa nhằm tăng tính trực quan và khả năng giải thích (Explainable AI) của hệ thống.")
    add_paragraph(doc, "- Thiết kế và phát triển ứng dụng Web hoàn chỉnh, cho phép người dùng tải ảnh lên, nhận kết quả chẩn đoán tức thời, xem phác đồ điều trị chi tiết và xuất phiếu báo cáo chẩn đoán dưới dạng PDF chuyên nghiệp.")

    add_heading(doc, "1.3. Đối tượng và phạm vi nghiên cứu", level=2)
    add_paragraph(doc, "- Đối tượng nghiên cứu: Hình ảnh các loại lá cây lúa bị nhiễm bệnh và lá khỏe mạnh được chụp trong điều kiện thực tế ngoài đồng ruộng.")
    add_paragraph(doc, "- Phạm vi dữ liệu: Hệ thống được huấn luyện trên một tập dữ liệu quy mô lớn gồm 24.272 hình ảnh, được tổng hợp, làm sạch và dán nhãn từ các nguồn dữ liệu uy tín (như tập dữ liệu Paddy Doctor và Kaggle).")

    add_heading(doc, "1.4. Bố cục của báo cáo", level=2)
    add_paragraph(doc, "Báo cáo được chia thành 5 chương chính: Chương 1 giới thiệu tổng quan; Chương 2 trình bày cơ sở lý thuyết về AI và xử lý ảnh; Chương 3 phân tích thiết kế hệ thống; Chương 4 mô tả quá trình triển khai, huấn luyện mô hình và xây dựng giao diện; Chương 5 tổng kết những kết quả đạt được và hướng phát triển.")

    # CHƯƠNG 2
    add_heading(doc, "CHƯƠNG 2: CƠ SỞ LÝ THUYẾT", level=1)
    
    add_heading(doc, "2.1. Tổng quan về Xử lý ảnh số (Digital Image Processing)", level=2)
    add_paragraph(doc, "Xử lý ảnh số là quá trình thao tác trên các điểm ảnh (pixel) nhằm trích xuất thông tin hoặc nâng cao chất lượng hình ảnh. Trong đề tài này, ảnh lá lúa được biểu diễn dưới hệ màu RGB. Các kỹ thuật tiền xử lý cơ bản bao gồm đổi kích thước ảnh (Resize) để phù hợp với kích thước đầu vào của mô hình AI, và lọc nhiễu.")

    add_heading(doc, "2.2. Trí tuệ nhân tạo (AI), Học máy và Học sâu", level=2)
    add_paragraph(doc, "Học sâu (Deep Learning) là một tập con của Học máy, sử dụng các mạng nơ-ron nhân tạo với nhiều lớp lẩn (hidden layers). Điểm mạnh vượt trội của Deep Learning trong nhận diện hình ảnh là khả năng tự động trích xuất đặc trưng (Feature Extraction) từ dữ liệu thô mà không cần sự can thiệp thủ công của con người.")

    add_heading(doc, "2.3. Mạng Nơ-ron Tích Chập (CNN)", level=2)
    add_paragraph(doc, "Mạng CNN là cấu trúc học sâu tối ưu nhất cho dữ liệu hình ảnh. Các lớp tích chập (Convolutional Layers) sử dụng các bộ lọc (filters) trượt qua ảnh để phát hiện các đường nét, kết cấu và hình khối của vết bệnh. Lớp gộp (Pooling) giúp giảm chiều dữ liệu, trong khi lớp kết nối đầy đủ (Fully Connected) đóng vai trò phân loại bệnh ở bước cuối cùng.")

    add_heading(doc, "2.4. Kiến trúc MobileNetV2 và Transfer Learning", level=2)
    add_paragraph(doc, "Thay vì huấn luyện mô hình từ đầu, đề tài sử dụng phương pháp Học chuyển giao (Transfer Learning) dựa trên kiến trúc MobileNetV2. MobileNetV2 sử dụng kỹ thuật Depthwise Separable Convolution, giúp giảm thiểu số lượng tham số đáng kể nhưng vẫn duy trì độ chính xác cao. Điều này đặc biệt phù hợp để triển khai trên các hệ thống Web và thiết bị di động có tài nguyên giới hạn.")

    add_heading(doc, "2.5. Kỹ thuật Class Activation Mapping (Grad-CAM)", level=2)
    add_paragraph(doc, "Đề tài ứng dụng kỹ thuật Grad-CAM để theo dõi luồng thông tin trong mạng CNN. Grad-CAM trích xuất gradient từ lớp tích chập cuối cùng để tạo ra Bản đồ nhiệt (Heatmap). Bản đồ này biểu diễn các vùng pixel mà mạng nơ-ron tập trung vào để đưa ra dự đoán. Kết hợp cùng thuật toán OpenCV, hệ thống có thể khoanh vùng (Bounding Box) vết bệnh một cách tự động và chính xác.")

    # CHƯƠNG 3, 4, 5 (Draft points)
    add_heading(doc, "CHƯƠNG 3: PHÂN TÍCH VÀ THIẾT KẾ HỆ THỐNG", level=1)
    add_paragraph(doc, "[Đoạn này sẽ bao gồm Đặc tả yêu cầu hệ thống, Sơ đồ Use Case, Sơ đồ tuần tự và Sơ đồ kiến trúc (Client-Server-AI Engine). Sẽ chèn các hình ảnh sơ đồ UML sau].", bold=True)
    
    add_heading(doc, "CHƯƠNG 4: TRIỂN KHAI VÀ ĐÁNH GIÁ KẾT QUẢ", level=1)
    add_paragraph(doc, "[Đoạn này sẽ trình bày quá trình làm sạch 24.272 ảnh, cấu hình huấn luyện model trên TensorFlow, đánh giá Confusion Matrix, và Demo các tính năng của Giao diện (Xuất PDF, Chatbot)].", bold=True)
    
    add_heading(doc, "CHƯƠNG 5: KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN", level=1)
    add_paragraph(doc, "[Đoạn này tổng kết mức độ chính xác của model và đề xuất mở rộng thành App Mobile trong tương lai].", bold=True)

    doc.save('D:\Phát hiện bệnh cây lúa\Do_An_Nhan_Dien_Benh_Lua_Draft.docx')
    print("Done")

if __name__ == '__main__':
    create_thesis_word()
